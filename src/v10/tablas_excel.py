"""V10-B6 · output/v10/tablas_atlas.xlsx (lo copia `make publicacion` a publicacion/tablas_atlas.xlsx).

Convenciones para usuarios no técnicos:
  - datos de entrada en azul, con su fuente y su fecha en la nota de la cabecera y en una fila de fuente;
  - cálculos de la propia hoja en negro (fórmulas de Excel, recalculables);
  - vínculos a otras pestañas en verde (fórmulas que leen otra hoja, e hipervínculos del índice);
  - sin paneles inmovilizados; zip con fechas fijas (dos ejecuciones dan el mismo sha256).
No estima nada: copia tablas del atlas (output/v10/atlas) y añade cálculos aritméticos sencillos.
"""
from __future__ import annotations

import datetime as dt
import re
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[2]
ATLAS = RAIZ / "output" / "v10" / "atlas"
SALIDA = RAIZ / "output" / "v10" / "tablas_atlas.xlsx"
FECHA_FIJA = dt.datetime(2026, 10, 14)
AZUL, NEGRO, VERDE = Font(color="1F4E99"), Font(color="000000"), Font(color="1E7B34")
NEGRITA = Font(bold=True)


def fijar_zip(ruta: Path) -> None:
    with zipfile.ZipFile(ruta) as z:
        datos = [(n, z.read(n)) for n in sorted(z.namelist())]
    fijo = FECHA_FIJA.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    datos = [(n, re.sub(rb"(<dcterms:(?:modified|created)[^>]*>)[^<]*", rb"\g<1>" + fijo, b)
              if n == "docProps/core.xml" else b) for n, b in datos]
    tmp = Path(str(ruta) + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in datos:
            zi = zipfile.ZipInfo(n, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            z.writestr(zi, b)
    tmp.replace(ruta)


def hoja_datos(wb, nombre: str, df: pd.DataFrame, fuente: str, fecha: str, capa: str,
               calculos: list[tuple[str, str]] | None = None) -> None:
    """Escribe df en azul (entrada); añade columnas calculadas en negro con fórmulas por fila.
    calculos: [(cabecera, plantilla de fórmula con {col_<nombre>} y {fila})]."""
    ws = wb.create_sheet(nombre)
    ws["A1"] = f"{nombre}: datos de entrada en azul; cálculos de esta hoja en negro"
    ws["A1"].font = NEGRITA
    ws["A2"] = f"Fuente: {fuente}. Fecha del dato: {fecha}. Capa: {capa}."
    ws["A2"].font = AZUL
    cab = 4
    cols = list(df.columns)
    letra = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}
    for j, c in enumerate(cols, 1):
        cel = ws.cell(row=cab, column=j, value=c)
        cel.font = Font(bold=True, color="1F4E99")
        cel.comment = Comment(f"Entrada. Fuente: {fuente}. Fecha: {fecha}.", "V10")
    for i, fila in enumerate(df.itertuples(index=False), cab + 1):
        for j, v in enumerate(fila, 1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            ws.cell(row=i, column=j, value=v).font = AZUL
    for k, (titulo, plantilla) in enumerate(calculos or [], len(cols) + 1):
        cel = ws.cell(row=cab, column=k, value=titulo)
        cel.font = Font(bold=True)
        cel.comment = Comment("Cálculo de esta hoja (fórmula).", "V10")
        for i in range(cab + 1, cab + 1 + len(df)):
            f = plantilla.format(fila=i, **{f"col_{c}": letra[c] for c in cols})
            ws.cell(row=i, column=k, value=f).font = NEGRO
    for j in range(1, len(cols) + len(calculos or []) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16


def main() -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "Indice"
    wb.properties.creator = "vivienda-espana-evidencia V10 (src/v10/tablas_excel.py)"
    wb.properties.lastModifiedBy = None
    wb.properties.created = FECHA_FIJA
    wb.properties.modified = FECHA_FIJA
    wb.properties.title = "Tablas del atlas descriptivo (V10)"
    ip = pd.read_csv(ATLAS / "b_ipva_nuevo_existente.csv", dtype={"territorio_codigo": str})
    ip = ip[["nivel", "territorio_codigo", "nombre", "periodo", "nuevo", "existente", "total"]]
    hojas = [
        ("IPVA_nuevo_existente", ip, "INE, IPVA (59004, 59005)", "2015-2024", "C4",
         [("brecha_nuevo_menos_existente", "=IF(AND(ISNUMBER({col_nuevo}{fila}),ISNUMBER({col_existente}{fila})),{col_nuevo}{fila}-{col_existente}{fila},\"\")")]),
    ]
    p = pd.read_csv(ATLAS / "d_espana_provincias.csv", dtype={"cod_provincia": str})
    p = p[["cod_provincia", "nombre", "pob_2025", "valor_tasado_2024", "alquiler_m2_2024", "renta_neta_persona_2023",
           "deficit_2021_2025_mediana"]]
    hojas.append(("Provincias", p, "MIVAU (valor tasado, SERPAVI); INE (censo anual, ADRH); V5-A4", "2023-2025", "C4",
                  [("alquiler_80m2_pct_renta", "=IF(AND(ISNUMBER({col_alquiler_m2_2024}{fila}),ISNUMBER({col_renta_neta_persona_2023}{fila})),100*{col_alquiler_m2_2024}{fila}*80*12/{col_renta_neta_persona_2023}{fila},\"\")"),
                   ("deficit_por_1000_hab", "=IF(AND(ISNUMBER({col_deficit_2021_2025_mediana}{fila}),ISNUMBER({col_pob_2025}{fila})),1000*{col_deficit_2021_2025_mediana}{fila}/{col_pob_2025}{fila},\"\")")]))
    c = pd.read_csv(ATLAS / "d_cv_municipios.csv", dtype={"cod_municipio": str})
    c = c[["cod_municipio", "nombre", "pob_2021", "pob_2025", "alquiler_m2_2019", "alquiler_m2_2024", "vut_ine_2026_05",
           "inmuebles_residenciales_catastro"]]
    hojas.append(("CV_municipios", c, "INE censo anual; MIVAU SERPAVI; INE VUT; Catastro", "2019-2026", "C4",
                  [("var_pob_pct", "=IF(AND(ISNUMBER({col_pob_2025}{fila}),{col_pob_2021}{fila}>0),100*({col_pob_2025}{fila}/{col_pob_2021}{fila}-1),\"\")"),
                   ("var_alquiler_pct", "=IF(AND(ISNUMBER({col_alquiler_m2_2024}{fila}),ISNUMBER({col_alquiler_m2_2019}{fila})),100*({col_alquiler_m2_2024}{fila}/{col_alquiler_m2_2019}{fila}-1),\"\")"),
                   ("vut_por_100_inmuebles", "=IF(AND(ISNUMBER({col_vut_ine_2026_05}{fila}),ISNUMBER({col_inmuebles_residenciales_catastro}{fila})),100*{col_vut_ine_2026_05}{fila}/{col_inmuebles_residenciales_catastro}{fila},\"\")")]))
    b = pd.read_csv(ATLAS / "c_valencia_barrio.csv", dtype={"cod_barrio": str})
    b = b[["cod_barrio", "nombre", "pob_2021", "pob_2025", "extranj_2021", "extranj_2025", "vut_2026_05", "viviendas_2021",
           "alquiler_m2_2019", "alquiler_m2_2024", "renta_persona_2023", "indice_presion_descriptivo"]]
    hojas.append(("Valencia_barrios", b, "INE (censo anual, ADRH, VUT, Censo 2021); MIVAU SERPAVI; Ajuntament de València",
                  "2019-2026", "C4 (recuentos de población por barrio repartidos por área)",
                  [("pct_extranj_2025", "=IF({col_pob_2025}{fila}>0,100*{col_extranj_2025}{fila}/{col_pob_2025}{fila},\"\")"),
                   ("vut_por_100_viv", "=IF({col_viviendas_2021}{fila}>0,100*{col_vut_2026_05}{fila}/{col_viviendas_2021}{fila},\"\")"),
                   ("var_alquiler_pct", "=IF(AND(ISNUMBER({col_alquiler_m2_2024}{fila}),ISNUMBER({col_alquiler_m2_2019}{fila})),100*({col_alquiler_m2_2024}{fila}/{col_alquiler_m2_2019}{fila}-1),\"\")")]))
    e = pd.read_csv(ATLAS / "a_quien_ecv.csv")
    e = e[e.anio == e.anio.max()][["anio", "grupo", "categoria", "indicador", "valor", "ic95_inf", "ic95_sup", "n_muestral", "capa"]]
    hojas.append(("Quien_ECV", e, "INE ECV (microdatos), cálculo propio", f"ECV {int(e.anio.max())}", "C1 o C4 según fila (columna capa)",
                  [("amplitud_ic", "=IF(ISNUMBER({col_ic95_sup}{fila}),{col_ic95_sup}{fila}-{col_ic95_inf}{fila},\"\")")]))
    for nombre, df, fuente, fecha, capa, calc in hojas:
        hoja_datos(wb, nombre, df, fuente, fecha, capa, calc)
    # índice con vínculos (verde) y resumen que lee otras pestañas (verde)
    ws["A1"] = "Tablas del atlas descriptivo (V10). Generado por src/v10/tablas_excel.py; no editar a mano."
    ws["A1"].font = NEGRITA
    ws["A2"] = "Azul: datos de entrada (fuente y fecha en la cabecera). Negro: cálculos de cada hoja. Verde: vínculos a otras pestañas."
    ws["A4"], ws["B4"], ws["C4"] = "Hoja", "Contenido", "Filas (vínculo)"
    for c_ in ("A4", "B4", "C4"):
        ws[c_].font = NEGRITA
    for i, (nombre, df, fuente, fecha, capa, _) in enumerate(hojas, 5):
        cel = ws.cell(row=i, column=1, value=nombre)
        cel.hyperlink = f"#'{nombre}'!A1"
        cel.font = Font(color="1E7B34", underline="single")
        ws.cell(row=i, column=2, value=f"{fuente} · {fecha} · capa {capa}")
        f = ws.cell(row=i, column=3, value=f"=COUNTA('{nombre}'!A5:A100000)")
        f.font = VERDE
    fila = 5 + len(hojas) + 1
    ws.cell(row=fila, column=1, value="Brecha IPVA España, último año (vínculo)").font = NEGRITA
    n = len(hojas[0][1])
    v = ws.cell(row=fila, column=3, value=f"=SUMIFS('IPVA_nuevo_existente'!H5:H{4 + n},'IPVA_nuevo_existente'!A5:A{4 + n},\"nacional\",'IPVA_nuevo_existente'!D5:D{4 + n},MAX('IPVA_nuevo_existente'!D5:D{4 + n}))")
    v.font = VERDE
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 90
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    fijar_zip(SALIDA)
    return {"hojas": len(hojas) + 1}


if __name__ == "__main__":
    print(main())
