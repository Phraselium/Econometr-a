"""V18 · Fase 6 · Documento de Trabajo (docs/v18/plan.md, fase 6) y tablas_v18.xlsx.

1. Ensambla output/v18/documento/plantilla/[0-9][0-9]_*.md (las escribe el redactor) en output/v18/documento/documento_v18.md y
   sustituye cada marcador con las funciones de src/v5/render.py contra output/v5/cifras_clave.csv:
   - {{id}}, {{id:campo}} como en render.py (valor, min, max, rango, cita, breve, sig, periodo, cobertura, fuentes, capa…);
   - {{id:clasificacion}}: la clasificacion_v18 de la cifra (K2; columna propia de cifras_clave, al lado de la capa);
   - {{id:min_sig}}, {{id:max_sig}}: extremos del intervalo con 3 cifras significativas (como «sig»; para intervalos pequeños);
     «sig» y sus extremos nunca usan notación científica (K12);
   - {{id:tres_dec}}, {{id:un_dec}}, {{id:abs_un_dec}} (con min_/max_): 3 o 1 decimales fijos; abs_un_dec da el valor absoluto
     (la frase de limitación de E2-R1, K9, da sus cifras así); sufijo _en en sig, tres_dec, un_dec, abs_un_dec y num: punto
     decimal, para el abstract en inglés (K12);
   - {{tabla.impactos}}: la tabla de impactos del resumen ejecutivo (factor | magnitud preferida | intervalo | clasificación |
     confianza | limitación principal), con las cifras V18-B5-E* de cifras_clave y el estado B5 de output/v18/f5/magnitudes.csv.
   Falla (código 1) si algún marcador no existe en cifras_clave, si una cifra pedida con :clasificacion no la tiene o si queda
   cualquier «{{…}}» sin resolver.
2. Escribe output/v18/documento/tablas_v18.xlsx (openpyxl; fechas y zip fijos): pestañas impactos, estimaciones, magnitudes y
   contabilidad. Convención: AZUL = dato introducido (copiado de cifras_clave o de magnitudes.csv; fuente en la nota de la celda),
   NEGRO = fórmula calculada en la hoja, VERDE = vínculo a otra pestaña. Sin paneles inmovilizados.
`make publicacion` (src/v7_1/publicacion.py) copia ambos a publicacion/documento_v18/. Determinista; sin red; solo lee output/.
Uso: python3 src/v18/f6_documento.py [--ck ruta/a/cifras_clave.csv] (por defecto, output/v5/cifras_clave.csv)
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

RAIZ = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ / "src" / "v5"))
sys.path.insert(0, str(RAIZ / "src" / "v10"))
import render as R  # noqa: E402
from tablas_excel import fijar_zip  # noqa: E402

ALCANCE_TERRITORIAL = "sin lectura de datos"   # solo lee cifras_clave y salidas V18 ya generadas (guarda de fuga v6)
DIR = RAIZ / "output" / "v18" / "documento"
PLANT = DIR / "plantilla"
DOC = DIR / "documento_v18.md"
XLSX = DIR / "tablas_v18.xlsx"
CK = RAIZ / "output" / "v5" / "cifras_clave.csv"
MAG = RAIZ / "output" / "v18" / "f5" / "magnitudes.csv"
FACTORES = [f"E{k}" for k in range(1, 10)]
FECHA_FIJA = dt.datetime(2026, 10, 20)
AZUL, NEGRO, VERDE, NEGRITA = Font(color="1F4E99"), Font(color="000000"), Font(color="1E7B34"), Font(bold=True)
TABLA = re.compile(r"\{\{tabla\.([a-z_]+)\}\}")
RESTO = re.compile(r"\{\{[^{}]*\}\}")
COMENTARIO = re.compile(r"<!--.*?-->\n?", re.S)
CAB = """# Factores de la subida del precio de la vivienda en España: reproducción, réplica y magnitudes (Documento de Trabajo V18)

Borja Romero, economista (firma individual, independiente). Documento de Trabajo de la versión V18 del proyecto (docs/v18/plan.md,
fase 6). Toda cifra procede de output/v5/cifras_clave.csv con su id. Toda cifra V18 es de capa C4 (nivel de evidencia EXPLORATORIO)
y lleva al lado su clasificación V18 (IDENTIFICADO, SUGESTIVO, ASOCIACIÓN, NO CONCLUYENTE o NO ESTIMABLE; docs/v18/decisiones.md,
K2). Pre-registro interno sin sello de tiempo externo (etiqueta plan-replicacion-v18). Generado por src/v18/f6_documento.py; no se
edita a mano.
"""


# ------------------------------------------------------------------ cifras y marcadores
def leer_ck(ruta: Path = CK) -> pd.DataFrame:
    ck = pd.read_csv(ruta).set_index("id")
    if "clasificacion_v18" not in ck.columns:
        ck["clasificacion_v18"] = pd.NA
    return ck


def _sig(x, en: bool = False) -> str:
    """Tres cifras significativas con coma decimal (punto en inglés) y signo menos tipográfico, sin notación científica (K12:
    −6,5e−05 pasa a −0,000065). Presentación; la cifra es la de cifras_clave."""
    if pd.isna(x):
        return "—"
    s = f"{float(x):.3g}"
    if "e" in s:   # %g usa exponente fuera de [1e-4, 1e3): se escribe en decimal con las mismas tres cifras significativas
        v = float(s)
        dec = max(0, 2 - math.floor(math.log10(abs(v)))) if v else 0
        s = f"{v:.{dec}f}"
        s = s.rstrip("0").rstrip(".") if "." in s else s
    return s.replace("-", "−") if en else s.replace("-", "−").replace(".", ",")


def _dec(x, n: int, en: bool = False) -> str:
    """n decimales fijos (K12: la frase de limitación de E2-R1 da sus cifras a 3 decimales), signo menos tipográfico."""
    s = f"{float(x):.{n}f}"
    if float(s) == 0:
        s = s.lstrip("-")
    return s.replace("-", "−") if en else s.replace("-", "−").replace(".", ",")


# Campos de marcador propios del Documento de Trabajo (además de los de render.py). Sufijo «_en»: punto decimal (abstract en inglés).
CAMPOS_PROPIOS = re.compile(r"^(min_|max_)?(sig|tres_dec|un_dec|abs_un_dec)(_en)?$|^(min_|max_)?num_en$")


def factor_nombre(ind: str) -> str:
    m = re.match(r"Magnitud preferida de (.+?) \((E\d+)\)", str(ind))
    return f"{m.group(1)[:1].upper()}{m.group(1)[1:]} ({m.group(2)})" if m else str(ind)


def confianza(estado: str, n: int) -> str:
    """Regla fija (sección 3 del documento): media solo con magnitud preferida de 3 o más estudios; muy baja sin error típico en
    la literatura; baja en el resto. Toda cifra V18 es C4: nunca alta."""
    if estado.startswith("magnitud preferida") and n >= 3:
        return "media"
    if estado.startswith("rango de la literatura"):
        return "muy baja"
    return "baja"


def limitacion(m: pd.Series) -> str:
    estado = str(m.estado_b5)
    if estado.startswith("magnitud de un solo estudio"):
        base = "un solo estudio en la unidad"
    elif estado.startswith("rango sin magnitud preferida (la decisión"):
        base = "la decisión cambia con la prior: rango sin magnitud preferida"
    elif estado.startswith("rango sin magnitud preferida"):
        base = "heterogeneidad alta entre estudios: rango sin magnitud preferida"
    elif estado.startswith("rango de la literatura"):
        lim = str(m.limitacion)
        n = re.search(r"rango de (un solo estudio|\d+ estudios)", lim)   # K12 (árbitro, hallazgo 10): cuántos estudios forman el rango
        base = "sin error típico en la literatura: " + (f"{n.group(0)}, sin intervalo estadístico" if n else "rango sin intervalo estadístico")
    else:
        base = "pocos estudios con error típico; heterogeneidad declarada"
    propia = str(m.clase_estimacion_propia)
    return f"{base}; estimación propia: {propia}" if propia and propia != "nan" else base


def tabla_impactos(ck: pd.DataFrame, mag: pd.DataFrame) -> tuple[str, list[str]]:
    falta = [f"V18-B5-{f}" for f in FACTORES if f"V18-B5-{f}" not in ck.index]
    if falta:
        return "", falta
    m = mag.set_index("factor")
    L = ["| Factor | Magnitud preferida | Intervalo | Clasificación | Confianza | Limitación principal |", "|---|---|---|---|---|---|"]
    unidades = []
    for f in FACTORES:
        r, mm = ck.loc[f"V18-B5-{f}"], m.loc[f]
        hay = pd.notna(r.valor) and str(mm.estado_b5).startswith("magnitud")
        mag_txt = _sig(r.valor) if hay else "sin magnitud preferida"
        k = int(mm.n_unidades_meta)
        if hay:   # K12 (árbitro, hallazgo 8): con dos estudios el efecto fijo no da un IC95 fiable (Monte Carlo de meta_bayes)
            tipo = "IC95" if k >= 3 else ("IC95 del único estudio" if k == 1 else "intervalo del efecto fijo con dos estudios; no es un IC95 fiable")
        else:
            tipo = "rango"
        inter = "—" if pd.isna(r["min"]) or pd.isna(r["max"]) else f"[{_sig(r['min'])}; {_sig(r['max'])}] ({tipo})"
        L.append(f"| {factor_nombre(r.indicador)} | {mag_txt} | {inter} | {r.clasificacion_v18} | "
                 f"{confianza(str(mm.estado_b5), int(mm.n_unidades_meta))} | {limitacion(mm)} |")
        unidades.append(f"{f}: {r.unidad}")
    L += ["", "Unidades: " + "; ".join(unidades) + ". Cifras: ids V18-B5-E1 a V18-B5-E9 de cifras_clave (capa C4). La clasificación es la de "
          "la estimación propia designada del factor; la magnitud preferida puede venir solo de la literatura (sección 4). "
          "Confianza: media solo con magnitud preferida de tres o más estudios; muy baja sin error típico en la literatura; baja en el resto."]
    return "\n".join(L), []


def render_texto(txt: str, ck: pd.DataFrame, mag: pd.DataFrame) -> tuple[str, list[str]]:
    falta: list[str] = []
    retiradas, odbl, usa_odbl = R.T10.retiradas(), R.T10.odbl(), []

    def tabla(m: re.Match) -> str:
        if m.group(1) != "impactos":
            falta.append(m.group(0))
            return m.group(0)
        t, f = tabla_impactos(ck, mag)
        falta.extend(f)
        return t or m.group(0)

    def sub(m: re.Match) -> str:
        i, campo = m.group(1), m.group(2)
        if i.startswith("proyecto.") and campo is None:
            v = R.PROYECTO.get(i.split(".", 1)[1])
            if isinstance(v, str):
                return v
            falta.append(i)
            return m.group(0)
        if i not in ck.index:
            falta.append(i)
            return m.group(0)
        if i in retiradas:
            return R.T10.marca(retiradas[i])
        if i in odbl:
            usa_odbl.append(i)
        r = ck.loc[i]
        if campo == "clasificacion":
            if pd.isna(r.clasificacion_v18):
                falta.append(f"{i}:clasificacion")
                return m.group(0)
            return str(r.clasificacion_v18)
        if campo and CAMPOS_PROPIOS.match(campo):
            # sig, tres_dec, un_dec, abs_un_dec, con prefijo min_/max_ (extremo del intervalo) y sufijo _en (punto decimal)
            en = campo.endswith("_en")
            base = campo[:-3] if en else campo
            col = "valor"
            if base.startswith(("min_", "max_")):
                col, base = base[:3], base[4:]
            if base == "num":   # num_en: el «num» de render.py con los separadores ingleses
                r2 = r.copy()
                r2["id"] = i
                s = R.texto(r2, "num" if col == "valor" else f"{col}_num")
                return s.translate(str.maketrans({",": ".", ".": ","})).replace("-", "−")
            x = r[col]
            if pd.isna(x):
                falta.append(f"{i}:{campo}")
                return m.group(0)
            if base == "sig":
                return _sig(x, en)
            if base == "tres_dec":
                return _dec(x, 3, en)
            if base == "un_dec":
                return _dec(x, 1, en)
            return _dec(abs(float(x)), 1, en)   # abs_un_dec: valor absoluto con un decimal («un 3,8 % por debajo»)
        r = r.copy()
        r["id"] = i   # render.muestra_rango lee el id de la fila (V17/V18: rango solo si tipo_rango = «rango»)
        return R.texto(r, campo)

    out = R.MARC.sub(sub, TABLA.sub(tabla, txt))
    falta += [f"sin resolver: {x}" for x in RESTO.findall(out) if x not in falta]
    if usa_odbl:
        out = out.rstrip("\n") + "\n\n---\n\n" + R.T10.nota_odbl() + "\n"
    return out, falta


def ensamblar(ck: pd.DataFrame, mag: pd.DataFrame) -> tuple[str, dict]:
    partes = sorted(PLANT.glob("[0-9][0-9]_*.md"))
    trozos, falta, ids = [CAB], [], set()
    for p in partes:
        bruto = p.read_text(encoding="utf-8")
        txt, f = render_texto(COMENTARIO.sub("", bruto), ck, mag)   # las instrucciones al redactor («<!-- … -->») no se publican
        txt = re.sub(r"\n{3,}", "\n\n", txt)
        falta += [f"{p.name}: {i}" for i in f]
        ids |= {m.group(1) for m in R.MARC.finditer(bruto) if not m.group(1).startswith("proyecto.")}
        trozos.append(txt.strip() + "\n")
    texto = "\n\n".join(trozos)
    usados = sorted(i for i in ids if i in ck.index)
    res = dict(rama="V18-F6", pregunta="Documento de Trabajo V18 (fase 6)", capa="C4 (cifras V18) y la de cada cifra citada",
               partes=[p.name for p in partes], marcadores=len(ids),
               ids_por_capa={str(k): int(v) for k, v in ck.loc[usados, "capa"].value_counts().items()},
               ids_v18_por_clasificacion={str(k): int(v) for k, v in ck.loc[[i for i in usados if i.startswith("V18-")], "clasificacion_v18"]
                                          .value_counts().items()},
               marcadores_sin_cifra=falta, notas="Ensamblado determinista; las cifras vienen de cifras_clave.")
    return texto, res


# ------------------------------------------------------------------ libro de tablas
def _cabecera(ws, titulo: str, fuente: str, columnas: list[tuple[str, Font]]) -> None:
    ws.append([titulo])
    ws["A1"].font = NEGRITA
    ws.append([f"Fuente: {fuente}. AZUL: dato introducido (fuente en la nota de la celda); NEGRO: fórmula de esta hoja; "
               "VERDE: vínculo a otra pestaña."])
    ws["A2"].font = AZUL
    ws.append([c for c, _ in columnas])
    for celda, (_, color) in zip(ws[3], columnas):
        celda.font = Font(bold=True, color=color.color)


def _nota(r: pd.Series, i: str) -> Comment:
    return Comment(f"Fuente: {r.fuentes}. Origen: {r.origen}. Capa {r.capa}; clasificación V18: {r.clasificacion_v18}.", "src/v18/f6_documento.py")


def _num(x):
    return None if pd.isna(x) else float(x)


def hoja_cifras(wb, nombre: str, v: pd.DataFrame, fuente: str) -> None:
    """Cifras V18 de cifras_clave en azul (nota con la fuente en el id); ancho del intervalo y si contiene el cero en negro."""
    ws = wb.create_sheet(nombre)
    cols = ["id", "indicador", "valor", "min", "max", "unidad", "periodo", "cobertura", "capa", "clasificacion_v18", "tipo_rango"]
    _cabecera(ws, f"{nombre}: cifras V18 de output/v5/cifras_clave.csv", fuente,
              [(c, AZUL) for c in cols] + [("ancho (max - min)", NEGRO), ("intervalo contiene el 0", NEGRO)])
    for i, r in v.iterrows():
        fila = [i, r.indicador, _num(r.valor), _num(r["min"]), _num(r["max"]), r.unidad, r.periodo, r.cobertura, r.capa,
                r.clasificacion_v18, r.tipo_rango]
        ws.append([None if (isinstance(x, float) and pd.isna(x)) else x for x in fila])
        n = ws.max_row
        for c in ws[n]:
            c.font = AZUL
        ws.cell(row=n, column=1).comment = _nota(r, i)
        ws.cell(row=n, column=12, value=f'=IF(AND(ISNUMBER(D{n}),ISNUMBER(E{n})),E{n}-D{n},"")').font = NEGRO
        ws.cell(row=n, column=13, value=f'=IF(AND(ISNUMBER(D{n}),ISNUMBER(E{n})),IF(AND(D{n}<=0,E{n}>=0),"sí","no"),"")').font = NEGRO


def libro(ck: pd.DataFrame, mag: pd.DataFrame) -> Workbook:
    v18 = ck[ck.index.str.startswith("V18-")]
    wb = Workbook()
    wi = wb.active
    wi.title = "impactos"
    # magnitudes (azul): V18-B5-E* de cifras_clave y estado B5 de magnitudes.csv
    wm = wb.create_sheet("magnitudes")
    cm = [("factor", AZUL), ("id", AZUL), ("indicador", AZUL), ("valor", AZUL), ("min", AZUL), ("max", AZUL), ("unidad", AZUL),
          ("clasificacion_v18", AZUL), ("estado_b5", AZUL), ("n_estudios_meta", AZUL), ("limitacion", AZUL),
          ("ancho (max - min)", NEGRO)]
    _cabecera(wm, "magnitudes: magnitud preferida por factor (B5)", "output/v5/cifras_clave.csv (V18-B5-E*) y output/v18/f5/magnitudes.csv", cm)
    m = mag.set_index("factor")
    for f in FACTORES:
        i = f"V18-B5-{f}"
        r, mm = ck.loc[i], m.loc[f]
        wm.append([f, i, r.indicador, _num(r.valor), _num(r["min"]), _num(r["max"]), r.unidad, r.clasificacion_v18, mm.estado_b5,
                   int(mm.n_unidades_meta), mm.limitacion])
        n = wm.max_row
        for c in wm[n]:
            c.font = AZUL
        wm.cell(row=n, column=2).comment = _nota(r, i)
        wm.cell(row=n, column=9).comment = Comment("Fuente: output/v18/f5/magnitudes.csv (estado_b5, n_unidades_meta, limitacion).",
                                                   "src/v18/f6_documento.py")
        wm.cell(row=n, column=12, value=f'=IF(AND(ISNUMBER(E{n}),ISNUMBER(F{n})),F{n}-E{n},"")').font = NEGRO
    # impactos (verde: vínculos a «magnitudes»; negro: confianza calculada en la hoja)
    ci = [("factor", VERDE), ("magnitud preferida", VERDE), ("min", VERDE), ("max", VERDE), ("clasificación", VERDE),
          ("estado B5", VERDE), ("n estudios", VERDE), ("confianza", NEGRO), ("limitación principal", VERDE)]
    _cabecera(wi, "impactos: tabla del resumen ejecutivo del Documento de Trabajo V18", "pestaña magnitudes (vínculos)", ci)
    for k, _ in enumerate(FACTORES):
        fm, n = 4 + k, wi.max_row + 1
        for col, ref in enumerate(["A", "D", "E", "F", "H", "I", "J"], start=1):
            wi.cell(row=n, column=col, value=f"=magnitudes!{ref}{fm}").font = VERDE
        wi.cell(row=n, column=8, value=f'=IF(AND(LEFT(F{n},18)="magnitud preferida",G{n}>=3),"media",'
                                       f'IF(LEFT(F{n},22)="rango de la literatura","muy baja","baja"))').font = NEGRO
        wi.cell(row=n, column=9, value=f"=magnitudes!K{fm}").font = VERDE
    est = v18[v18.index.str.match(r"^V18-(E|F)\d")]
    hoja_cifras(wb, "estimaciones", est, "output/v5/cifras_clave.csv (ids V18-E* y V18-F*; origen output/v18/f6/hechos.json)")
    con = v18[v18.index.str.startswith(("V18-C5-", "V18-C5I-"))]
    hoja_cifras(wb, "contabilidad", con, "output/v5/cifras_clave.csv (ids V18-C5-*: contabilidad principal; V18-C5I-*: ilustración "
                "con magnitudes de la literatura, no es una atribución)")
    for ws in wb.worksheets:
        ws.freeze_panes = None
    wb.properties.creator = "vivienda-espana-evidencia V18 (src/v18/f6_documento.py)"
    wb.properties.lastModifiedBy = None
    wb.properties.created = FECHA_FIJA
    wb.properties.modified = FECHA_FIJA
    wb.properties.title = "Tablas V18"
    return wb


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--ck", type=Path, default=CK, help="cifras_clave.csv (por defecto output/v5/cifras_clave.csv)")
    a = ap.parse_args(argv)
    ck, mag = leer_ck(a.ck), pd.read_csv(MAG)
    texto, res = ensamblar(ck, mag)
    if res["marcadores_sin_cifra"] or not res["partes"]:
        print(f"f6_documento: {len(res['partes'])} partes; {len(res['marcadores_sin_cifra'])} marcadores sin cifra")
        for x in res["marcadores_sin_cifra"][:40]:
            print("  -", x)
        return 1
    DIR.mkdir(parents=True, exist_ok=True)
    DOC.write_text(texto, encoding="utf-8")
    (DIR / "resultado.json").write_text(json.dumps(res, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    wb = libro(ck, mag)
    wb.save(XLSX)
    fijar_zip(XLSX)
    print(f"f6_documento: {len(res['partes'])} partes, {res['marcadores']} cifras distintas -> {DOC.relative_to(RAIZ)}; "
          f"{XLSX.relative_to(RAIZ)} ({', '.join(ws.title for ws in wb.worksheets)})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
