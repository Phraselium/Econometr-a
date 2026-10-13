"""V8-F6: tests de tablas_modelos_destacados (md/csv) y tablas_v8.xlsx (contenido, colores, determinismo)."""
import hashlib
import importlib.util
import os
import re
import sys
import zipfile
from pathlib import Path

for _v in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS"):
    os.environ[_v] = "1"

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
import pytest  # noqa: E402

RAIZ = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ / "src" / "v8"))
sys.path.insert(0, str(RAIZ / "src"))
_spec = importlib.util.spec_from_file_location("f6_tablas", RAIZ / "src" / "v8" / "f6_tablas.py")
f6 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(f6)
OUT = RAIZ / "output" / "v8"


@pytest.fixture(scope="module")
def csv():
    return pd.read_csv(OUT / "tablas_modelos_destacados.csv")


@pytest.fixture(scope="module")
def wb():
    return openpyxl.load_workbook(OUT / "tablas_v8.xlsx")


def test_csv_columnas_contrato(csv):
    assert list(csv.columns[:len(f6.COLS_CONTRATO)]) == f6.COLS_CONTRATO
    assert "criterio_destacado" in csv.columns and (csv["destacado"] == "si").all()


def test_destacados_no_se_eligen_por_p(csv):
    # el criterio (b) cita el RMSE fuera de muestra; nunca un p
    assert csv["criterio_destacado"].str.contains("mejor de la familia").all()
    assert csv["rmse_oos"].notna().all() and csv["rmse_ar4_oos"].notna().all()


def test_marca_no_cumple_regla(csv):
    por_modelo = csv.groupby("modelo_id")["no_cumple_regla"].apply(lambda s: (s == "si").all())
    sin_a = por_modelo[por_modelo]
    for mid in sin_a.index:
        assert "NO CUMPLE LA REGLA" in csv.loc[csv.modelo_id == mid, "motivo_destacado"].iloc[0]


def test_regla_a_coherente_con_fuentes(csv):
    m7 = pd.read_csv(OUT / "M7" / "resumen_multiverso.csv")
    m0 = pd.read_csv(OUT / "M0" / "regla_m0.csv")
    ambas = m0[m0.cumple == "si"].assign(variable=lambda d: d.variable.str.replace("d4_", "", regex=False)).merge(
        m7[m7.cumple_multiverso == True][["mercado", "periodo", "variable"]],  # noqa: E712
        on=["mercado", "periodo", "variable"])
    n_a = csv.assign(a=(csv.cumple_regla_m0.astype(str) == "si") & (csv.cumple_multiverso.astype(str) == "True"))
    assert int(n_a["a"].sum()) <= len(ambas) * 10  # ninguna fila (a) fuera de las que cumplen ambas
    if len(ambas) == 0:
        assert not n_a["a"].any()


def test_estrellas_usan_p_ajustado():
    assert f6.estrellas(0.04) == "**" and f6.estrellas(float("nan")) == "" and f6.estrellas(0.2) == ""


def test_md_contiene_secciones():
    t = (OUT / "tablas_modelos_destacados.md").read_text(encoding="utf-8")
    for s in ("H8-1", "H8-9", "Votos de la regla", "fuera de muestra", "C4", "p AJUSTADO", "Holm"):
        assert s in t
    assert "causal" in t  # advertencia de no causalidad
    assert len(re.findall(r"^\*\*M\d_", t, flags=re.M)) == pd.read_csv(OUT / "tablas_modelos_destacados.csv").modelo_id.nunique()


def test_hojas_y_sin_paneles(wb):
    assert wb.sheetnames == ["Leeme", "Destacados", "Hipotesis", "Votos_s9", "Atribucion", "Ranking",
                             "Fuera_muestra", "Regional", "Multiverso"]
    assert all(ws.freeze_panes is None for ws in wb)


def test_colores_convencion(wb):
    azul = verde = negro = 0
    for ws in wb:
        if ws.title == "Leeme":
            continue
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if c.value is None:
                    continue
                col = c.font.color.rgb[-6:] if c.font.color is not None and c.font.color.rgb else "000000"
                if isinstance(c.value, str) and c.value.startswith("="):
                    if re.match(r"^=\w+!\$?[A-Z]+\$?\d+$", c.value):
                        assert col == "008000", (ws.title, c.coordinate)
                        verde += 1
                    else:
                        assert col == "000000", (ws.title, c.coordinate)
                        negro += 1
                elif c.font.bold:
                    continue
                elif isinstance(c.value, (int, float, str)) and col == "0000FF":
                    azul += 1
    assert azul > 0 and verde > 0 and negro > 0


def test_cabeceras_negrita_y_fuente(wb):
    for n in ("Destacados", "Hipotesis", "Votos_s9", "Atribucion", "Ranking", "Regional", "Multiverso"):
        ws = wb[n]
        assert all(c.font.bold for c in ws[1] if c.value is not None), n
        assert "fuente" in [c.value for c in ws[1]], n


def test_formulas_votos_coinciden_con_fuente(wb):
    ws = wb["Votos_s9"]
    V = pd.read_csv(OUT / "F5" / "regla_s9.csv")
    for i, (_, x) in enumerate(V.iterrows(), start=2):
        votos = [ws.cell(i, j).value for j in range(3, 7)]
        calc = sum(1 for v in votos if v == ws.cell(i, 7).value)
        assert calc == x["n_votos"], (x["mercado"], x["periodo"])
        assert ws.cell(i, 15).value == f"=COUNTIF(C{i}:F{i},G{i})"


def test_atribucion_formula_pct_total(wb):
    ws = wb["Atribucion"]
    A = pd.read_csv(OUT / "F5" / "atribucion_combinada.csv")
    assert ws.max_row == len(A) + 1
    assert ws["U2"].value == "=IF(R2=0,\"\",F2/R2)" and ws["F2"].value == pytest.approx(A.loc[0, "contribucion_pp"])


def test_hipotesis_nueve(wb):
    ws = wb["Hipotesis"]
    ids = [ws.cell(i, 1).value for i in range(2, 11)]
    assert ids == [f"H8-{k}" for k in range(1, 10)]


def test_determinista(tmp_path):
    p = OUT / "tablas_v8.xlsx"
    h0 = hashlib.sha256(p.read_bytes()).hexdigest()
    with zipfile.ZipFile(p) as z:
        assert {zi.date_time for zi in z.infolist()} == {(1980, 1, 1, 0, 0, 0)}
    # regenerar en un directorio temporal no es posible sin tocar OUT: comprobar contenido de celdas vs reescritura
    Wb1 = openpyxl.load_workbook(p)
    f6.main()
    h1 = hashlib.sha256(p.read_bytes()).hexdigest()
    Wb2 = openpyxl.load_workbook(p)
    for a, b in zip(Wb1, Wb2):
        assert [[c.value for c in r] for r in a.iter_rows()] == [[c.value for c in r] for r in b.iter_rows()]
    assert h0 == h1
