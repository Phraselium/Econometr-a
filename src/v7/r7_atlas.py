"""R7 V7: entregables del atlas regional. Determinista, sin red, solo lee ficheros versionados."""
import re
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

RAIZ = Path(__file__).resolve().parents[2]
O = RAIZ / "output" / "v7"
AT = O / "atlas"
PANEL = RAIZ / "data/processed/v7/panel_ccaa_q_v7.csv"
PROV = RAIZ / "data/processed/v7/panel_prov_q_v7.csv"
VLC = RAIZ / "data/processed/valencia.csv"
PERIODOS = [("2014T1-2019T4", "2014Q1", "2019Q4"), ("2020T1-2022T2", "2020Q1", "2022Q2"),
            ("2022T3-último", "2022Q3", None)]
ROT = "informativo, muestra ya vista en v1"
MERC = ["alquiler", "compra"]
AZUL, VERDE, NEGRO = "0000FF", "008000", "000000"


def slug(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")


def pct(x, d=1):
    return "NA" if pd.isna(x) else f"{x:+.{d}f} %".replace(".", ",")


def num(x, d=2):
    return "NA" if pd.isna(x) else f"{x:.{d}f}".replace(".", ",")


def rd(p, **k):
    return pd.read_csv(p, **k)


def cargar():
    d = {}
    d["panel"] = rd(PANEL, dtype={"codigo_ine_ccaa": str})
    d["vel"] = {m: rd(O / f"R4/{m}/velocidad_ajuste_region.csv", dtype={"codigo_ine_ccaa": str}) for m in MERC}
    d["diag"] = {m: rd(O / f"R4/{m}/diagnosticos_region.csv", dtype={"codigo_ine_ccaa": str}) for m in MERC}
    d["des"] = {m: rd(O / f"R4/{m}/descomposicion.csv", dtype={"codigo_ine_ccaa": str}) for m in MERC}
    d["cp"] = {m: rd(O / f"R4/{m}/corto_plazo_region.csv", dtype={"codigo_ine_ccaa": str}) for m in MERC}
    d["hip"] = rd(O / "R4/hipotesis_confirmatorias.csv")
    d["best"] = rd(O / "R5/mejor_por_region.csv", dtype={"codigo_ine_ccaa": str})
    d["imp"] = rd(O / "R6/impactos.csv", dtype={"cc": str})
    d["rob"] = rd(O / "R6/robustas_por_ccaa.csv", dtype={"cc": str})
    d["pv"] = rd(O / "R4b/provincia_valencia.csv", dtype={"cod_prov": str})
    d["prov"] = rd(PROV, dtype={"cod_prov": str})
    d["vlc"] = rd(VLC)
    return d


def valor(df, col, q):
    s = df.loc[df.trimestre == q, col]
    return float(s.iloc[0]) if len(s) and pd.notna(s.iloc[0]) else float("nan")


def variaciones(panel, cc):
    """Devuelve filas (serie, periodo, q0, v0, q1, v1) para IPV e IPC alquiler."""
    p = panel[panel.codigo_ine_ccaa == cc].sort_values("trimestre")
    ult = p.dropna(subset=["ipv", "ipc_alquiler"]).trimestre.max()
    filas = []
    for col, nom in [("ipv", "IPV (INE)"), ("ipc_alquiler", "IPC alquiler (INE)")]:
        for per, a, b in PERIODOS:
            b = b or ult
            filas.append((nom, per, a, valor(p, col, a), b, valor(p, col, b)))
    return filas, ult


def split_region(path):
    txt = path.read_text(encoding="utf-8")
    cab, *specs = re.split(r"\n(?=## Especificacion)", txt)
    out = []
    for s in specs:
        pre, _, des = s.partition("Descomposicion historica")
        out.append((pre.strip(), ("Descomposicion historica" + des).strip() if des else ""))
    return cab.strip().splitlines()[0], out


def dem(t):
    return re.sub(r"^## ", "#### ", t, flags=re.M)


def ficha(d, cc, nombre):
    L = []
    a = L.append
    filas, ult = variaciones(d["panel"], cc)
    a(f"# Atlas regional: {nombre} ({cc})\n")
    a(f"*Capa C4 (fuente única INE a nivel de CCAA; {ROT}). Asociación, sin lectura causal. "
      f"Datos hasta {ult}; panel_ccaa_q_v7.*\n")
    a("## 1. Qué subió y cuánto\n")
    a("| serie | periodo | inicio | valor inicio | fin | valor fin | variación |\n|---|---|---|---|---|---|---|")
    for s, per, q0, v0, q1, v1 in filas:
        a(f"| {s} | {per} | {q0} | {num(v0)} | {q1} | {num(v1)} | {pct((v1 / v0 - 1) * 100)} |")
    a("\nVariación = valor fin / valor inicio − 1, sobre el índice trimestral de cada fecha. "
      "El IPC de alquiler mide el stock de contratos; no es el alquiler de contratos nuevos.\n")
    a("## 2. Ecuación y tabla de cada modelo (ECM regional)\n")
    a("Los coeficientes de corto plazo tienen inferencia con control BH (q = 0,10). La velocidad de ajuste "
      "(φ), la semivida y θ se muestran como estimaciones **sin inferencia válida**: R4 no confirmó ninguna "
      "de las 7 hipótesis de largo plazo.\n")
    desc = {}
    for m in MERC:
        _, specs = split_region(O / f"R4/{m}/region_{cc}.md")
        a(f"### Mercado: {m}\n")
        for pre, des in specs:
            a(dem(pre) + "\n")
            desc.setdefault(m, []).append(des)
        v = d["vel"][m]
        v = v[v.codigo_ine_ccaa == cc]
        a("| spec | φ (sin inferencia válida) | semivida (trim.) |\n|---|---|---|")
        for _, r in v.iterrows():
            a(f"| {r.spec} | {num(r.phi, 3)} | {num(r.semivida_trim, 1)} |")
        a("")
    a("## 3. Diagnósticos\n")
    a("| mercado | spec | Breusch-Godfrey p | Breusch-Pagan p | Jarque-Bera p | RESET p | CUSUM p | VIF máx | R² aj. | "
      "φ mitad 1 | φ mitad 2 | cambio de signo de φ |\n|---|---|---|---|---|---|---|---|---|---|---|---|")
    for m in MERC:
        g = d["diag"][m]
        for _, r in g[g.codigo_ine_ccaa == cc].iterrows():
            a(f"| {m} | {r.spec} | {num(r.BG_p, 3)} | {num(r.BP_p, 3)} | {num(r.JB_p, 3)} | {num(r.RESET_p, 3)} | "
              f"{num(r.CUSUM_p, 3)} | {num(r.VIF_max)} | {num(r.r2_adj)} | {num(r.phi_mitad1, 3)} | "
              f"{num(r.phi_mitad2, 3)} | {r.cambio_signo_phi} |")
    a("\nUn p bajo en Breusch-Godfrey indica autocorrelación residual; en ese caso los errores estándar de la "
      "tabla de corto plazo son menos fiables.\n")
    a("## 4. Descomposición por periodo (contabilidad descriptiva, con residuo)\n")
    a("Coeficientes no confirmados: la magnitud se da como rango entre estimadores (MG, PMG, CS-DL) y es "
      "asociación. Variación en log. «no atribuible = True»: el residuo domina o el rango abarca la variación.\n")
    a("**Ojo:** la descomposición se hace sobre datos con demeaning transversal, así que la variación es **relativa a la media de las 17 CCAA** "
      "(log; positiva si la región subió más que la media). Las subidas absolutas están en la sección 1.\n")
    a("| mercado | periodo | variación relativa a la media de las CCAA (log) | residuo no explicado [rango] | ¿no atribuible? |\n|---|---|---|---|---|")
    for m in MERC:   # v7 (orquestador): solo observado y residuo, para fichas de 2-4 páginas; el detalle en region_<cod>.md
        g = d["des"][m]
        g = g[(g.codigo_ine_ccaa == cc) & (g.spec == "A") & (g.componente == "residuo no explicado")]
        for _, r in g.iterrows():
            a(f"| {m} | {r.periodo} | {num(r.dy_obs, 4)} | [{num(r.rango_min, 4)}; {num(r.rango_max, 4)}] | "
              f"{'sí' if r.no_atribuible else 'no'} |")
    a("")
    a("Las tablas completas (incluida la descomposición del ECM propio de la región) están en "
      f"output/v7/R4/<mercado>/region_{cc}.md.\n")
    a("## 5. Backtest (R5)\n")
    a("| mercado | mejor modelo | RMSE ×100 | base fuerte | ¿bate a la base fuerte? | base más exigente | ¿bate a la más exigente? |\n"
      "|---|---|---|---|---|---|---|")
    for m in MERC:
        r = d["best"][(d["best"].codigo_ine_ccaa == cc) & (d["best"].mercado == m)].iloc[0]
        a(f"| {m} | {r.mejor_modelo} | {num(r.rmse_mejor_x100, 3)} | {num(r.rmse_base_fuerte_x100, 3)} | "
          f"{'sí' if r.bate_base_fuerte else 'no'} | {num(r.rmse_base_exigente_x100, 3)} | "
          f"{'sí' if r.bate_base_exigente else 'no'} |")
    a(f"\nOrigen móvil 2016Q1-2026Q1, h = 1-4 agrupados. {ROT}; la selección del mejor modelo es en muestra "
      "(optimista). Solo \"predice mejor\" quien bate a la base más exigente.\n")
    a("## 6. Propuestas (R6; condicional, C4)\n")
    a("«Si se hiciera X, con estos supuestos…». Mediana e intervalo al 90 % (p5-p95) de la variación %, "
      "a 10 años (2 y 5 años en output/v7/R6/impactos.csv), con la elasticidad de oferta de la rejilla (v5). Esfuerzo = variación % del esfuerzo "
      "neto del beneficiario. Robusto = esfuerzo baja en todo el Monte Carlo y el precio de mercado no sube "
      "en ninguna extracción.\n")
    im = d["imp"][(d["imp"].cc == cc) & (d["imp"].fuente_es == "rejilla")]
    rb = d["rob"][d["rob"].cc == cc]
    a("| instrumento | mercado | h | precio o alquiler de mercado: mediana [p5, p95] | esfuerzo: mediana [p5, p95] | robusto |\n"
      "|---|---|---|---|---|---|")
    for _, r in im[im.horizonte == 10].sort_values(["instrumento", "mercado"]).iterrows():   # 2 y 5 años en impactos.csv
        a(f"| {INSTR.get(r.instrumento, r.instrumento)} | {r.mercado} | {r.horizonte} | {pct(r.dprecio_med)} [{pct(r.dprecio_p5)}, {pct(r.dprecio_p95)}] | "
          f"{pct(r.esfuerzo_med)} [{pct(r.esfuerzo_p5)}, {pct(r.esfuerzo_p95)}] | {'sí' if r.robusto else 'no'} |")
    ro = rb[rb.robusto]
    a("\n**Combinaciones robustas en esta región:** "
      + ("; ".join(f"{INSTR.get(r.instrumento, r.instrumento)} ({r.mercado})" for _, r in ro.iterrows()) or "ninguna")
      + " (con la rejilla de elasticidades; la elasticidad regional no fue utilizable).\n")
    a("Con la elasticidad regional (DOLS) el resultado puede diferir; esa elasticidad no pasó los diagnósticos en "
      "ninguna de las 17 CCAA, de modo que se mantiene la rejilla y se reportan ambas en impactos.csv.\n")
    if cc == "10":
        provincial(d, a)
    a("## 7. Qué no se puede afirmar\n")
    for t in ["Que haya una relación de largo plazo estable entre precios o alquileres y empleo, población, "
              "hipotecas o paro: ninguna hipótesis confirmatoria de R4 se confirmó (p de Holm ≥ 0,05).",
              "Que la velocidad de ajuste de la región sea distinta de cero o de la del resto: no hay inferencia válida individual.",
              "Que la descomposición atribuya la variación a un factor: el residuo domina y los coeficientes no están confirmados.",
              "Que el mejor modelo del backtest vaya a predecir mejor en el futuro: muestra ya vista en v1 y selección en muestra.",
              "Que una propuesta produzca los impactos de la tabla: son escenarios condicionales con supuestos (C4), no estimaciones identificadas.",
              "Nada de lo anterior es causal: la evidencia queda por debajo del nivel CAUSAL."]:
        a(f"- {t}")
    a(f"\n---\n*Capas: variaciones de índices C4 (fuente única regional); ECM, descomposición y backtest C4; "
      f"propuestas C4 condicional. {ROT}. Validación diferida pendiente: los modelos están congelados "
      "(R5) y se evaluarán con `make evaluate_v7` sobre datos posteriores.*")
    return "\n".join(L) + "\n"


INSTR = {"a_construccion": "construcción adicional", "b_vacias": "movilización de vacías", "c_vivienda_publica": "vivienda pública o asequible",
         "d_turisticos": "regulación local de turísticos", "e_ayuda_demanda": "ayudas a la demanda",
         "f1_rebaja_ITP_IVA_vivienda_nueva": "rebaja de ITP/IVA (vivienda nueva)", "f2_impuestos_construccion": "rebaja de impuestos y costes de construcción"}


def _zip_determinista(p):
    """Reescribe el .xlsx (zip) con fecha fija en cada entrada y orden estable: md5 idéntico entre ejecuciones."""
    import zipfile
    with zipfile.ZipFile(p) as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    import re as _re   # openpyxl fija «modified» a la hora de guardado: se normaliza en docProps/core.xml
    items = [(n, _re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(<)", rb"\g<1>2026-10-12T00:00:00Z\g<2>", d) if n == "docProps/core.xml" else d)
             for n, d in items]
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre, datos in items:
            zi = zipfile.ZipInfo(nombre, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, datos)


def provincial(d, a):
    a("### Apartado provincial: provincias de Alicante (03), Castellón (12) y Valencia (46), y València ciudad\n")
    a("Coeficientes de R4b (provincia_valencia.csv; 21 observaciones, "
      "especificación core; asociación, C4):\n")
    a("| provincia | dependiente | Δln población | Δln ocupados EPA | Δln importe hipotecario | n |\n|---|---|---|---|---|---|")
    for _, r in d["pv"].iterrows():
        a(f"| {r.cod_prov} | {r.dependiente} | {num(r.d_ln_pob_total, 3)} | {num(r.d_ln_ocupados_epa, 3)} | "
          f"{num(r.d_ln_hip_importe, 3)} | {r.n_obs} |")
    a("\nN pequeño y sin inferencia validada: se leen como descriptivos. Variación del valor tasado "
      "(€/m², Ministerio):\n")
    a("| territorio | periodo | inicio | fin | variación |\n|---|---|---|---|---|")
    pq = d["prov"]
    ult = pq.dropna(subset=["valor_tasado"]).trimestre.max()
    for cp, nom in [("03", "Alicante"), ("12", "Castellón"), ("46", "Valencia (provincia)")]:
        s = pq[pq.cod_prov == cp].set_index("trimestre").valor_tasado
        for per, q0, q1 in PERIODOS:
            q1 = q1 or ult
            v0, v1 = s.get(q0, float("nan")), s.get(q1, float("nan"))
            a(f"| {nom} ({cp}) | {per} | {q0} {num(v0, 0)} | {q1} {num(v1, 0)} | {pct((v1 / v0 - 1) * 100)} |")
    v = d["vlc"]
    s = v[(v.territorio == "València (municipio)") & (v.variable == "p_tasado")].set_index("periodo").valor
    ult = s.index.max()
    for per, q0, q1 in PERIODOS:
        q1 = q1 or ult
        v0, v1 = s.get(q0, float("nan")), s.get(q1, float("nan"))
        a(f"| València ciudad (municipio) | {per} | {q0} {num(v0, 0)} | {q1} {num(v1, 0)} | {pct((v1 / v0 - 1) * 100)} |")
    a("\nDatos de València ciudad (data/processed/valencia.csv): valor tasado trimestral, transacciones, "
      "viviendas turísticas desde 2020T3 y 19 distritos con SERPAVI anual. No hay IPV ni IPC de alquiler "
      "municipal en el INE a esta escala; no se estima un ECM municipal (N insuficiente).\n")


def xlsx(d, nombres):
    wb = Workbook()
    res = wb.active
    res.title = "Resumen"
    az, vd, ng = Font(color=AZUL), Font(color=VERDE), Font(color=NEGRO)
    res.append(["CCAA", "Variación IPV 2022T3-último", "Variación IPC alquiler 2022T3-último",
                "Mejor modelo compra (RMSE×100)", "Mejor modelo alquiler (RMSE×100)"])
    hojas = {}
    for cc, nom in nombres:
        ws = wb.create_sheet(f"{cc}_{slug(nom)}"[:31])
        hojas[cc] = ws.title
        filas, ult = variaciones(d["panel"], cc)
        ws["A1"] = f"{nom} ({cc}): {ROT}; C4; asociación"
        ws.append([])
        ws.append(["Serie", "Periodo", "Inicio", "Valor inicio", "Fin", "Valor fin", "Variación", "Fuente"])
        for i, (s, per, q0, v0, q1, v1) in enumerate(filas, start=4):
            ws.append([s, per, q0, v0, q1, v1, f"=F{i}/D{i}-1",
                       f"panel_ccaa_q_v7.csv, INE, hasta {ult}"])
            for col in "CDEF":
                ws[f"{col}{i}"].font = az
            ws[f"D{i}"].comment = Comment("Fuente: data/processed/v7/panel_ccaa_q_v7.csv (INE)", "R7")
            ws[f"G{i}"].font = ng
            ws[f"G{i}"].number_format = "0.0%"
        ws.append([])
        ws.append(["Mercado", "Mejor modelo", "RMSE mejor ×100", "RMSE base más exigente ×100", "Diferencia", "¿Bate a la base más exigente?", "Fuente"])
        r0 = ws.max_row + 1
        for k, m in enumerate(MERC):
            r = d["best"][(d["best"].codigo_ine_ccaa == cc) & (d["best"].mercado == m)].iloc[0]
            i = r0 + k
            ws.append([m, r.mejor_modelo, r.rmse_mejor_x100, r.rmse_base_exigente_x100, f"=D{i}-C{i}",
                       f'=IF(E{i}>0,"sí","no")', "output/v7/R5/mejor_por_region.csv; " + ROT])
            for col in "BCD":
                ws[f"{col}{i}"].font = az
        ws.append([])
        ws.append(["Mercado", "Spec", "φ (sin inferencia válida)", "Semivida (trim.)", "Fuente"])
        for m in MERC:
            v = d["vel"][m]
            for _, r in v[v.codigo_ine_ccaa == cc].iterrows():
                i = ws.max_row + 1
                ws.append([m, r.spec, r.phi, f'=IF(AND(C{i}<0,C{i}>-1),LN(0.5)/LN(1+C{i}),"NA")',
                           f"output/v7/R4/{m}/velocidad_ajuste_region.csv"])
                ws[f"C{i}"].font = az
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and c.font.color is None:
                    c.font = ng
        ws.freeze_panes = None
    for i, (cc, nom) in enumerate(nombres, start=2):
        h = hojas[cc]
        res.append([nom, f"='{h}'!G6", f"='{h}'!G9", f"='{h}'!B13", f"='{h}'!B12"])
        for col in "BCDE":
            res[f"{col}{i}"].font = vd
        res[f"B{i}"].number_format = res[f"C{i}"].number_format = "0.0%"
    n = len(nombres) + 2
    res[f"A{n}"] = "Media simple de las 17 CCAA"
    for col in "BC":
        res[f"{col}{n}"] = f"=AVERAGE({col}2:{col}{n - 1})"
        res[f"{col}{n}"].font = ng
        res[f"{col}{n}"].number_format = "0.0%"
    res[f"A{n + 2}"] = ("Enlaces en verde a cada hoja; la media es un cálculo de esta hoja. Fuente de las series: "
                        "panel_ccaa_q_v7.csv (INE). " + ROT)
    for ws in wb.worksheets:
        ws.freeze_panes = None
    import datetime as _dt
    fijo = _dt.datetime(2026, 10, 12)
    wb.properties.created = wb.properties.modified = fijo   # v7: metadatos fijos (determinismo)
    wb.save(O / "tablas_regionales.xlsx")
    _zip_determinista(O / "tablas_regionales.xlsx")


def resumen(d, nombres):
    L = []
    a = L.append
    a("# Atlas regional: comparación entre regiones\n")
    a(f"*Capa C4; {ROT}; asociación, sin lectura causal. Datos hasta 2026Q2 (panel v7).*\n")
    a("## 1. Largo plazo: ninguna hipótesis confirmada\n")
    a("| id | mercado | variable | MG | PMG | CS-DL | rango entre estimadores | veredicto |\n|---|---|---|---|---|---|---|---|")
    for _, r in d["hip"].iterrows():
        a(f"| {r['id']} | {r.mercado} | {r.variable} | {num(r.MG_dm, 3)} | {num(r.PMG_dm, 3)} | {num(r.CS_DL, 3)} | "
          f"[{num(r.rango_min, 3)}; {num(r.rango_max, 3)}] | {r.veredicto} |")
    a("\nLas 7 hipótesis quedan sin confirmar (p de Holm ≥ 0,05) y en varias los estimadores discrepan en el signo; "
      "se dan ambos. No se identifica un largo plazo estable ni común (V15).\n")
    a("## 2. Velocidades de ajuste (sin inferencia válida)\n")
    a("| CCAA | φ alquiler A | semivida | φ compra A | semivida |\n|---|---|---|---|---|")
    for cc, nom in nombres:
        f = []
        for m in MERC:
            v = d["vel"][m]
            r = v[(v.codigo_ine_ccaa == cc) & (v.spec == "A")].iloc[0]
            f += [num(r.phi, 3), num(r.semivida_trim, 1)]
        a(f"| {nom} | " + " | ".join(f) + " |")
    a("\nSemivida en trimestres; NA si φ ≥ 0. Solo la comparación del promedio con el valor crítico simulado es válida, y no es significativa.\n")
    a("## 3. Corto plazo heterogéneo y residuo\n")
    for m in MERC:
        c = d["cp"][m]
        c = c[c.coef != "phi"]
        a(f"- {m}: {int(c.rechaza_bh10.sum())} de {len(c)} coeficientes de corto plazo región × spec × variable "
          "con q BH < 0,10 (el resto sin señal).")
    a("\nRango del residuo no explicado entre estimadores, 2022T3-último, especificación A:\n")
    a("Variación **relativa a la media de las 17 CCAA** (datos con demeaning transversal; log). No es la subida absoluta, que está en cada ficha (sección 1).\n")
    a("| CCAA | mercado | variación relativa a la media | residuo mín. | residuo máx. |\n|---|---|---|---|---|")
    for cc, nom in nombres:
        for m in MERC:
            g = d["des"][m]
            g = g[(g.codigo_ine_ccaa == cc) & (g.spec == "A") & (g.periodo == "2022T3-ultimo")
                  & (g.componente == "residuo no explicado")]
            if len(g):
                r = g.iloc[0]
                a(f"| {nom} | {m} | {num(r.dy_obs, 3)} | {num(r.rango_min, 3)} | {num(r.rango_max, 3)} |")
    a("\n## 4. Backtest (R5; selección en muestra, optimista)\n")
    b = d["best"]
    for m in MERC:
        g = b[b.mercado == m]
        a(f"- {m}: el mejor modelo bate a la base fuerte en {int(g.bate_base_fuerte.sum())} de 17 CCAA y a la base más "
          f"exigente en {int(g.bate_base_exigente.sum())} de 17 ({', '.join(nombre for cc, nombre in nombres if cc in set(g[g.bate_base_exigente].codigo_ine_ccaa)) or 'ninguna'}). "
          "PMG no la bate en ningún mercado.")
    a("\n## 5. Recomendaciones robustas por región (condicional, C4)\n")
    a("Robusto: esfuerzo a la baja en todo el Monte Carlo y precio de mercado sin subida en ninguna extracción (V18). "
      "Se evalúan instrumentos, no actores. «Rejilla» indica que la elasticidad regional no fue utilizable.\n")
    a("| CCAA | alquiler | compra |\n|---|---|---|")
    for cc, nom in nombres:
        r = d["rob"][(d["rob"].cc == cc) & d["rob"].robusto]
        f = [", ".join(sorted(INSTR.get(i, i) for i in r[r.mercado == m].instrumento)) or "ninguno" for m in ("alquiler", "precio")]
        a(f"| {nom} | {f[0]} | {f[1]} |")
    a("\n## English summary (about 150 words)\n")
    a("This atlas compares 17 Spanish regions using quarterly panel data to 2026Q2. It is informative only: "
      "the sample was already seen in v1, and every result sits in layer C4 (exploratory, association, no causal "
      "reading). None of the seven long-run hypotheses on prices and rents was confirmed (Holm-adjusted p above "
      "0.05), and the three estimators often disagree, so ranges are reported. Speeds of adjustment differ "
      "across regions but carry no valid individual inference. Short-run movements are heterogeneous and the "
      "unexplained residual dominates the historical accounting. In the backtest, the best model beats the "
      "most demanding baseline in 6 of 17 regions for rents and 5 of 17 for purchases; selection is in-sample, "
      "so this is optimistic. Policy scenarios are conditional (if X were done, under stated assumptions) and "
      "evaluate instruments, not actors. Deferred validation on new data remains pending.\n")
    (O / "atlas_resumen.md").write_text("\n".join(L), encoding="utf-8")


REFS = [
    "Pesaran, M. H., Shin, Y. y Smith, R. P. (1999). Pooled mean group estimation of dynamic heterogeneous panels. JASA. NO VERIFICADA (cuartil no verificado).",
    "Pesaran, M. H. (2006). Estimation and inference in large heterogeneous panels with a multifactor error structure. Econometrica. NO VERIFICADA (cuartil no verificado).",
    "Pesaran, M. H. (2007). A simple panel unit root test in the presence of cross-section dependence. J. Applied Econometrics. NO VERIFICADA (cuartil no verificado).",
    "Chudik, A., Mohaddes, K., Pesaran, M. H. y Raissi, M. (2016). Long-run effects in large heterogeneous panel data models with cross-sectionally correlated errors. Advances in Econometrics. NO VERIFICADA (cuartil no verificado).",
    "Webb, M. D. (2014). Reworking wild bootstrap based inference for clustered errors. Queen's Economics Department WP 1315. NO VERIFICADA.",
    "Driscoll, J. C. y Kraay, A. C. (1998). Consistent covariance matrix estimation with spatially dependent panel data. Review of Economics and Statistics. NO VERIFICADA (cuartil no verificado).",
    "Holm, S. (1979). A simple sequentially rejective multiple test procedure. Scandinavian J. Statistics. NO VERIFICADA (cuartil no verificado).",
    "Benjamini, Y. y Hochberg, Y. (1995). Controlling the false discovery rate. JRSS B. NO VERIFICADA (cuartil no verificado).",
    "Diebold, F. X. y Mariano, R. S. (1995). Comparing predictive accuracy. J. Business & Economic Statistics. NO VERIFICADA (cuartil no verificado).",
    "Bai, J. y Perron, P. (2003). Computation and analysis of multiple structural change models. J. Applied Econometrics. NO VERIFICADA (cuartil no verificado).",
    "Breusch-Godfrey (autocorrelación), Breusch-Pagan (heterocedasticidad), Jarque-Bera (normalidad), RESET de Ramsey (forma funcional), CUSUM de Brown-Durbin-Evans (estabilidad), VIF (colinealidad): contrastes estándar. NO VERIFICADA (referencias no consultadas en Crossref en este encargo).",
]


def anexo():
    L = ["# Anexo de ecuaciones y contrastes (V7)\n",
         f"*Capa C4; {ROT}; asociación, sin lectura causal. Generado desde output/v7/R4, R4b, R5 y R6.*\n",
         "## A. Contrastes y su uso\n",
         "| contraste | uso | familia de control |\n|---|---|---|",
         "| Raíces unitarias de panel (CIPS) y por región | orden de integración | informativo; I(2) no se descarta en 8 de 17 CCAA |",
         "| Cointegración de panel y por región | existencia de largo plazo | Holm por familia |",
         "| t de MG, PMG, DFE y CS-DL con bootstrap salvaje (Webb) y Driscoll-Kraay | hipótesis H7 | Holm |",
         "| Breusch-Godfrey, Breusch-Pagan, Jarque-Bera, RESET, CUSUM, VIF | diagnósticos por región | descriptivo |",
         "| Bai-Perron | cambios estructurales | descriptivo |",
         "| Diebold-Mariano y permutación | backtest R5 | Holm |",
         "| Coeficientes de corto plazo por región | exploratorio | BH, q = 0,10 |\n"]
    for m in MERC:
        t = (O / f"R4/{m}/ecuaciones.md").read_text(encoding="utf-8")
        L.append(f"## B. Ecuaciones R4, {m}\n\n" + dem(re.sub(r"^# .*\n", "", t, count=1)).replace("#### ", "### ", 0))
    t = (O / "R6/metodo.md").read_text(encoding="utf-8")
    L.append("## C. Ecuaciones R6 (traslado e incidencia, condicional)\n\n" + re.sub(r"^# .*\n", "", t, count=1))
    L.append("## D. Referencias\n")
    L += [f"- {r}" for r in REFS]
    (O / "anexo_ecuaciones.md").write_text("\n".join(L) + "\n", encoding="utf-8")


def guia():
    t = f"""# Guía de lectura para no especialistas

*Capa C4; {ROT}. Todo lo que sigue es asociación y escenarios condicionales, no causalidad.*

## Qué es un ECM
Un modelo de corrección del error (ECM) separa dos movimientos de un precio o un alquiler. El primero es el
**corto plazo**: cuánto cambia el índice cuando cambian este trimestre el empleo, la población o las hipotecas.
El segundo es la **corrección**: si el precio se ha alejado de lo que «le correspondería» según esas variables,
qué fracción de esa distancia se recupera cada trimestre. Esa fracción es la velocidad de ajuste (φ). La
**semivida** es el tiempo que tarda la distancia en reducirse a la mitad.

Para que la corrección tenga sentido hace falta que exista una relación de **largo plazo** estable entre el
precio y esas variables. Con los datos regionales de 2007 a 2026, no se encontró esa relación estable ni común a
las regiones.

## Qué significa «no confirmada»
Antes de mirar los datos se fijaron siete hipótesis de largo plazo (por ejemplo, «más población se asocia a
precios más altos»). «No confirmada» significa que, tras corregir por hacer varios contrastes a la vez (Holm), la
evidencia no alcanza el umbral del 5 %. No significa que la hipótesis sea falsa: significa que estos datos no
permiten sostenerla. Cuando los tres estimadores usados dan signos distintos, se muestran todos y el rango.

Por eso las velocidades de ajuste de cada región se publican con el rótulo **«sin inferencia válida»**: son
estimaciones de lo que ocurrió en la muestra, sin intervalo de confianza que permita compararlas.

## Por qué las propuestas van en condicional
Las tablas de propuestas responden a «si se hiciera X, con estos supuestos, el precio o el esfuerzo variarían en
este rango». Los supuestos (elasticidades de oferta y demanda, reparto entre alquiler y compra) no están
identificados con los datos; por eso la capa es C4 y el intervalo al 90 % refleja la incertidumbre de esos
supuestos, no un error de medida. Una propuesta se marca robusta solo si el esfuerzo baja en todas las
simulaciones y el precio de mercado no sube en ninguna. Se evalúan instrumentos, no actores.

## Cómo leer una ficha
1. **Qué subió y cuánto**: variación de índices del INE entre fechas concretas.
2. **Ecuación y tabla**: los coeficientes de corto plazo tienen intervalo; φ y θ no.
3. **Diagnósticos**: avisan de cuándo una tabla es menos fiable.
4. **Descomposición**: reparte la variación entre componentes y un **residuo**; si el residuo domina, la
   variación queda sin explicar.
5. **Backtest**: un modelo «predice mejor» solo si supera a la línea base más exigente. La muestra ya se usó en
   v1 y el mejor modelo se eligió dentro de ella, así que el resultado es optimista.
6. **Qué no se puede afirmar**: los límites de lectura de esa región.

## Pendiente
La validación fuera de muestra con datos posteriores sigue pendiente; los modelos están congelados para
evaluarla (`make evaluate_v7`).
"""
    (O / "guia_lectura.md").write_text(t, encoding="utf-8")


def main():
    d = cargar()
    nombres = sorted(d["imp"][["cc", "ccaa"]].drop_duplicates().itertuples(index=False, name=None))
    AT.mkdir(parents=True, exist_ok=True)
    for f in AT.glob("*.md"):
        f.unlink()
    for cc, nom in nombres:
        (AT / f"{cc}_{slug(nom)}.md").write_text(ficha(d, cc, nom), encoding="utf-8")
    resumen(d, nombres)
    xlsx(d, nombres)
    anexo()
    guia()
    print("ok", len(nombres), "fichas")


if __name__ == "__main__":
    main()
