"""V7.1-S1 · Atlas, Excel y verificador actualizados con dosis y coste (sin análisis nuevo). Lee output/v7 y output/v7_1/S1.

Escribe:
  output/v7_1/atlas/<cc>_<nombre>.md        fichas v7 + sección «6 bis. Dosis y coste (v7.1)»
  output/v7_1/atlas/atlas_resumen.md         resumen v7 + tabla nacional de dosis y coste
  output/v7_1/atlas/tablas_regionales.xlsx   hojas v7 + 3 hojas de dosis y coste (azul = dato de entrada con fuente,
                                             negro = cálculo de la hoja, verde = enlace a otra hoja; sin paneles inmovilizados)
  output/v7_1/fichas/fichas_verificador.json y output/v7_1/verificador_v7_1.md (V7-V2 con dosis y coste)
"""
from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.worksheet.formula import ArrayFormula

RAIZ = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ / "src" / "v7"))
from r7_atlas import _zip_determinista  # noqa: E402

V7 = RAIZ / "output" / "v7"
S1 = RAIZ / "output" / "v7_1" / "S1"
O = RAIZ / "output" / "v7_1" / "atlas"
AZUL, VERDE, NEGRO = "FF0000FF", "FF008000", "FF000000"
ORDEN = ["b_vacias", "c_vivienda_publica", "d_turisticos", "a_construccion", "f2_impuestos_construccion", "e_ayuda_demanda",
         "f1_rebaja_ITP_IVA_vivienda_nueva"]
MERC = {"precio": "compra", "alquiler": "alquiler"}


def es(x, d=1, signo=False):
    if pd.isna(x):
        return "—"
    s = f"{x:+,.{d}f}" if signo else f"{x:,.{d}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def seccion(t: pd.DataFrame, cc: str) -> str:
    g = t[(t.cc == cc) & (t.horizonte == 10)]
    L = ["## 6 bis. Dosis y coste de las propuestas (v7.1; condicional, C4)", "",
         "Misma simulación que la sección 6 (output/v7_1/S1/dosis_coste.csv, que reproduce R6). Dosis acumulada a 10 años (en e, f1 y f2, por año); "
         "coste fiscal anual aproximado, sin efectos inducidos; mediana [p5, p95]. El signo lo fija en gran parte la fórmula de incidencia: lo "
         "informativo son las magnitudes relativas.", "",
         "| instrumento | mercado | dosis | precio de mercado | coste anual (M€) | por 10.000 viviendas | por 100 M€ anuales (mercado) | coste por punto (M€/año) |",
         "|---|---|---|---|---|---|---|---|"]
    for inst in ORDEN:
        for mer in ("precio", "alquiler"):
            r = g[(g.instrumento == inst) & (g.mercado == mer)]
            if r.empty:
                continue
            r = r.iloc[0]
            if pd.notna(r.coste_anual_meur_med):
                cs = f"{es(r.coste_anual_meur_med, 0)} [{es(r.coste_anual_meur_p5, 0)}, {es(r.coste_anual_meur_p95, 0)}]"
            elif r.estado_coste.startswith("coste anual"):
                cs = "el de compra"
            else:
                cs = "no estimable"
            e10 = "—" if pd.isna(r.efecto_10k_med) else f"{es(r.efecto_10k_med, 2, True)} pp [{es(r.efecto_10k_p5, 2, True)}, {es(r.efecto_10k_p95, 2, True)}]"
            e100 = "—" if pd.isna(r.efecto_100meur_mercado_med) else f"{es(r.efecto_100meur_mercado_med, 2, True)} pp [{es(r.efecto_100meur_mercado_p5, 2, True)}, {es(r.efecto_100meur_mercado_p95, 2, True)}]"
            cpp = "—" if pd.isna(r.coste_por_punto_meur_med) else f"{es(r.coste_por_punto_meur_med, 0)} [{es(r.coste_por_punto_meur_p5, 0)}, {es(r.coste_por_punto_meur_p95, 0)}]"
            L.append(f"| {r.nombre} | {MERC[mer]} | {es(r.dosis_med, 0)} | {es(r.dprecio_med, 1, True)} % | {cs} | {e10} | {e100} | {cpp} |")
    L += ["", "«No estimable»: no hay coste unitario verificado ni versionado (motivos en output/v7_1/S1/resumen.md). En las ayudas a la demanda "
          "y en la rebaja de ITP/IVA el coste por punto se refiere al precio neto que paga el beneficiario; el precio de mercado sube. El efecto por "
          "100 M€ de la rebaja de impuestos a la construcción está sobrestimado por construcción (la rebaja se traslada a todo el mercado y el coste "
          "solo cubre la obra nueva).", ""]
    return "\n".join(L)


def fichas_atlas(t: pd.DataFrame) -> None:
    O.mkdir(parents=True, exist_ok=True)
    for p in sorted((V7 / "atlas").glob("[0-9][0-9]_*.md")):
        cc = p.name[:2]
        txt = p.read_text(encoding="utf-8")
        marca = "## 7. Qué no se puede afirmar"
        assert marca in txt, p
        txt = txt.replace(marca, seccion(t, cc) + "\n" + marca, 1)
        txt = txt.replace("- Que una propuesta produzca los impactos de la tabla:",
                          "- Que un instrumento sea «más eficiente por euro» que otro: los costes de los instrumentos de cantidad no son estimables y los "
                          "demás dependen de supuestos (C4).\n- Que una propuesta produzca los impactos de la tabla:", 1)
        (O / p.name).write_text(txt, encoding="utf-8")
    r = (V7 / "atlas_resumen.md").read_text(encoding="utf-8")
    s1 = (S1 / "resumen.md").read_text(encoding="utf-8").split("## Resumen nacional", 1)[1].split("## Dosis por instrumento", 1)[0]
    lect = (S1 / "resumen.md").read_text(encoding="utf-8").split("## Lectura", 1)[1]
    r += ("\n\n## Dosis y coste de las propuestas (v7.1; condicional, C4)\n\nMediana entre CCAA y rango entre CCAA" + s1.split("\n", 1)[1]
          + "Definiciones y motivos de «no estimable»: output/v7_1/S1/resumen.md. Ficha por CCAA: sección 6 bis.\n\n### Lectura" + lect)
    (O / "atlas_resumen.md").write_text(r, encoding="utf-8")
    for f in ("anexo_ecuaciones.md", "guia_lectura.md"):
        shutil.copyfile(V7 / f, O / f)


def excel(t: pd.DataFrame) -> None:
    wb = load_workbook(V7 / "tablas_regionales.xlsx")
    az, vd, ng = Font(color=AZUL), Font(color=VERDE), Font(color=NEGRO)
    fuente = "Fuente: output/v7_1/S1/dosis_coste.csv (reproducción de output/v7/R6; Monte Carlo 4.000 extracciones; C4)"
    d = wb.create_sheet("DC_datos")
    cols = ["cc", "ccaa", "nombre", "mercado", "dosis_med", "dprecio_med", "dneto_med", "coste_anual_meur_med", "coste_anual_meur_p5",
            "coste_anual_meur_p95", "efecto_10k_med", "efecto_100meur_mercado_med", "coste_por_punto_meur_med", "estado_coste"]
    d.append(["Dosis y coste a 10 años por CCAA (v7.1, C4). Datos de entrada en azul; fuente en la nota de B2 y en la fila final."])
    d.append(["CCAA (código INE)", "CCAA", "Instrumento", "Mercado", "Dosis (viviendas)", "Δ precio de mercado (%)", "Δ precio neto (%)",
              "Coste anual (M€) mediana", "p5", "p95", "Efecto por 10.000 viv. (pp)", "Efecto por 100 M€/año (pp)", "Coste por punto (M€/año)",
              "Estado del coste"])
    d["B2"].comment = Comment(fuente, "V7.1")
    g = t[t.horizonte == 10].copy()
    g["mercado"] = g.mercado.map(MERC)
    g = g.sort_values(["cc", "instrumento", "mercado"])
    for row in g[cols].itertuples(index=False):
        d.append([None if (isinstance(v, float) and pd.isna(v)) else (round(v, 4) if isinstance(v, float) else v) for v in row])
        i = d.max_row
        for c in "ABCDEFGHIJKLMN":
            d[f"{c}{i}"].font = az
    n_ult = d.max_row
    d.append([])
    d.append([fuente])
    c = wb.create_sheet("DC_calculos")
    c.append(["Cálculos de esta hoja (negro) a partir de DC_datos (enlaces en verde)."])
    c.append(["CCAA", "Instrumento", "Mercado", "Coste anual (M€)", "Coste a 10 años (M€)", "Δ precio de mercado (%)",
              "Cociente de medianas: pp por 100 M€ anuales", "Nota"])
    for i in range(3, n_ult + 1):
        k = c.max_row + 1
        c.append([f"=DC_datos!B{i}", f"=DC_datos!C{i}", f"=DC_datos!D{i}", f"=DC_datos!H{i}",
                  f'=IF(ISNUMBER(D{k}),10*D{k},"no estimable")', f"=DC_datos!F{i}",
                  f'=IF(ISNUMBER(D{k}),F{k}*100/D{k},"—")', "cociente de medianas (difiere algo de la mediana del cociente de DC_datos)"])
        for col in "ABCDF":
            c[f"{col}{k}"].font = vd
        for col in "EG":
            c[f"{col}{k}"].font = ng
    r = wb.create_sheet("DC_resumen")
    r.append(["Resumen nacional a 10 años: mediana entre CCAA (cálculo en negro sobre enlaces a DC_datos)."])
    r.append(["Instrumento", "Mercado", "Dosis mediana", "Δ precio de mercado mediano (%)", "Coste anual mediano (M€)", "Efecto por 10.000 viv. (pp)",
              "Efecto por 100 M€/año (pp)"])
    for inst in ORDEN:
        for mer in ("precio", "alquiler"):
            if g[(g.instrumento == inst) & (g.mercado == MERC[mer])].empty:
                continue
            nom = g[g.instrumento == inst].nombre.iloc[0]
            k = r.max_row + 1
            rango = lambda col: f"DC_datos!{col}3:{col}{n_ult}"  # noqa: E731
            r.append([nom, MERC[mer]])
            for col_r, col in zip("CDEFG", "EFHKL"):
                ref = f"{col_r}{k}"
                r[ref] = ArrayFormula(ref, f'=IFERROR(MEDIAN(IF(({rango("C")}="{nom}")*({rango("D")}="{MERC[mer]}")*ISNUMBER({rango(col)}),{rango(col)})),"—")')
                r[ref].font = ng
    r.append([])
    r.append(["Fórmulas matriciales (MEDIAN(IF(...))) sobre DC_datos; los valores publicables están en output/v7_1/S1/hechos.json → cifras_clave."])
    for ws in wb.worksheets:
        ws.freeze_panes = None
    import datetime as _dt
    wb.properties.created = wb.properties.modified = _dt.datetime(2026, 10, 12)
    wb.save(O / "tablas_regionales.xlsx")
    _zip_determinista(O / "tablas_regionales.xlsx")


def verificador(t: pd.DataFrame) -> None:
    f = json.loads((V7 / "fichas" / "fichas_verificador.json").read_text())
    g = t[(t.horizonte == 10) & (t.mercado == "precio")]
    e = g[g.instrumento == "e_ayuda_demanda"]
    for x in f:
        if x["id"] == "V7-V2":
            x["magnitud"] += (f" Dosis: todas las compraventas de 2025 (mediana {es(e.dosis_med.median(), 0)} por CCAA) con una ayuda del 5-15 % del precio; "
                              f"coste fiscal anual aproximado {es(e.coste_anual_meur_med.median(), 0)} M€ por CCAA (mediana; rango entre CCAA "
                              f"{es(e.coste_anual_meur_med.min(), 0)}-{es(e.coste_anual_meur_med.max(), 0)}); por cada 100 M€ anuales el precio de mercado "
                              f"sube {es(e.efecto_100meur_mercado_med.median(), 2)} puntos (mediana). El signo lo fija la fórmula de incidencia (C4).")
            x["evidencia"] = x["evidencia"] + ["output/v7_1/S1/dosis_coste.csv", "output/v7_1/S1/resumen.md"]
            x["version"] = "v7.1"
    out = RAIZ / "output" / "v7_1" / "fichas"
    out.mkdir(parents=True, exist_ok=True)
    (out / "fichas_verificador.json").write_text(json.dumps(f, ensure_ascii=False, indent=1), encoding="utf-8")
    ver = ["# Verificador V7.1 (fichas de V7, con dosis y coste en V7-V2)", "",
           "Fichas añadidas en V7, actualizadas en v7.1. Complementan el verificador de v6.2 (output/v5/verificador), cuyos recuentos no cambian.", ""]
    for x in f:
        ver += [f"## {x['id']} · {x['tema']}", "", f"**Afirmación evaluada:** {x['enunciado']}", "", f"**Veredicto:** {x['veredicto']} ({x['capa']})", "",
                f"**Magnitud:** {x['magnitud']}", "", f"**Intervalo:** {x['intervalo']}", "", f"**Límites:** {x['limites']}", "",
                f"**Regla:** {x['regla']}", "", "**Evidencia:** " + "; ".join(x["evidencia"]), ""]
    (RAIZ / "output" / "v7_1" / "verificador_v7_1.md").write_text("\n".join(ver), encoding="utf-8")


def main():
    t = pd.read_csv(S1 / "dosis_coste.csv", dtype={"cc": str})
    fichas_atlas(t)
    excel(t)
    verificador(t)
    print("s1_atlas: fichas, resumen, Excel y verificador v7.1 escritos")


if __name__ == "__main__":
    main()
