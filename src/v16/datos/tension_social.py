"""V16-F1 · Normalización de la familia tension_social al formato largo de data/base/hechos.parquet.

Determinista (SEED=20261010, sin aleatoriedad) y sin red. Lee solo data/raw/v16/tension_social/ y data/base/publica/territorios.parquet.
Salidas:
- data/processed/v16/tension_social.parquet: solo INE (redistribuible = sí): EMCR municipal 2021-2024, cifras oficiales del padrón
  1996-2025 y padrón por sección censal 2022 por continentes. NADA del CGPJ.
- data/processed/v16/tension_social_terceros.parquet: CGPJ (redistribuible = False, fuente dudosa; NO se versiona), solo si existen
  los libros en _terceros/.
- data/processed/v16/tension_social_partidos_municipios_terceros.parquet: correspondencia partido judicial -> municipios (código INE)
  y lista de nombres del CGPJ descartados (NO se versiona).
- data/processed/v16/tension_social_resultado.json y docs/v16/fuentes/tension_social_variables.csv.
capa = C4 (dato bruto). posterior_prereg = False (el corte de V16 no está fijado).
"""
from __future__ import annotations

import csv
import gzip
import io
import json
import re
import sys
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd

SEED = 20261010
RAIZ = Path(__file__).resolve().parents[3]
ENT = RAIZ / "data" / "raw" / "v16" / "tension_social"
TER = ENT / "_terceros"
SAL = RAIZ / "data" / "processed" / "v16"
DICC = RAIZ / "docs" / "v16" / "fuentes" / "tension_social_variables.csv"
COLS = ["territorio_codigo", "nivel", "periodo", "frecuencia", "variable", "desagregacion", "valor", "unidad", "fuente", "edicion",
        "fecha_dato", "capa", "licencia", "redistribuible", "metodo_dato", "origen", "posterior_prereg", "muestra_sellada_v6"]

# Tabla de correspondencia versionada: nombre de provincia del CGPJ (mayúsculas sin tildes) -> código INE de provincia.
# Ceuta y Melilla no tienen fila propia en las series provinciales (el CGPJ las cuenta en Cádiz y Málaga).
CGPJ_PROV = {
    "ALMERIA": "04", "CADIZ": "11", "CORDOBA": "14", "GRANADA": "18", "HUELVA": "21", "JAEN": "23", "MALAGA": "29", "SEVILLA": "41",
    "HUESCA": "22", "TERUEL": "44", "ZARAGOZA": "50", "ASTURIAS": "33", "ILLES BALEARS": "07", "LAS PALMAS": "35",
    "SANTA CRUZ DE TENERIFE": "38", "CANTABRIA": "39", "AVILA": "05", "BURGOS": "09", "LEON": "24", "PALENCIA": "34",
    "SALAMANCA": "37", "SEGOVIA": "40", "SORIA": "42", "VALLADOLID": "47", "ZAMORA": "49", "ALBACETE": "02", "CIUDAD REAL": "13",
    "CUENCA": "16", "GUADALAJARA": "19", "TOLEDO": "45", "BARCELONA": "08", "GIRONA": "17", "LLEIDA": "25", "TARRAGONA": "43",
    "ALICANTE": "03", "CASTELLON": "12", "VALENCIA": "46", "BADAJOZ": "06", "CACERES": "10", "A CORUNA": "15", "LUGO": "27",
    "OURENSE": "32", "PONTEVEDRA": "36", "MADRID": "28", "MURCIA": "30", "NAVARRA": "31", "ARABA/ALAVA": "01", "GIPUZKOA": "20",
    "BIZKAIA": "48", "LA RIOJA": "26",
}
# Partidos judiciales de Ceuta y Melilla (el CGPJ los incluye en los bloques de Cádiz y Málaga).
CGPJ_PJ_ESPECIAL = {"CEUTA": "51", "MELILLA": "52"}

CGPJ_SERIES_PROV = [  # hoja, variable, descripcion
    ("Lanzamientos pract. Total prov", "cgpj_lanzamientos_total_numero", "Lanzamientos practicados, total"),
    ("Lanzamientos E.hipotecaria prov", "cgpj_lanzamientos_hipotecaria_numero", "Lanzamientos practicados por ejecución hipotecaria"),
    ("Lanzamientos L.A.U. prov", "cgpj_lanzamientos_lau_numero", "Lanzamientos practicados por la Ley de Arrendamientos Urbanos"),
    ("Lanzamientos. Otros prov", "cgpj_lanzamientos_otros_numero", "Otros lanzamientos practicados"),
    (" Ej.Hipot por provincias", "cgpj_ejecuciones_hipotecarias_numero", "Ejecuciones hipotecarias presentadas"),
]
UNI_LANZ = "lanzamientos practicados (número)"
UNI_EJH = "ejecuciones hipotecarias presentadas (número)"
UNI_PERS = "personas"
UNI_PERS_1E = "personas a 1 de enero"


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s).upper())
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.replace("'", "").replace("-", " ").replace(",", " ")
    return re.sub(r"\s+", " ", s).strip()


def variantes(nombre: str) -> set[str]:
    """Variantes normalizadas de un nombre (separa lenguas con «/», recoloca «(A)» y «-LA» como artículo inicial)."""
    out = {norm(str(nombre).upper().replace("/", " "))}  # forma completa con las lenguas unidas
    for p in str(nombre).upper().split("/"):
        p = p.strip()
        m = re.match(r"^(.*)\s*\((A|O|EL|LA|LOS|LAS|L|ES|SA)\)$", p)
        if m:
            p = m.group(2) + " " + m.group(1).strip()
        m = re.match(r"^(.*)-(LA|EL|LOS|LAS)$", p)
        if m:
            p = m.group(2) + " " + m.group(1)
        out.add(norm(p))
    return out


def manifiesto() -> dict:
    m = {}
    with open(ENT / "_manifest.csv", encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f):
            m[r["fichero"]] = r
    return m


def construir(df: pd.DataFrame, rel: str, man: dict) -> pd.DataFrame:
    r = man[rel]
    df = df.copy()
    df["fuente"] = r["fuente"]
    df["edicion"] = r["edicion"] + " (descarga " + r["fecha_descarga"] + ")"
    df["licencia"] = r["licencia"]
    df["redistribuible"] = r["redistribuible"].strip().lower() == "sí"
    df["origen"] = rel
    df["capa"] = "C4"
    df["posterior_prereg"] = False
    if "metodo_dato" not in df:
        df["metodo_dato"] = "observado"
    if "desagregacion" not in df:
        df["desagregacion"] = ""
    return df


def sellada(df: pd.DataFrame) -> pd.Series:
    sys.path.insert(0, str(RAIZ / "src"))
    import holdout  # solo devuelve la lista de códigos; no abre nada sellado
    prov = {str(p).zfill(2) for p in holdout.provincias_selladas_v6()}
    cod = df.territorio_codigo
    pj = df.nivel.eq("partido_judicial") & cod.str[3:5].isin(prov)  # decisión conservadora: el partido judicial hereda la marca
    reg = df.nivel.isin(["provincia", "municipio", "distrito", "seccion"]) & cod.str[:2].isin(prov)
    return reg | pj


def territorios() -> pd.DataFrame:
    return pd.read_parquet(RAIZ / "data/base/publica/territorios.parquet")


DESCARTES: dict[str, list] = {}


def filtra_codigos(df: pd.DataFrame, nivel: str, t: pd.DataFrame, etiqueta: str) -> pd.DataFrame:
    """Conserva códigos presentes en territorios (vigentes o con vigente_2025 False); lista y descarta el resto."""
    tt = t[t.nivel == nivel]
    ok = set(tt.territorio_codigo)
    mal = sorted(set(df.loc[df.nivel == nivel, "territorio_codigo"]) - ok)
    no_vig = sorted(set(df.loc[df.nivel == nivel, "territorio_codigo"]) & set(tt[~tt.vigente_2025].territorio_codigo))
    DESCARTES[etiqueta] = {"descartados_no_en_territorios": mal, "conservados_con_vigente_2025_false": no_vig}
    return df[(df.nivel != nivel) | df.territorio_codigo.isin(ok)]


def num(s: pd.Series) -> pd.Series:
    return pd.to_numeric(s.str.replace(".", "", regex=False), errors="coerce")


# ---------------------------------------------------------------- INE
def leer_ine_csv(rel: str) -> pd.DataFrame:
    with gzip.open(RAIZ / rel, "rt", encoding="utf-8-sig") as f:
        return pd.read_csv(f, sep=";", dtype=str)


def territorio_ine(d: pd.DataFrame) -> pd.DataFrame:
    """Total Nacional / Provincias / Municipios -> territorio_codigo y nivel (solo código INE, antes del primer espacio)."""
    mun = d.Municipios.str[:5]
    prv = d.Provincias.str[:2]
    d["territorio_codigo"] = np.where(d.Municipios.notna(), mun, np.where(d.Provincias.notna(), prv, "00"))
    d["nivel"] = np.where(d.Municipios.notna(), "municipio", np.where(d.Provincias.notna(), "provincia", "nacional"))
    return d


def slug(s: pd.Series) -> pd.Series:
    return s.map(lambda x: re.sub(r"[^a-z0-9]+", "_", "".join(c for c in unicodedata.normalize("NFKD", str(x).lower())
                                                              if not unicodedata.combining(c))).strip("_"))


def emcr(man: dict, t: pd.DataFrame) -> pd.DataFrame:
    # (tabla, variable, columna de desagregación adicional, nombre de la desagregación)
    tablas = [
        (69696, "ine_emcr_inmigraciones_exterior_personas", None, None),
        (69711, "ine_emcr_emigraciones_exterior_personas", None, None),
        (69743, "ine_emcr_migraciones_interiores_llegadas_personas", "Nacionalidad", "nacionalidad"),
        (69745, "ine_emcr_migraciones_interiores_salidas_personas", "Nacionalidad", "nacionalidad"),
        (69744, "ine_emcr_migraciones_interiores_llegadas_personas", "País de nacimiento", "pais_nacimiento"),
        (69746, "ine_emcr_migraciones_interiores_salidas_personas", "País de nacimiento", "pais_nacimiento"),
        (69767, "ine_emcr_saldo_migratorio_personas", "Tipo de saldo", "tipo_saldo"),
    ]
    partes = []
    for tab, var, col, nom in tablas:
        rel = f"data/raw/v16/tension_social/emcr/ine_emcr_tabla_{tab}.csv.gz"
        d = territorio_ine(leer_ine_csv(rel))
        d["valor"] = num(d["Total"])
        d = d.dropna(subset=["valor"]).copy()
        d["periodo"] = d.Periodo
        d["frecuencia"] = "A"
        d["fecha_dato"] = d.Periodo + "-12-31"
        d["variable"] = var
        d["unidad"] = UNI_PERS + " (flujo anual)"
        des = "sexo=" + slug(d.Sexo)
        if col:
            des = des + f";{nom}=" + slug(d[col])
        d["desagregacion"] = des
        d["metodo_dato"] = "observado (tabla %d)" % tab
        partes.append(construir(d[["territorio_codigo", "nivel", "periodo", "frecuencia", "variable", "desagregacion", "valor", "unidad",
                                   "fecha_dato", "metodo_dato"]], rel, man))
    h = pd.concat(partes, ignore_index=True)
    h = filtra_codigos(h, "municipio", t, "emcr")
    return h


def padron_cifras(man: dict, t: pd.DataFrame) -> pd.DataFrame:
    rel = "data/raw/v16/tension_social/padron_cifras_oficiales/ine_cifras_oficiales_padron_municipio_29005.csv.gz"
    d = leer_ine_csv(rel)
    d["territorio_codigo"] = d.Municipios.str[:5]
    d["nivel"] = "municipio"
    d["valor"] = num(d["Total"])
    d = d.dropna(subset=["valor"]).copy()
    d["periodo"] = d.Periodo
    d["frecuencia"] = "A"
    d["fecha_dato"] = d.Periodo + "-01-01"
    d["variable"] = "ine_padron_poblacion_personas"
    d["unidad"] = UNI_PERS_1E
    d["desagregacion"] = "sexo=" + slug(d.Sexo)
    d = construir(d[["territorio_codigo", "nivel", "periodo", "frecuencia", "variable", "desagregacion", "valor", "unidad", "fecha_dato"]],
                  rel, man)
    return filtra_codigos(d, "municipio", t, "padron_cifras")


def leer_px(ruta: Path) -> pd.DataFrame:
    txt = ruta.read_bytes().decode("iso-8859-15")
    cab, datos = txt.split("DATA=", 1)
    vals = {}
    for m in re.finditer(r'VALUES\("([^"]+)"\)=((?:"[^"]*"\s*,?\s*)+);', cab):
        vals[m.group(1)] = re.findall(r'"([^"]*)"', m.group(2))
    stub = re.findall(r'"([^"]+)"', re.search(r"STUB=([^;]+);", cab).group(1))
    head = re.findall(r'"([^"]+)"', re.search(r"HEADING=([^;]+);", cab).group(1))
    assert len(stub) == 2 and len(head) == 1, ruta
    v = pd.to_numeric(pd.Series(datos.replace(";", " ").split()), errors="coerce").to_numpy()
    s0, s1, h0 = vals[stub[0]], vals[stub[1]], vals[head[0]]
    assert len(v) == len(s0) * len(s1) * len(h0), (ruta, len(v), len(s0), len(s1), len(h0))
    idx = pd.MultiIndex.from_product([s0, s1, h0], names=["sexo", "seccion", "nacionalidad"])
    return pd.DataFrame({"valor": v}, index=idx).reset_index()


def padron_seccion(man: dict, t: pd.DataFrame) -> pd.DataFrame:
    partes = []
    for f in sorted((ENT / "padron_seccion_2022_nacionalidad_continentes").glob("*.px")):
        rel = f"data/raw/v16/tension_social/padron_seccion_2022_nacionalidad_continentes/{f.name}"
        d = leer_px(f)
        d = d[d.seccion != "TOTAL"].dropna(subset=["valor"]).copy()
        d["territorio_codigo"], d["nivel"] = d.seccion, "seccion"
        d["periodo"], d["frecuencia"], d["fecha_dato"] = "2022", "A", "2022-01-01"
        d["variable"] = "ine_padron_seccion_poblacion_personas"
        d["unidad"] = UNI_PERS_1E
        d["desagregacion"] = "sexo=" + slug(d.sexo) + ";nacionalidad=" + slug(d.nacionalidad)
        partes.append(construir(d[["territorio_codigo", "nivel", "periodo", "frecuencia", "variable", "desagregacion", "valor", "unidad",
                                   "fecha_dato"]], rel, man))
    h = pd.concat(partes, ignore_index=True)
    return filtra_codigos(h, "seccion", t, "padron_seccion")


# ---------------------------------------------------------------- CGPJ (terceros)
def fin_trimestre(p: str) -> str:
    y, q = int(p[:4]), int(p[-1])
    return f"{y}-{3 * q:02d}-{[31, 30, 30, 31][q - 1]:02d}"


def cgpj_provincial(man: dict, t: pd.DataFrame):
    import openpyxl
    rel = "data/raw/v16/tension_social/_terceros/cgpj_series_provincias_2T2026.xlsx"
    wb = openpyxl.load_workbook(RAIZ / rel, read_only=True, data_only=True)
    codigos = set(t[t.nivel == "provincia"].territorio_codigo)
    assert set(CGPJ_PROV.values()) <= codigos
    partes, anual_total, chk = [], None, {}
    for hoja, var, _ in CGPJ_SERIES_PROV:
        filas = list(wb[hoja].iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(filas) if any(isinstance(x, str) and re.fullmatch(r"\d\d-T[1-4]", x) for x in r))
        cab = filas[hi]
        cols = [(j, "20" + h[:2] + "T" + h[-1]) for j, h in enumerate(cab) if isinstance(h, str) and re.fullmatch(r"\d\d-T[1-4]", h)]
        regs, tot = [], None
        for r in filas[hi + 1:]:
            nom = r[1]
            if nom is None:
                break
            nom_n = norm(nom) if nom != "ARABA/ALAVA" else "ARABA/ALAVA"
            if nom_n == "TOTAL":
                tot = r
                break
            for j, p in cols:
                if isinstance(r[j], (int, float)):
                    regs.append((CGPJ_PROV[nom_n if nom_n in CGPJ_PROV else nom], p, float(r[j])))
        d = pd.DataFrame(regs, columns=["territorio_codigo", "periodo", "valor"])
        if tot is not None:
            tot_n = {p: tot[j] for j, p in cols if isinstance(tot[j], (int, float))}
            s = d.groupby("periodo").valor.sum()
            chk[var] = int(sum(abs(s[p] - v) > 0.5 for p, v in tot_n.items() if p in s.index))
        d["nivel"], d["frecuencia"] = "provincia", "T"
        d["variable"] = var
        d["unidad"] = UNI_EJH if "ejecuciones" in var else UNI_LANZ
        d["fecha_dato"] = d.periodo.map(fin_trimestre)
        partes.append(construir(d, rel, man))
    return pd.concat(partes, ignore_index=True), chk


def nomenclator() -> pd.DataFrame:
    d = pd.read_csv(TER / "derivado_pj_municipio.csv", dtype=str)
    d["partido_judicial_n"] = d.partido_judicial_n.astype(int)
    d["vars"] = d.municipio.map(variantes)
    return d


def bloques_provincia(R, T, prov_nombres, P, cand, especial):
    """Reparte las filas de partidos judiciales entre las provincias (en el orden de la hoja provincial) minimizando el error
    relativo entre la suma de bloque y el total provincial anual (13 años) más una penalización por nombres cuyo municipio de
    igual nombre solo existe en otra provincia. Determinista."""
    n, m = len(R), len(prov_nombres)
    C = np.vstack([np.zeros(T.shape[1]), np.cumsum(T, 0)])
    codes = [CGPJ_PROV[p] for p in prov_nombres]
    pen = np.zeros((m, n + 1))
    for i, c in enumerate(codes):
        mal = np.array([1.0 if (cand[r] and c not in cand[r] and not especial[r]) else 0.0 for r in range(n)])
        pen[i, 1:] = np.cumsum(mal)
    INF = 1e18
    D = np.full((m + 1, n + 1), INF)
    B = np.zeros((m + 1, n + 1), int)
    D[0, 0] = 0
    for i in range(m):
        den = P[i].sum()
        for b in range(1, n + 1):
            a = np.arange(max(0, b - 90), b)
            ok = D[i, a] < INF
            if not ok.any():
                continue
            a = a[ok]
            c = D[i, a] + np.abs(C[b] - C[a] - P[i]).sum(1) / den + (pen[i, b] - pen[i, a])
            k = int(np.argmin(c))
            D[i + 1, b], B[i + 1, b] = c[k], a[k]
    b, out = n, []
    for i in range(m, 0, -1):
        a = B[i, b]
        out.append((i - 1, a, b))
        b = a
    return out[::-1], float(D[m, n])


def cgpj_pj(man: dict, t: pd.DataFrame):
    import openpyxl
    rel = "data/raw/v16/tension_social/_terceros/cgpj_lanzamientos_partidos_judiciales_2013_2025.xlsx"
    wb = openpyxl.load_workbook(RAIZ / rel, read_only=True, data_only=True)
    filas = list(wb.worksheets[0].iter_rows(values_only=True))
    anios = [int(x) for x in filas[5] if x is not None]
    assert anios == list(range(2013, 2026)), anios
    R = [r for r in filas[7:] if r[0]]
    nombres = [str(r[0]).strip() for r in R]
    T = np.array([[r[1 + 4 * k] or 0 for k in range(13)] for r in R], float)
    # totales provinciales (suma trimestral de la hoja provincial) para repartir en bloques
    wp = openpyxl.load_workbook(RAIZ / "data/raw/v16/tension_social/_terceros/cgpj_series_provincias_2T2026.xlsx", read_only=True,
                                data_only=True)
    fp = list(wp["Lanzamientos pract. Total prov"].iter_rows(values_only=True))
    hdr = fp[6]
    prov, P = [], []
    for r in fp[7:]:
        if r[1] == "TOTAL":
            break
        prov.append(r[1] if r[1] == "ARABA/ALAVA" else norm(r[1]))
        P.append([sum(v for h, v in zip(hdr, r) if isinstance(h, str) and re.fullmatch(r"\d\d-T[1-4]", h) and h.startswith(str(y)[2:])
                      and isinstance(v, (int, float))) for y in anios])
    P = np.array(P, float)
    nm = nomenclator()
    # candidatos por fila: provincias en las que el nombre coincide con algún municipio
    idx: dict[str, set] = {}
    for c, vs in zip(nm.cod_provincia, nm.vars):
        for v in vs:
            idx.setdefault(v, set()).add(c)
    cand, especial = [], []
    for n_ in nombres:
        vs = variantes(n_)
        cand.append(set().union(*[idx.get(v, set()) for v in vs]))
        especial.append(norm(n_) in CGPJ_PJ_ESPECIAL)
    bl, coste = bloques_provincia(R, T, prov, P, cand, especial)
    prov_fila = [None] * len(R)
    for i, a, b in bl:
        for r in range(a, b):
            prov_fila[r] = CGPJ_PROV[prov[i]]
    # cruce nombre -> número del partido, coincidencia única de nombre normalizado dentro de la provincia
    estados, asign = [], {}
    for r, n_ in enumerate(nombres):
        pr = CGPJ_PJ_ESPECIAL.get(norm(n_), prov_fila[r])
        sub = nm[nm.cod_provincia == pr]
        vs = variantes(n_)
        hit = sub[sub.vars.map(lambda x: bool(x & vs))]
        if len(hit) == 1:
            estados.append((n_, pr, int(hit.partido_judicial_n.iloc[0]), "coincidencia_unica"))
            asign.setdefault((pr, int(hit.partido_judicial_n.iloc[0])), []).append(r)
        else:
            estados.append((n_, pr, None, "sin_coincidencia" if len(hit) == 0 else "ambigua_%d" % len(hit)))
    dup = {k for k, v in asign.items() if len(v) > 1}
    filas_v, desc = [], []
    for r, (n_, pr, pj, est) in enumerate(estados):
        if pj is not None and (pr, pj) in dup:
            est = "descartada_pj_duplicado"
            pj = None
        if pj is None:
            desc.append((n_, pr, est))
            continue
        cod = f"PJ-{pr}-{pj}"
        for k, y in enumerate(anios):
            for j, var in enumerate(["cgpj_lanzamientos_total_numero", "cgpj_lanzamientos_hipotecaria_numero", "cgpj_lanzamientos_lau_numero",
                                     "cgpj_lanzamientos_otros_numero"]):
                v = R[r][1 + 4 * k + j]
                if isinstance(v, (int, float)):
                    filas_v.append((cod, str(y), var, float(v)))
    d = pd.DataFrame(filas_v, columns=["territorio_codigo", "periodo", "variable", "valor"])
    d["nivel"], d["frecuencia"], d["unidad"] = "partido_judicial", "A", UNI_LANZ
    d["fecha_dato"] = d.periodo + "-12-31"
    d["metodo_dato"] = "observado (partido identificado por nombre de cabecera; cruce con el nomenclátor a 30/12/2025)"
    d = construir(d, rel, man)
    # 2025: la suma por provincia de los partidos difiere más de un 10 % de la serie provincial en 24 provincias (p. ej. PJ-32-2 = 2944)
    d.loc[d.periodo == "2025", "metodo_dato"] += "; 2025 discrepante con la serie provincial del CGPJ (ver ficha)"
    # correspondencia partido -> municipios (todos los partidos del nomenclátor) y descartes del CGPJ
    nombre_por_pj = {(pr, pj): nombres[rs[0]] for (pr, pj), rs in asign.items() if (pr, pj) not in dup}
    cm = nm[["cod_provincia", "partido_judicial_n", "cod_municipio_ine", "municipio"]].copy()
    cm["territorio_codigo"] = "PJ-" + cm.cod_provincia + "-" + cm.partido_judicial_n.astype(str)
    cm["nombre_cgpj"] = [nombre_por_pj.get((a, b)) for a, b in zip(cm.cod_provincia, cm.partido_judicial_n)]
    cm["estado"] = np.where(cm.nombre_cgpj.notna(), "partido_con_serie_cgpj", "partido_sin_serie_cgpj")
    tm = set(t[t.nivel == "municipio"].territorio_codigo)
    cm["municipio_en_territorios"] = cm.cod_municipio_ine.isin(tm)
    cm = cm.rename(columns={"municipio": "municipio_nomenclator"})
    dd = pd.DataFrame(desc, columns=["nombre_cgpj", "cod_provincia", "estado"])
    dd["estado"] = "cgpj_descartado:" + dd.estado
    cm = pd.concat([cm, dd], ignore_index=True)
    cm["partido_judicial_n"] = cm.partido_judicial_n.astype("Int64")
    info = {"bloques_coste": coste, "filas_cgpj": len(R), "con_serie": len(R) - len(desc), "descartadas": desc,
            "partidos_nomenclator": int(nm.groupby("cod_provincia").partido_judicial_n.nunique().sum())}
    return d, cm, info


# ---------------------------------------------------------------- diccionario
DESC = {
    "ine_emcr_inmigraciones_exterior_personas": "EMCR: inmigraciones procedentes del extranjero por municipio (tabla 69696)",
    "ine_emcr_emigraciones_exterior_personas": "EMCR: emigraciones con destino al extranjero por municipio (tabla 69711)",
    "ine_emcr_migraciones_interiores_llegadas_personas": "EMCR: llegadas desde otro municipio español (destino); por nacionalidad (69743) o país de nacimiento (69744)",
    "ine_emcr_migraciones_interiores_salidas_personas": "EMCR: salidas hacia otro municipio español (origen); por nacionalidad (69745) o país de nacimiento (69746)",
    "ine_emcr_saldo_migratorio_personas": "EMCR: saldo migratorio total, exterior e interior (tabla 69767)",
    "ine_padron_poblacion_personas": "Cifras oficiales del padrón por municipio y sexo, 1 de enero (tabla 29005)",
    "ine_padron_seccion_poblacion_personas": "Padrón a 1/1/2022 por sección censal, sexo y nacionalidad (continentes)",
    "cgpj_lanzamientos_total_numero": "CGPJ: lanzamientos practicados, total",
    "cgpj_lanzamientos_hipotecaria_numero": "CGPJ: lanzamientos practicados como consecuencia de ejecución hipotecaria",
    "cgpj_lanzamientos_lau_numero": "CGPJ: lanzamientos practicados como consecuencia de la Ley de Arrendamientos Urbanos",
    "cgpj_lanzamientos_otros_numero": "CGPJ: otros lanzamientos practicados",
    "cgpj_ejecuciones_hipotecarias_numero": "CGPJ: ejecuciones hipotecarias presentadas (juzgados de primera instancia)",
}
RUPT = {
    "ine_emcr": "EMCR 2021-2024: ruptura con la EVR (1988-2021, v11): otra operación y metodología; no unir sin análisis. Municipios según la relación de cada año",
    "ine_padron_poblacion": "1 de enero; municipios fusionados o segregados sin código en territorios descartados; sin dato en huecos del INE",
    "ine_padron_seccion": "seccionado de 1/1/2022 frente al seccionado 2025 de territorios; secciones ausentes descartadas; nacionalidad por continentes",
    "cgpj_lanzamientos": "terceros (licencia dudosa), no versionado. Provincia: Cádiz incluye Ceuta y Málaga incluye Melilla; partido judicial: cambios de demarcación 2013-2025",
    "cgpj_ejecuciones": "terceros (licencia dudosa), no versionado; serie desde 2007T1; provincia con Ceuta y Melilla incluidas en Cádiz y Málaga",
}


def diccionario(frames: list[pd.DataFrame]) -> pd.DataFrame:
    filas = []
    for df in frames:
        for (v, f, n), g in df.groupby(["variable", "frecuencia", "nivel"]):
            clave = next((k for k in sorted(RUPT, key=len, reverse=True) if v.startswith(k)), "")
            filas.append({"variable": v, "descripcion": DESC.get(v, v), "unidad": g.unidad.iloc[0], "nivel": n, "frecuencia": f,
                          "periodo_min": g.periodo.min(), "periodo_max": g.periodo.max(), "fuente": g.fuente.iloc[0],
                          "rupturas": RUPT.get(clave, "")})
    return pd.DataFrame(filas)


def ordenar(h: pd.DataFrame) -> pd.DataFrame:
    h["muestra_sellada_v6"] = sellada(h)
    return h[COLS].sort_values(["variable", "nivel", "territorio_codigo", "desagregacion", "periodo"], kind="mergesort").reset_index(drop=True)


def main() -> int:
    man = manifiesto()
    t = territorios()
    h = ordenar(pd.concat([emcr(man, t), padron_cifras(man, t), padron_seccion(man, t)], ignore_index=True))
    SAL.mkdir(parents=True, exist_ok=True)
    h.to_parquet(SAL / "tension_social.parquet", index=False)
    ruta_t = SAL / "tension_social_terceros.parquet"
    ruta_c = SAL / "tension_social_partidos_municipios_terceros.parquet"
    tt, info, chk = None, None, None
    if (TER / "cgpj_series_provincias_2T2026.xlsx").exists() and (TER / "cgpj_lanzamientos_partidos_judiciales_2013_2025.xlsx").exists() \
            and (TER / "derivado_pj_municipio.csv").exists():
        try:
            import openpyxl  # noqa: F401
            pv, chk = cgpj_provincial(man, t)
            pj, cm, info = cgpj_pj(man, t)
            tt = ordenar(pd.concat([pv, pj], ignore_index=True))
            tt["redistribuible"] = False
            tt.to_parquet(ruta_t, index=False)
            cm.to_parquet(ruta_c, index=False)
        except ImportError as e:
            print("CGPJ no procesado (falta openpyxl):", e)
    else:
        for r in (ruta_t, ruta_c):
            if r.exists():
                r.unlink()
    diccionario([h] + ([tt] if tt is not None else [])).to_csv(DICC, index=False)
    res = {"rama": "v16-f1-datos-tension_social",
           "pregunta": "Normalización al formato largo: EMCR municipal 2021-2024, padrón municipal 1996-2025 y por sección 2022 (INE, versionado); "
                       "lanzamientos y ejecuciones hipotecarias del CGPJ por provincia y partido judicial (terceros, no versionado)",
           "capa": "C4", "N": int(len(h)), "metodo": "normalización",
           "notas": f"N_terceros={0 if tt is None else len(tt)} (no versionado, redistribuible=False). posterior_prereg=False porque el corte de V16 no "
                    "está fijado. Descartes de códigos: " + json.dumps({k: {"descartados": len(v["descartados_no_en_territorios"]),
                                                                        "conservados_no_vigentes": len(v["conservados_con_vigente_2025_false"])}
                                                                       for k, v in DESCARTES.items()}, ensure_ascii=False)
                    + ("" if info is None else f". Partidos judiciales del CGPJ: {info['con_serie']} de {info['filas_cgpj']} con cruce único; "
                       f"{len(info['descartadas'])} descartados (lista en la ficha)")}
    (SAL / "tension_social_resultado.json").write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"filas versionadas={len(h)} terceros={0 if tt is None else len(tt)}")
    print(json.dumps(DESCARTES, ensure_ascii=False)[:3000])
    if info is not None:
        print("control sumas provincia vs TOTAL (trimestres discrepantes):", chk)
        print({k: v for k, v in info.items() if k != "descartadas"})
        for x in info["descartadas"]:
            print("DESCARTADO", x)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
