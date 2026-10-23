"""Consolidación final · `make revision_externa_2_0`: paquete de la revisión externa de la publicación 2.0 (revision_externa/2_0/).

1. Muestra de 40 cifras del documento final (las que citan sus plantillas, con valor o rango), estratificada por capa y por clase V18
   con SEED=20261010. Estratos: C1, C2, C3s, C4 sin clase V18 (versiones anteriores a V18) y, dentro de las cifras V18 (todas C4),
   SUGESTIVO, ASOCIACIÓN, NO CONCLUYENTE, NO ESTIMABLE y sin clase propia (literatura, diagnóstico, robustez, coprincipal). Cada
   estrato recibe min(4, fichas disponibles); las plazas que sobran van, de una en una, a los estratos con más fichas libres. Una
   cifra por ficha, como en la 1.0 (src/v14/muestra_revision.py: grupo).
2. revision_externa/2_0/muestra.csv y hoja_revision.xlsx (misma hoja que la 1.0, con la clase V18 junto a la capa).
3. revision_externa/2_0/resumen_ejecutivo.pdf (del resumen ejecutivo publicado) y guia.pdf (de revision_externa/2_0/guia.md), con
   pandoc y Chromium sin interfaz. El PDF no es byte a byte reproducible (lleva la fecha de creación); el texto sí.
Sin red; no cambia ninguna cifra (solo lee cifras_clave y las plantillas del documento final).
"""
from __future__ import annotations

import glob
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd

RAIZ = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ))
from src.v14 import muestra_revision as M  # noqa: E402

SEED = 20261010
N_TOTAL, MINIMO = 40, 4
SAL = RAIZ / "revision_externa" / "2_0"
PLANT = RAIZ / "docs" / "documento_final" / "plantilla"
RESUMEN = RAIZ / "publicacion" / "documento_final" / "resumen_ejecutivo.md"
MARC = re.compile(r"\{\{([A-Za-z0-9_.\-]+)(?::[a-z_]+)?\}\}")
SIN_CLASE_PROPIA = ("literatura (sin clase propia)", "diagnóstico o variante (sin clase propia)", "robustez a posteriori",
                    "sin clasificar (coprincipal informativa)")
ESTRATOS = ["C1", "C2", "C3s", "C4 sin clase V18", "V18 SUGESTIVO", "V18 ASOCIACIÓN", "V18 NO CONCLUYENTE", "V18 NO ESTIMABLE",
            "V18 sin clase propia"]
CSS = """body{font-family:"DejaVu Sans",Arial,sans-serif;font-size:7.6pt;line-height:1.28;margin:0}
h1{font-size:11pt;margin:0 0 4pt}h2{font-size:9.5pt;margin:7pt 0 3pt}p,li{margin:2pt 0}
table{border-collapse:collapse;width:100%;margin:4pt 0}th,td{border:0.5pt solid #999;padding:1.5pt 3pt;vertical-align:top}
th{background:#e7eef7}blockquote{margin:3pt 8pt;color:#333}@page{size:A4 landscape;margin:9mm 10mm}
td:nth-child(1){width:10%}td:nth-child(2){width:31%}td:nth-child(3){width:33%}td:nth-child(4){width:10%}td:nth-child(5){width:16%}
h2[id^="executive-summary"]{break-before:page}"""


def citadas() -> set[str]:
    ids = {f"V18-B5-E{k}" for k in range(1, 10)}   # {{tabla.impactos}}
    for p in sorted(PLANT.glob("[0-9][0-9]_*.md")):
        ids |= {m.group(1) for m in MARC.finditer(p.read_text(encoding="utf-8")) if not m.group(1).startswith(("proyecto.", "tabla."))}
    return ids


def estrato(r: pd.Series) -> str:
    if r.capa != "C4" or not str(r.id).startswith("V18-"):
        return "C4 sin clase V18" if r.capa == "C4" else r.capa
    c = r.clasificacion_v18
    return "V18 sin clase propia" if c in SIN_CLASE_PROPIA else f"V18 {c}"


def muestra() -> pd.DataFrame:
    ck = pd.read_csv(RAIZ / "output" / "v5" / "cifras_clave.csv")
    ck = ck[ck.id.isin(citadas()) & (ck.valor.notna() | ck["min"].notna())].copy()
    anexo = M.fichas_anexo()
    ck["grupo"] = [M.grupo(r, anexo) for r in ck.itertuples()]
    ck["estrato"] = [estrato(r) for _, r in ck.iterrows()]
    rng = np.random.default_rng(SEED)
    libres = {}
    for e in ESTRATOS:
        x = ck[ck.estrato == e].sort_values("id")
        libres[e] = x.iloc[rng.permutation(len(x))].drop_duplicates("grupo")   # una cifra por ficha
    cupo = {e: min(MINIMO, len(libres[e])) for e in ESTRATOS}
    while sum(cupo.values()) < N_TOTAL:
        e = max(ESTRATOS, key=lambda k: (len(libres[k]) - cupo[k], -ESTRATOS.index(k)))
        if len(libres[e]) - cupo[e] <= 0:
            break
        cupo[e] += 1
    m = pd.concat([libres[e].head(cupo[e]) for e in ESTRATOS]).reset_index(drop=True)
    m["ficha_anexo"] = [f"{anexo[i][0]} · metodologia/anexo_metodologico.md#{anexo[i][1]}" if i in anexo else f"script de origen: {o}"
                        for i, o in zip(m.id, m.origen)]
    m["enunciado"] = [f"{r.indicador}: {M.R.texto(r, 'cita')}" for _, r in m.iterrows()]
    m["organismo"] = [next((o for o, _ in M.ORGANISMOS if o.lower() in str(f).lower()), "propio (elaboración del proyecto)") for f in m.fuentes]
    return m


def chromium() -> str:
    for c in [os.environ.get("CHROMIUM", ""), *sorted(glob.glob("/opt/pw-browsers/chromium_headless_shell-*/chrome-linux/headless_shell")),
              *sorted(glob.glob("/opt/pw-browsers/chromium-*/chrome-linux/chrome")), "chromium", "chromium-browser", "google-chrome"]:
        if c and (Path(c).exists() or subprocess.run(["which", c], capture_output=True).returncode == 0):
            return c
    raise SystemExit("revision_externa: no hay Chromium (variable CHROMIUM) para generar los PDF")


def pdf(md: str, titulo: str, destino: Path) -> None:
    with tempfile.TemporaryDirectory() as d:
        src, css, htm = Path(d) / "t.md", Path(d) / "e.css", Path(d) / "t.html"
        src.write_text(md, encoding="utf-8")
        css.write_text(CSS, encoding="utf-8")
        subprocess.run(["pandoc", str(src), "-f", "markdown-implicit_figures", "-t", "html5", "-s", "--css", str(css),
                        "--embed-resources", "--metadata", f"pagetitle={titulo}", "-o", str(htm)], check=True)
        subprocess.run([chromium(), "--headless", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={destino}", htm.as_uri()], check=True, capture_output=True)


def main() -> int:
    SAL.mkdir(parents=True, exist_ok=True)
    m = muestra()
    cols = ["id", "estrato", "capa", "clasificacion_v18", "grupo", "indicador", "valor", "min", "max", "unidad", "periodo", "fuentes",
            "organismo", "ficha_anexo", "origen"]
    m[cols].to_csv(SAL / "muestra.csv", index=False)
    h = m.copy()
    h["capa"] = [f"{c} · {e}" for c, e in zip(h.capa, h.estrato)]   # la clase V18 va con la capa en la hoja
    M.hoja(h, SAL / "hoja_revision.xlsx")
    from openpyxl import load_workbook   # la pestaña «leeme» de la 1.0 describe otra muestra: se reescribe
    wb = load_workbook(SAL / "hoja_revision.xlsx")
    le = wb["leeme"]
    for fila, txt in enumerate(("Hoja para la revisión externa de la publicación 2.0 (documento final; preparada, sin publicar).",
                                "Una fila por cifra: 40 cifras citadas en el documento final, de output/v5/cifras_clave.csv, estratificadas por capa "
                                "y por clase V18 (columna «capa»), una por ficha, SEED=20261010.",
                                "Colores: azul = dato fijo copiado de cifras_clave (fuente en la columna «fuente(s)»); negro = cálculo; verde = viene de otra pestaña.",
                                "Rellene solo las tres últimas columnas de «revision». Guía: revision_externa/2_0/guia.md (y guia.pdf)."), start=1):
        le.cell(fila, 1, txt)
    wb.properties.creator = "vivienda-espana-evidencia 2.0 (src/documento_final/revision_externa.py)"
    wb.save(SAL / "hoja_revision.xlsx")
    M.fijar_zip(SAL / "hoja_revision.xlsx")
    res = RESUMEN.read_text(encoding="utf-8")
    res = res.split("\n", 2)[-1] if res.startswith("<!--") else res
    pdf(res, "Resumen ejecutivo del documento final", SAL / "resumen_ejecutivo.pdf")
    pdf((SAL / "guia.md").read_text(encoding="utf-8"), "Guía para la revisión externa (2.0)", SAL / "guia.pdf")
    print("revision_externa_2_0:", m.estrato.value_counts().reindex(ESTRATOS).fillna(0).astype(int).to_dict(),
          f"({len(m)} cifras, {m.grupo.nunique()} fichas distintas); PDF en {SAL.relative_to(RAIZ)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
