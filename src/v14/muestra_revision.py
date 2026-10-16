"""V14-B3 · `make muestra_revision`: paquete para la revisión externa (revision_externa/).

1. Muestra aleatoria estratificada por capa (C1, C2, C3s, C4) de 40 filas de output/v5/cifras_clave.csv, con SEED=20261010.
   - Reparto: 10 por capa; si una capa tiene menos grupos distintos, sus plazas pasan a C4.
   - Se excluyen los duplicados de la misma ficha: como mucho una cifra por grupo. El grupo es la ficha del anexo
     metodológico si la cifra la tiene; si no, la ficha del verificador (columna `ficha`); si tampoco, el id sin el sufijo
     de variante (-est/-pholm/-min/…).
2. revision_externa/hoja_revision.xlsx: una fila por cifra, con columnas vacías para el revisor.
   - Datos fijos en azul con su fuente; cálculos en negro; lo que viene de otra pestaña, en verde.
   - Sin paneles inmovilizados.
3. revision_externa/muestra.csv: la muestra en texto plano.
Determinista y sin red; no cambia ninguna cifra (solo lee cifras_clave).
"""
from __future__ import annotations

import datetime as dt
import re
import sys
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ / "src" / "v5"))
import render as R  # noqa: E402

SEED = 20261010
N_TOTAL, CAPAS = 40, ("C1", "C2", "C3s", "C4")
SAL = RAIZ / "revision_externa"
ANEXO = RAIZ / "metodologia" / "anexo_metodologico.md"
FECHA_FIJA = dt.datetime(2026, 10, 16)
AZUL, NEGRO, VERDE = Font(color="1F4E99"), Font(color="000000"), Font(color="1E7B34")
NEGRITA, CABECERA = Font(bold=True), PatternFill("solid", fgColor="E7EEF7")
SUFIJO = re.compile(r"-(est|pholm|p|p05|bh|min|max|lo|hi|ic95|ic|n|num|valor|central|media|mediana|p16|p84|\d+)$", re.I)
# organismo -> URL de su portal estadístico o de datos abiertos (se resuelve en la pestaña «fuentes»)
ORGANISMOS = [
    ("INE", "https://www.ine.es"), ("MIVAU", "https://www.mivau.gob.es/el-ministerio/informacion-estadistica"),
    ("Ministerio", "https://www.mivau.gob.es/el-ministerio/informacion-estadistica"), ("Catastro", "https://www.catastro.hacienda.gob.es"),
    ("Eurostat", "https://ec.europa.eu/eurostat"), ("BdE", "https://www.bde.es/webbe/es/estadisticas/"),
    ("Banco de España", "https://www.bde.es/webbe/es/estadisticas/"), ("BCE", "https://data.ecb.europa.eu"),
    ("OCDE", "https://data-explorer.oecd.org"), ("BIS", "https://data.bis.org"), ("GVA", "https://dadesobertes.gva.es"),
    ("Generalitat Valenciana", "https://dadesobertes.gva.es"), ("Notariado", "https://www.notariado.org/portal/estad%C3%ADsticas"),
    ("Registradores", "https://www.registradores.org/actualidad/portal-estadistico-registral"), ("AEAT", "https://sede.agenciatributaria.gob.es"),
    ("SERPAVI", "https://www.mivau.gob.es/vivienda/alquila-bien-es-tu-derecho/serpavi"), ("SEPE", "https://www.sepe.es"),
    ("Seguridad Social", "https://www.seg-social.es"), ("Ajuntament de València", "https://www.valencia.es/estadistica"),
    ("Incasòl", "https://analisi.transparenciacatalunya.cat"), ("CGPJ", "https://www.poderjudicial.es/cgpj/es/Temas/Estadistica-Judicial/"),
    ("Interior", "https://www.interior.gob.es"), ("BOE", "https://www.boe.es"), ("Inside Airbnb", "https://insideairbnb.com"),
]


def _slug(h: str) -> str:
    return re.sub(r"[^\w\- ]", "", h.strip().lower()).replace(" ", "-")


def fichas_anexo() -> dict[str, tuple[str, str]]:
    """{id: (ficha, ancla)} de las fichas «### Ax.» del anexo metodológico."""
    out = {}
    for ln in ANEXO.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^###\s+((A\d+(?:-A\d+)?)\.\s+(.+?)):", ln)
        if m:
            titulo = ln.lstrip("#").strip()
            for i in re.findall(r"[A-Za-z0-9_]+(?:-[A-Za-z0-9_]+)*", m.group(3)):
                out.setdefault(i, (m.group(2), _slug(titulo)))
    return out


def grupo(r: pd.Series, anexo: dict) -> str:
    if r.id in anexo:
        return "anexo:" + anexo[r.id][0]
    if isinstance(r.ficha, str) and r.ficha not in ("—", ""):
        return "ficha:" + r.ficha
    i = r.id
    while SUFIJO.search(i):   # V8-SVAR-…-media y -mediana, VER-A23-V2-1 y -2: misma ficha
        i = SUFIJO.sub("", i)
    return "id:" + i


def muestra() -> pd.DataFrame:
    ck = pd.read_csv(RAIZ / "output" / "v5" / "cifras_clave.csv")
    anexo = fichas_anexo()
    ck["grupo"] = [grupo(r, anexo) for r in ck.itertuples()]
    rng = np.random.default_rng(SEED)
    cupo = {c: N_TOTAL // len(CAPAS) for c in CAPAS}
    elegidas = []
    for c in CAPAS:
        x = ck[ck.capa == c].sort_values("id")
        x = x.iloc[rng.permutation(len(x))].drop_duplicates("grupo")   # una cifra por ficha
        if len(x) < cupo[c]:
            cupo["C4"] += cupo[c] - len(x)
        elegidas.append((c, x))
    filas = [x.head(cupo[c]) for c, x in elegidas]
    m = pd.concat(filas).reset_index(drop=True)
    m["ficha_anexo"] = [f"{anexo[i][0]} · metodologia/anexo_metodologico.md#{anexo[i][1]}" if i in anexo else f"script de origen: {o}"
                        for i, o in zip(m.id, m.origen)]
    m["enunciado"] = [f"{r.indicador}: {R.texto(r, 'cita')}" for _, r in m.iterrows()]
    m["organismo"] = [next((o for o, _ in ORGANISMOS if o.lower() in str(f).lower()), "propio (elaboración del proyecto)") for f in m.fuentes]
    return m


def fijar_zip(ruta: Path) -> None:
    with zipfile.ZipFile(ruta) as z:
        datos = [(n, z.read(n)) for n in sorted(z.namelist())]
    fijo = FECHA_FIJA.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    datos = [(n, re.sub(rb"(<dcterms:(?:modified|created)[^>]*>)[^<]*", rb"\g<1>" + fijo, b) if n == "docProps/core.xml" else b) for n, b in datos]
    tmp = Path(str(ruta) + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in datos:
            zi = zipfile.ZipInfo(n, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            z.writestr(zi, b)
    tmp.replace(ruta)


def hoja(m: pd.DataFrame, ruta: Path) -> None:
    wb = Workbook()
    wb.properties.creator = "vivienda-espana-evidencia V14 (src/v14/muestra_revision.py)"
    wb.properties.lastModifiedBy = None
    wb.properties.created = wb.properties.modified = FECHA_FIJA
    wb.properties.title = "Revisión externa: muestra de cifras"
    ws = wb.active
    ws.title = "revision"
    fu = wb.create_sheet("fuentes")
    le = wb.create_sheet("leeme")
    # pestaña «fuentes»: organismo y URL (datos fijos, azul)
    fu.append(["organismo", "url"])
    for c in fu[1]:
        c.font, c.fill = NEGRITA, CABECERA
    orgs = sorted(set(m.organismo))
    url = dict(ORGANISMOS)
    for o in orgs:
        fu.append([o, url.get(o, "—")])
        for c in fu[fu.max_row]:
            c.font = AZUL
    n_org = fu.max_row
    cab = ["id", "enunciado (tal como se publica)", "valor", "mínimo", "máximo", "unidad", "periodo", "capa", "fuente(s)", "organismo",
           "URL de la fuente", "ficha del anexo", "amplitud del rango", "valor comprobado (revisor)", "¿coincide? (S/N)", "comentario (revisor)"]
    ws.append(cab)
    for c in ws[1]:
        c.font, c.fill, c.alignment = NEGRITA, CABECERA, Alignment(wrap_text=True, vertical="top")
    for r in m.itertuples():
        f = ws.max_row + 1
        ws.append([r.id, r.enunciado, r.valor, r.min, r.max, r.unidad, r.periodo, r.capa, r.fuentes, r.organismo,
                   f'=IFERROR(INDEX(fuentes!$B$2:$B${n_org},MATCH(J{f},fuentes!$A$2:$A${n_org},0)),"—")',
                   r.ficha_anexo, f'=IF(AND(ISNUMBER(D{f}),ISNUMBER(E{f})),E{f}-D{f},"")', None, None, None])
        for col in range(1, 11):
            ws.cell(f, col).font = AZUL          # datos fijos copiados de cifras_clave, con su fuente en la columna I
        ws.cell(f, 11).font = VERDE              # viene de la pestaña «fuentes»
        ws.cell(f, 12).font = AZUL
        ws.cell(f, 13).font = NEGRO              # cálculo
        for col in range(1, 17):
            ws.cell(f, col).alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"S,N"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"O2:O{ws.max_row}")
    for col, ancho in zip("ABCDEFGHIJKLMNOP", (22, 70, 12, 12, 12, 14, 14, 7, 40, 22, 40, 50, 12, 18, 12, 40)):
        ws.column_dimensions[col].width = ancho
    for txt in ("Hoja para la revisión externa (versión pública del proyecto «Vivienda en España: base de evidencia reproducible»).",
                "Una fila por cifra de la muestra (40 cifras de output/v5/cifras_clave.csv, estratificadas por capa, SEED=20261010).",
                "Colores: azul = dato fijo copiado de cifras_clave (fuente en la columna «fuente(s)»); negro = cálculo; verde = viene de otra pestaña.",
                "Rellene solo las tres últimas columnas de «revision». Guía: revision_externa/guia.md."):
        le.append([txt])
    le.column_dimensions["A"].width = 130
    ws.freeze_panes = None
    wb.save(ruta)
    fijar_zip(ruta)


def main() -> int:
    SAL.mkdir(exist_ok=True)
    m = muestra()
    cols = ["id", "capa", "grupo", "indicador", "valor", "min", "max", "unidad", "periodo", "fuentes", "organismo", "ficha_anexo", "origen"]
    m[cols].to_csv(SAL / "muestra.csv", index=False)
    hoja(m, SAL / "hoja_revision.xlsx")
    print("muestra_revision:", m.capa.value_counts().reindex(list(CAPAS)).to_dict(), f"({len(m)} cifras, {m.grupo.nunique()} fichas distintas)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
