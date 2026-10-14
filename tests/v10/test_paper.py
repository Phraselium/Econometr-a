"""V10-B6: el manuscrito de Quarto es ejecutable (bloques de código y código en línea), sus figuras existen, sus citas
están en referencias.bib y las secciones de v9 están marcadas como pendientes. Excel del atlas determinista."""
from __future__ import annotations

import os
import re
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[2]
QMD = RAIZ / "paper" / "manuscrito.qmd"


def test_manuscrito_ejecutable_y_completo():
    txt = QMD.read_text(encoding="utf-8")
    bloques = re.findall(r"```\{python\}\n(.*?)```", txt, re.S)
    entorno: dict = {}
    cwd = os.getcwd()
    os.chdir(QMD.parent)
    try:
        for b in bloques:
            exec(compile(b, "manuscrito.qmd", "exec"), entorno)   # noqa: S102 (código propio del manuscrito)
        for expr in re.findall(r"`\{python\} ([^`]+)`", txt):
            assert str(eval(expr, entorno)), expr   # noqa: S307
    finally:
        os.chdir(cwd)
    for fig in re.findall(r"\]\((\.\./figuras/[^)]+)\)", txt):
        assert (QMD.parent / fig).exists(), fig
    bib = (RAIZ / "docs/v9_prep/referencias.bib").read_text(encoding="utf-8")
    claves = set(re.findall(r"@\w+\{([^,]+),", bib))
    for c in re.findall(r"@([A-Za-z][\w]+)", re.sub(r"```.*?```", "", txt, flags=re.S)):
        if c in ("article", "book", "techreport", "unpublished", "refs"):
            continue
        assert c in claves, c
    assert txt.count("[PENDIENTE]") >= 6 and "## Abstract" in txt


def test_excel_atlas_determinista():
    import hashlib
    import sys
    sys.path.insert(0, str(RAIZ / "src" / "v10"))
    import tablas_excel
    ruta = tablas_excel.SALIDA
    tablas_excel.main()
    a = hashlib.sha256(ruta.read_bytes()).hexdigest()
    tablas_excel.main()
    assert a == hashlib.sha256(ruta.read_bytes()).hexdigest()
    from openpyxl import load_workbook
    wb = load_workbook(ruta)
    assert all(ws.freeze_panes is None for ws in wb)
    ws = wb["IPVA_nuevo_existente"]
    assert ws["A5"].font.color.rgb.endswith("1F4E99") and str(ws["H5"].value).startswith("=")
    assert wb["Indice"]["C5"].font.color.rgb.endswith("1E7B34")
