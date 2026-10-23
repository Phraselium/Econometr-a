"""V14 · Series de terceros no redistribuibles (Notariado, Registradores, AEAT, CGPJ) en tablas propias de output/.

- `columnas_terceros(cabecera)`: columnas cuyo nombre contiene, como palabra (separada por «_» o espacios), notariado,
  registradores, cgn, cgpj o aeat. «p_hipotesis_preregistrada» o «especificaciones_registradas» no cuentan.
- `REVISADAS_SIN_SERIES`: tablas revisadas a mano cuyas columnas con ese nombre son resultados propios, no series de la
  fuente. Por ejemplo, la cuota de R² de Shapley de un modelo con el precio de Registradores.
- `redactar(origen, destino)`: copia la tabla sin esas columnas. Así el árbol público y el depósito conservan cada ruta
  publicada (enlaces del verificador, del anexo y de consulta/) sin ninguna serie de esas fuentes (D5 y D4).
- `REQUIEREN_TERCEROS`: ficheros de datos que solo existen tras `make datos_terceros` y la reconstrucción de los paneles.
  tests/conftest.py omite, con su motivo, los tests que los necesitan cuando faltan.
- V14.2 · `accion(rel, p)`: qué se hace con cada fichero de output/ en el árbol público y en el paquete, según la fuente y la
  política de la 1.0 (proyecto.json, terceros_1_0; src/v14/terceros_1_0.py):
  - «redactar»: se quitan las columnas que son la serie de una fuente que no se redistribuye (por su nombre, o por la lista
    `COLUMNAS_SERIE` cuando el nombre no la nombra);
  - «vaciar»: `TABLAS_SERIE` (ficheros que son la serie o la reproducen) quedan con la cabecera o una nota, salvo
    «redistribuir»; `AGREGADOS_PROPIOS` (sumas propias de la serie), solo con «excluir»;
  - «copiar»: lo demás. CGPJ no tiene escenario: nunca se redistribuye.
  Las listas salen del escaneo de V14.2 (src/v14/escaneo_terceros.py: columnas, nombres de fichero y valores).
"""
from __future__ import annotations

import csv
import gzip
import io
import re
from pathlib import Path

PALABRA = re.compile(r"(?:^|[_\s])(notariado|registradores|registr|cgn|cgpj|aeat)(?:[_\s]|$)", re.I)
REVISADAS_SIN_SERIES = {
    "output/v5/B3/tablas/shapley_R2.csv": "cuota de R² (Shapley) de cada familia en modelos con el precio de Registradores: resultado propio",
    "output/v5/B3/_smoke/tablas/shapley_R2.csv": "ídem (ejecución de prueba)",
}
REQUIEREN_TERCEROS = (
    "data/processed/nacional_q.csv", "data/processed/panel_ccaa_a.csv", "data/processed/v2/train/nacional_q_v2.csv",
    "data/processed/v8/nacional_q_v8.csv", "data/processed/v8/cv_prov_q_v8.csv",
    "data/raw/pdf/notariado_cgn_extranjeros_semestral.csv", "data/raw/pdf/registradores_opendata_compraventas.csv",
)


def columnas_terceros(cabecera: list[str]) -> list[str]:
    return [c for c in cabecera if PALABRA.search(c)]


def _abrir(p: Path):
    if p.name.endswith(".csv.gz"):   # V14.1: también las tablas comprimidas (antes solo se revisaban los .csv)
        return io.TextIOWrapper(gzip.open(p, "rb"), encoding="utf-8", errors="ignore", newline="")
    return p.open(encoding="utf-8", errors="ignore", newline="")


def es_tabla(p: Path) -> bool:
    return p.name.endswith((".csv", ".csv.gz"))


def cabecera(p: Path) -> list[str]:
    if not es_tabla(p):
        return []
    with _abrir(p) as fh:
        return next(csv.reader(fh), [])


FUENTE_COL = (("notariado", re.compile(r"(?:^|[_\s])(notariado|cgn)(?:[_\s]|$)", re.I)),
              ("registradores", re.compile(r"(?:^|[_\s])(registradores|registr)(?:[_\s]|$)", re.I)),   # el producto (ODbL o ERI) lo decide fuente_columna
              ("aeat", re.compile(r"(?:^|[_\s])aeat(?:[_\s]|$)", re.I)),
              ("cgpj", re.compile(r"(?:^|[_\s])cgpj(?:[_\s]|$)", re.I)))
# V14.2: columnas que son la serie de una fuente aunque su nombre no la nombre (escaneo por valores)
COLUMNAS_SERIE = {
    "output/v5/CA/tablas/contado_ratio_hipotecas_compraventas.csv": {"reg_compraventas_viv": "registradores_odbl", "reg_hipotecas_viv": "registradores_eri"},
    "output/v5/A23/tablas/tamano_implicito_registradores.csv": {"pm2": "registradores_odbl", "importe_medio_vivienda": "registradores_odbl"},
}
# V14.2: ficheros de output/ que son la serie de la fuente o la reproducen (por nombre de fichero y por valores)
TABLAS_SERIE = {
    "output/f6/notariado_municipio_vs_provincia.csv": "notariado", "output/f6/notariado_municipio_vs_provincia.md": "notariado",
    "output/f6/notariado_provincia_trimestral.csv": "notariado", "output/f6/notariado_provincia_trimestral.md": "notariado",
    "output/v4/M4/tablas/extranjeros_notariado_cv_provincias.csv": "notariado",
    "output/v12/f0/aeat_declarantes_inmuebles_aprox.csv": "aeat", "output/v12/f0/aeat_inmuebles_2013_2015.csv": "aeat",
    "output/v5/CB/tablas/aeat_irpf_capital_inmobiliario_2019_2024.csv": "aeat",
    "output/v5/CB/tablas/aeat_viviendas_arrendadas_ccaa_declarante_2024.csv": "aeat",
    "output/v5/CB/tablas/cgpj_lanzamientos_ccaa.csv": "cgpj",
    # K16 (2026-10-23): recuentos del CGPJ casi sin transformar en un resultado.json; licencia del CGPJ no verificada: se excluye
    "output/v16/auditoria/tension_social/resultado.json": "cgpj", "output/v5/CB/tablas/cgpj_verbales_posesorios_ocupacion.csv": "cgpj",
}
# V14.2: agregados propios de la serie (sumas de la Comunitat Valenciana): se publican con «agregados», se vacían con «excluir»
AGREGADOS_PROPIOS = {"output/v12/b5/aeat_cv_declarantes_inmuebles.csv": "aeat", "output/v12/b5/aeat_cv_2013_2015.csv": "aeat"}
NOMBRE_FUENTE = {"notariado": "Notariado", "registradores_odbl": "Registradores (datos abiertos)", "registradores_eri": "Registradores (Anuario ERI)",
                 "aeat": "AEAT", "cgpj": "CGPJ"}


ERI_COL = re.compile(r"extranj|_ext_|_ext$|pct_publicado|cobertura_ccaa|hipotecas", re.I)   # V14.3: columnas del Anuario ERI


def fuente_columna(c: str) -> str | None:
    f = next((f for f, pat in FUENTE_COL if pat.search(c)), None)
    if f == "registradores":   # V14.3: producto de Registradores (docs/v14/registradores_origen.md)
        return "registradores_eri" if ERI_COL.search(c) else "registradores_odbl"
    return f


def _pol(pol):
    from src.v14 import terceros_1_0 as T10
    return pol or T10.politica()


def accion(rel: str, p: Path, pol: dict | None = None) -> tuple[str, object]:
    """("copiar", None), ("redactar", [columnas]) o ("vaciar", fuente) para el fichero versionado rel."""
    if not rel.startswith("output/"):
        return ("copiar", None)
    pol = _pol(pol)
    redis = lambda f: pol.get(f) == "redistribuir"  # noqa: E731  (CGPJ y otras: nunca)
    if (f := TABLAS_SERIE.get(rel)) and not redis(f):
        return ("vaciar", f)
    if (f := AGREGADOS_PROPIOS.get(rel)) and pol.get(f) == "excluir":
        return ("vaciar", f)
    if rel in REVISADAS_SIN_SERIES or not es_tabla(p):
        return ("copiar", None)
    extra = COLUMNAS_SERIE.get(rel, {})
    cols = [c for c in cabecera(p) if (f := fuente_columna(c) or extra.get(c)) and not redis(f)]
    return ("redactar", cols) if cols else ("copiar", None)


def a_redactar(rel: str, p: Path, pol: dict | None = None) -> list[str]:
    """Columnas a quitar de la tabla rel (vacío si no hay que tocarla). Compatibilidad: ver `accion`."""
    a, v = accion(rel, p, pol)
    return v if a == "redactar" else []


def redactar(origen: Path, destino: Path, quitar: list[str]) -> None:
    with _abrir(origen) as fi:
        r = csv.reader(fi)
        cab = next(r)
        idx = [i for i, c in enumerate(cab) if c not in set(quitar)]
        destino.parent.mkdir(parents=True, exist_ok=True)
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator="\n")
        w.writerow([cab[i] for i in idx])
        for fila in r:
            w.writerow([fila[i] if i < len(fila) else "" for i in idx])
    _escribir(destino, buf.getvalue().encode("utf-8"))


def _escribir(destino: Path, datos: bytes) -> None:
    destino.parent.mkdir(parents=True, exist_ok=True)
    if destino.name.endswith(".gz"):
        with destino.open("wb") as fh, gzip.GzipFile(filename="", mode="wb", fileobj=fh, mtime=0) as gz:   # determinista
            gz.write(datos)
    else:
        destino.write_bytes(datos)


NOTA_VACIADO = ("# Contenido no redistribuido en la versión pública\n\nEste fichero reproduce la serie de {f}, que no se redistribuye en "
                "la versión pública (proyecto.json, terceros_1_0; docs/v14/escaneo_terceros.md y docs/v14/licencias.md). Se regenera en el "
                "repositorio de desarrollo con `make datos_terceros` y el script que lo escribe.\n")


def vaciar(origen: Path, destino: Path, fuente: str) -> None:
    """Tabla: solo la cabecera. Texto: una nota que explica por qué está vacío y cómo se regenera."""
    if es_tabla(origen):
        buf = io.StringIO()
        csv.writer(buf, lineterminator="\n").writerow(cabecera(origen))
        _escribir(destino, buf.getvalue().encode("utf-8"))
    else:
        _escribir(destino, NOTA_VACIADO.format(f=NOMBRE_FUENTE.get(fuente, fuente)).encode("utf-8"))


def aplicar(rel: str, origen: Path, destino: Path, pol: dict | None = None) -> str | None:
    """Copia rel a destino según `accion`. Devuelve el motivo si lo cambia (para EXCLUIDOS.csv / excluidos_arbol.csv)."""
    import shutil
    a, v = accion(rel, origen, pol)
    destino.parent.mkdir(parents=True, exist_ok=True)
    if a == "redactar":
        redactar(origen, destino, v)
        return "redactada: sin las columnas " + " ".join(v)
    if a == "vaciar":
        vaciar(origen, destino, v)
        return f"vaciada: es la serie de {NOMBRE_FUENTE.get(v, v)} (o un agregado propio con terceros_1_0 = excluir)"
    shutil.copyfile(origen, destino)
    return None


# ---------------------------------------------------------------- V14.2: escaneo de todos los formatos
FUENTE_TXT = re.compile(r"notariado|\bCGN\b|registradores|\bAEAT\b|agencia tributaria|aeat_irpf", re.I)
COLS_LARGAS = ("variable", "serie", "indicador")


def escanear(p: Path, nombre: str | None = None, datos: bytes | None = None) -> list[str]:
    """Columnas, hojas, claves o variables con nombre de Notariado, Registradores o AEAT en el fichero p (o en `datos`)."""
    import json as _json
    import tempfile
    import zipfile
    nombre = nombre or p.name
    hits: list[str] = []
    if datos is not None:   # miembro de un zip: se escribe a un temporal con la misma extensión
        suf = "".join(Path(nombre).suffixes[-2:]) if nombre.endswith(".csv.gz") else Path(nombre).suffix
        with tempfile.NamedTemporaryFile(suffix=suf, delete=False) as t:
            t.write(datos)
        try:
            return escanear(Path(t.name), nombre)
        finally:
            Path(t.name).unlink()
    try:
        if nombre.endswith((".csv", ".csv.gz")):
            hits += [f"columna {c}" for c in columnas_terceros(cabecera(p))]
        elif nombre.endswith(".gz"):
            with gzip.open(p, "rt", encoding="utf-8", errors="ignore") as fh:
                linea = fh.readline()
            hits += [f"columna {c}" for c in columnas_terceros(next(csv.reader([linea]), []))]
        elif nombre.endswith(".parquet"):
            import pyarrow.parquet as pq
            esq = pq.read_schema(p)
            hits += [f"columna {c}" for c in columnas_terceros(list(esq.names))]
            for col in COLS_LARGAS:
                if col in esq.names:
                    vals = pq.read_table(p, columns=[col]).column(col).unique().to_pylist()
                    hits += [f"{col} = {v}" for v in vals if isinstance(v, str) and PALABRA.search(v)]
        elif nombre.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(p, read_only=True, data_only=True)
            for ws in wb.worksheets:
                if PALABRA.search(ws.title.replace(" ", "_")):
                    hits.append(f"hoja {ws.title}")
                for fila in ws.iter_rows(min_row=1, max_row=6, values_only=True):
                    hits += [f"hoja {ws.title}: cabecera {c}" for c in fila if isinstance(c, str) and PALABRA.search(c.replace(" ", "_"))]
            wb.close()
        elif nombre.endswith(".zip"):
            with zipfile.ZipFile(p) as z:
                for m in z.namelist():
                    if m.endswith((".csv", ".gz", ".parquet", ".xlsx", ".json", ".zip")):
                        hits += [f"{m}: {h}" for h in escanear(p, m, z.read(m))]
        elif nombre.endswith(".json"):
            def claves(o, ruta=""):
                if isinstance(o, dict):
                    for k, v in o.items():
                        if isinstance(k, str) and PALABRA.search(k):
                            hits.append(f"clave {ruta}{k}")
                        claves(v, f"{ruta}{k}.")
                elif isinstance(o, list):
                    for x in o[:2000]:
                        claves(x, ruta)
            claves(_json.loads(p.read_text(encoding="utf-8", errors="ignore")))
        elif nombre.endswith(".duckdb"):
            import duckdb
            con = duckdb.connect(str(p), read_only=True)
            for t, c in con.execute("SELECT table_name, column_name FROM information_schema.columns").fetchall():
                if PALABRA.search(c):
                    hits.append(f"tabla {t}: columna {c}")
            con.close()
    except Exception as e:  # noqa: BLE001
        hits.append(f"ERROR al leer: {type(e).__name__}: {e}")
    return sorted(set(hits))
