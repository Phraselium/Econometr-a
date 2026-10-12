from pathlib import Path

from openpyxl import load_workbook

O = Path(__file__).resolve().parents[2] / "output" / "v7"


def test_fichas():
    assert len(list((O / "atlas").glob("*.md"))) == 17


def test_xlsx():
    wb = load_workbook(O / "tablas_regionales.xlsx")
    assert len(wb.sheetnames) == 18
    assert all(ws.freeze_panes is None for ws in wb.worksheets)
    ws = wb.worksheets[1]
    celdas = [c for r in ws.iter_rows() for c in r if c.value is not None]
    assert any(c.font.color is not None and c.font.color.rgb == "000000FF" for c in celdas)
    assert any(isinstance(c.value, str) and c.value.startswith("=") for c in celdas)
    verdes = [c for r in wb["Resumen"].iter_rows() for c in r if c.font.color is not None and c.font.color.rgb == "00008000"]
    assert verdes
