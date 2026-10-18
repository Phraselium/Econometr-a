"""V16 · Fase 7 · tablas_v16.xlsx con la convención de colores del repositorio. Sin red, determinista (fechas y zip fijos).

Convención: AZUL = dato copiado de un fichero (la columna «origen» o «fuente» da la ruta); NEGRO = fórmula de la hoja; VERDE = vínculo
a otra hoja. Sin paneles inmovilizados.
Hojas:
  LEEME        convención de colores y fuentes;
  Cifras       todas las cifras V15-* y V16-* de output/v5/cifras_clave.csv (azul);
  Resultados   un renglón por resultado.json de V15 y de las preguntas de V16 (azul);
  Capas        recuento de cifras por prefijo y capa, con fórmulas CONTAR.SI.CONJUNTO sobre «Cifras» (negro) y vínculos (verde).
Escribe output/v16/tablas_v16.xlsx; `make publicacion` lo copia a publicacion/tablas_v16.xlsx.
"""
from __future__ import annotations

import datetime as dt
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

RAIZ = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ / "src" / "v10"))
from tablas_excel import fijar_zip  # noqa: E402

SALIDA = RAIZ / "output" / "v16" / "tablas_v16.xlsx"
FECHA_FIJA = dt.datetime(2026, 10, 18)
AZUL, NEGRO, VERDE, NEGRITA = Font(color="1F4E99"), Font(color="000000"), Font(color="1E7B34"), Font(bold=True)
CAPAS = ["C1", "C2", "C3", "C3s", "C4"]


def _hoja(wb, nombre, df, fuente):
    ws = wb.create_sheet(nombre)
    ws.append([f"Fuente: {fuente}"])
    ws["A1"].font = AZUL
    ws.append(list(df.columns))
    for c in ws[2]:
        c.font = NEGRITA
    for fila in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in fila])
        for c in ws[ws.max_row]:
            c.font = AZUL
    return ws


def _resultados() -> pd.DataFrame:
    filas = []
    for p in sorted((RAIZ / "output" / "v15").glob("*/resultado.json")) + sorted((RAIZ / "output" / "v16").glob("q*/resultado.json")):
        r = json.loads(p.read_text(encoding="utf-8"))
        filas.append(dict(fichero=str(p.relative_to(RAIZ)), rama=str(r.get("rama", "")), pregunta=str(r.get("pregunta", ""))[:300],
                          capa=str(r.get("capa", "")), nivel_evidencia=str(r.get("nivel_evidencia", "")),
                          N=json.dumps(r.get("N"), ensure_ascii=False)[:200], metodo=str(r.get("metodo", ""))[:300]))
    return pd.DataFrame(filas)


def main() -> int:
    ck = pd.read_csv(RAIZ / "output" / "v5" / "cifras_clave.csv", dtype={"valor": float, "min": float, "max": float})
    v = ck[ck.id.str.startswith(("V15-", "V16-"))].reset_index(drop=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "LEEME"
    for t, f in (("Tablas V16 (src/v16/f7_tablas.py; no se edita a mano)", NEGRITA),
                 ("AZUL: dato copiado de un fichero del repositorio; la fila 1 de cada hoja da la ruta.", AZUL),
                 ("NEGRO: fórmula calculada en la hoja.", NEGRO),
                 ("VERDE: vínculo a otra hoja del libro.", VERDE),
                 ("Todas las cifras de V15 y V16 son C4 salvo las C1 descriptivas; pre-registro interno sin sello de tiempo externo.", NEGRO)):
        ws.append([t])
        ws.cell(row=ws.max_row, column=1).font = f
    _hoja(wb, "Cifras", v[["id", "indicador", "valor", "min", "max", "unidad", "periodo", "cobertura", "fuentes", "capa", "fecha_dato", "ficha", "origen"]],
          "output/v5/cifras_clave.csv (ids V15-* y V16-*)")
    _hoja(wb, "Resultados", _resultados(), "output/v15/*/resultado.json y output/v16/q*/resultado.json")
    wc = wb.create_sheet("Capas")
    n = len(v) + 2
    wc.append(["prefijo"] + CAPAS + ["total"])
    for c in wc[1]:
        c.font = NEGRITA
    for k, pref in enumerate(["V15-", "V16-"], start=2):
        wc.cell(row=k, column=1, value=pref).font = NEGRO
        for j, capa in enumerate(CAPAS, start=2):
            wc.cell(row=k, column=j, value=f'=COUNTIFS(Cifras!$A$3:$A${n},"{pref}*",Cifras!$J$3:$J${n},"{capa}")').font = NEGRO
        wc.cell(row=k, column=len(CAPAS) + 2, value=f"=SUM(B{k}:{chr(64 + len(CAPAS) + 1)}{k})").font = NEGRO
    wc.cell(row=5, column=1, value="=Cifras!A1").font = VERDE
    wc.cell(row=6, column=1, value="=Resultados!A1").font = VERDE
    wb.properties.creator = "vivienda-espana-evidencia V16 (src/v16/f7_tablas.py)"
    wb.properties.lastModifiedBy = None
    wb.properties.created = FECHA_FIJA
    wb.properties.modified = FECHA_FIJA
    wb.properties.title = "Tablas V16"
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    fijar_zip(SALIDA)
    print(f"f7_tablas: {SALIDA.relative_to(RAIZ)} ({len(v)} cifras V15/V16)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
