"""V12-B8 · atlas_v12.md (render de docs/v12/plantillas/atlas_v12.md con output/v5/cifras_clave.csv) y tablas_v12.xlsx.

publicacion/atlas_v12.md y publicacion/tablas_v12.xlsx solo se generan con `make publicacion` (src/v7_1/publicacion.py
copia estas salidas). Formato del Excel: entradas en AZUL con su fuente en la fila 2 de cada hoja, cálculos en NEGRO
como fórmulas de Excel, vínculos entre hojas en VERDE; sin paneles inmovilizados.
"""
from __future__ import annotations

import importlib.util
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

RAIZ = Path(__file__).resolve().parents[2]
O = RAIZ / "output" / "v12"
AZUL, NEGRO, VERDE = Font(color="0000FF"), Font(color="000000"), Font(color="008000")
NEGRITA = Font(bold=True)


def atlas() -> Path:
    spec = importlib.util.spec_from_file_location("render_v5", RAIZ / "src/v5/render.py")
    r = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(r)
    ck = pd.read_csv(RAIZ / "output/v5/cifras_clave.csv").set_index("id")
    txt, faltan = r.render(RAIZ / "docs/v12/plantillas/atlas_v12.md", ck)
    assert not faltan, f"marcadores sin resolver: {faltan}"
    dest = O / "atlas_v12.md"
    dest.write_text(txt, encoding="utf-8")
    return dest


def hoja(wb: Workbook, nombre: str, fuente: str, entradas: pd.DataFrame, calculos: list[tuple[str, str]]) -> dict:
    """Entradas (azul) desde la fila 4; fórmulas (negro) a la derecha; `calculos` = [(cabecera, plantilla con {r} = fila)]."""
    ws = wb.create_sheet(nombre)
    ws["A1"], ws["A2"] = nombre, f"Fuente de las entradas (azul): {fuente}. Cálculos en negro (fórmulas)."
    ws["A1"].font = NEGRITA
    cols = list(entradas.columns)
    for j, c in enumerate(cols + [h for h, _ in calculos], start=1):
        ws.cell(row=3, column=j, value=c).font = NEGRITA
    for i, fila in enumerate(entradas.itertuples(index=False), start=4):
        for j, v in enumerate(fila, start=1):
            ws.cell(row=i, column=j, value=None if pd.isna(v) else (v.item() if hasattr(v, "item") else v)).font = AZUL
        for k, (_, f) in enumerate(calculos, start=len(cols) + 1):
            ws.cell(row=i, column=k, value=f.format(r=i)).font = NEGRO
    ws.freeze_panes = None
    return {"hoja": nombre, "primera": 4, "ultima": 3 + len(entradas), "cols": cols + [h for h, _ in calculos]}


def letra(i: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(i)


def tablas() -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    res = wb.create_sheet("Resumen")
    # B1
    n = pd.read_csv(O / "b1/contabilidad_periodos_nacional.csv")
    n = n[n.nac == "TOTAL"][["periodo", "anios", "poblacion_inicio", "crecimiento", "vegetativo", "saldo_exterior", "saldo_interior"]]
    b1 = hoja(wb, "B1_poblacion", "INE: padrón y censo anual, MNP, EVR/EMCR (output/v12/b1)", n,
              [("residuo", "=D{r}-E{r}-F{r}-G{r}"), ("saldo_exterior_por_mil_anual", "=1000*F{r}/C{r}/B{r}"),
               ("vegetativo_por_mil_anual", "=1000*E{r}/C{r}/B{r}")])
    # B3
    d = pd.read_csv(O / "b3/descomposicion_2011_2021_provincia.csv", dtype={"cpro": str})
    d = d[["cpro", "hogares_2011", "hogares_2021", "poblacion_A", "estructura_A", "jefatura_A"]]
    b3 = hoja(wb, "B3_hogares", "INE, Censos 2011 y 2021 (microdatos); output/v12/b3", d,
              [("crecimiento", "=C{r}-B{r}"), ("control_suma_A", "=D{r}+E{r}+F{r}-(C{r}-B{r})"), ("pct_jefatura_A", "=100*F{r}/ABS(C{r}-B{r})")])
    # B4
    p = pd.read_csv(O / "b4/presion_provincia_2021_2025.csv", dtype={"cpro": str})[["cpro", "hogares_nuevos_ecp", "obra_neta_catastro", "terminadas_mivau"]]
    b4 = hoja(wb, "B4_oferta", "INE ECP (hogares); Catastro URBANA; MIVAU (output/v12/b4)", p,
              [("ratio_catastro_hogares", '=IF(AND(ISNUMBER(C{r}),B{r}>0),C{r}/B{r},"")'),
               ("ratio_mivau_hogares", '=IF(AND(ISNUMBER(D{r}),B{r}>0),D{r}/B{r},"")'),
               ("diferencia_catastro_mivau_pct", '=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r}),D{r}>0),100*(C{r}/D{r}-1),"")')])
    # B5
    t = pd.read_csv(O / "b5/titularidad_cuotas.csv", dtype={"cod_territorio": str})
    t = t[t.nivel.isin(["nacional", "provincia"]) & (t.anio == t.anio.max())][["nivel", "cod_territorio", "viviendas", "viv_mas10", "viv_mas100", "viv_PJ", "viv_PE"]]
    b5 = hoja(wb, "B5_titularidad", "Catastro, titularidad (uso residencial; tramos dentro del ámbito); output/v12/b5", t,
              [("pct_mas10", "=100*D{r}/C{r}"), ("pct_mas100", "=100*E{r}/C{r}"), ("pct_PJ", "=100*F{r}/C{r}"), ("pct_PE", "=100*G{r}/C{r}")])
    # B6
    b = pd.read_csv(O / "b6/ecv_brecha_misma_edad_renta.csv")[["indicador", "extranjera", "espanola_cruda", "espanola_misma_edad_y_renta", "ic95_inf", "ic95_sup"]]
    b6 = hoja(wb, "B6_condiciones", "INE, ECV (microdatos); bootstrap por hogar (output/v12/b6)", b,
              [("diferencia_cruda", "=B{r}-C{r}"), ("diferencia_misma_edad_renta", "=B{r}-D{r}")])
    # B7
    v = pd.read_csv(O / "b7/cv_municipios_vut_ine_registro.csv", dtype={"cod_municipio": str})
    u = [c for c in v.columns if c.startswith("vut_ine_") and c != "vut_ine_2024_11"][0]
    v = v[["cod_municipio", u, "registro_gva_2026-10-14", "parque_ine"]]
    b7 = hoja(wb, "B7_turismo", "INE (viviendas turísticas, experimental); registro GVA (output/v12/b7)", v,
              [("ratio_registro_ine", '=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r}),B{r}>0),C{r}/B{r},"")'),
               ("pct_ine_parque", '=IF(AND(ISNUMBER(B{r}),ISNUMBER(D{r}),D{r}>0),100*B{r}/D{r},"")')])
    # Resumen: vínculos (verde) a celdas de las otras hojas
    res["A1"], res["A2"] = "Resumen V12 (vínculos a las hojas de cálculo)", "Celdas en verde: vínculos a otras hojas. Descriptivo (C4)."
    res["A1"].font = NEGRITA
    filas = [("Crecimiento de la población, último periodo", f"='{b1['hoja']}'!D{b1['ultima']}"),
             ("Saldo con el extranjero, último periodo", f"='{b1['hoja']}'!F{b1['ultima']}"),
             ("Saldo vegetativo, último periodo", f"='{b1['hoja']}'!E{b1['ultima']}"),
             ("Hogares nuevos 2021-2024 (provincias)", f"=SUM('{b4['hoja']}'!B{b4['primera']}:B{b4['ultima']})"),
             ("Obra nueva neta Catastro 2021-2024 (provincias con dato)", f"=SUM('{b4['hoja']}'!C{b4['primera']}:C{b4['ultima']})"),
             ("Viviendas de titulares con > 10 (España, %)", f"='{b5['hoja']}'!{letra(len(t.columns) + 1)}{b5['primera']}"),
             ("Viviendas de personas jurídicas (España, %)", f"='{b5['hoja']}'!{letra(len(t.columns) + 3)}{b5['primera']}"),
             ("Crecimiento de hogares 2011-2021 (suma de provincias)", f"=SUM('{b3['hoja']}'!{letra(len(d.columns) + 1)}{b3['primera']}:{letra(len(d.columns) + 1)}{b3['ultima']})")]
    for i, (et, f) in enumerate(filas, start=4):
        res.cell(row=i, column=1, value=et)
        res.cell(row=i, column=2, value=f).font = VERDE
    del b6, b7
    fu = wb.create_sheet("Fuentes")
    fu["A1"] = "Fuentes y decisiones: docs/v12/decisiones.md; cifras publicables: output/v5/cifras_clave.csv (ids V12-)."
    for ws in wb.worksheets:
        ws.freeze_panes = None
    spec = importlib.util.spec_from_file_location("tablas_v10", RAIZ / "src/v10/tablas_excel.py")
    tx = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(tx)
    wb.properties.creator = "Borja Romero"
    wb.properties.created = wb.properties.modified = tx.FECHA_FIJA
    dest = O / "tablas_v12.xlsx"
    wb.save(dest)
    tx.fijar_zip(dest)   # fechas fijas en el zip: dos ejecuciones dan el mismo sha256
    return dest


def main() -> int:
    print(atlas(), tablas())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
