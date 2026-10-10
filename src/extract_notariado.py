"""Extraccion Notariado (Consejo General del Notariado + Colegio Notarial de Valencia).

Escalera (se documenta en docs/fallidas/notariado.md):
  1. API/visor: el Portal Estadistico del Notariado (penotariado.com) exige login para /private/statistics;
     el visor CIEN abierto (notariado.org/liferay/web/cien/estadisticas-al-completo) ya no publica compraventa
     de vivienda. -> no recuperable por API abierta.
  2. XLSX abierto: anexo del informe semestral "Compraventa de vivienda por extranjeros" del CGN (2S25),
     publicado en la nota de prensa del CIEN (origen = xls).
  3. PDF con texto (pdfplumber): estadisticas del Colegio Notarial de Valencia / Centro Tecnologico del
     Notariado (extranjeros por provincia y nacionalidad, trimestral; actos de compraventa mensuales por
     provincia; extranjeros por municipio, anual) (origen = pdf).

Salidas (data/raw/pdf/):
  notariado_cgn_extranjeros_semestral.csv   xls   Espana/CCAA/nacionalidad, semestral 2007-2025
  notariado_cv_prov_trimestral.csv          pdf   Valencia/Alicante/Castellon, trimestral 2018T1-2025T4
  notariado_cv_actos_mensual.csv            pdf   actos de compraventa de inmuebles por provincia, mensual
  notariado_cv_municipios_anual.csv         pdf   viviendas compradas por extranjeros por municipio y nacionalidad
  notariado_validacion.csv                  cuadres y contraste con INE / MIVAU
Originales sin editar en data/raw/pdf/originales/. Renders PNG para revision visual en data/raw/pdf/validacion/.
Idempotente (descargas con cache). No se corrige ningun valor: lo que no cuadra se marca.

Cache a nivel de salida: si los 5 CSV de salida existen y no hay FORCE=1, el script termina sin red ni extraccion
(find_doc hace peticiones HTTP; con FORCE=1 se rehace todo).

Municipios (parse_muni): las columnas de cada pagina se detectan con la geometria real de las palabras (x0 de la
primera palabra tras la cabecera '... - hasta 4T AAAA' de la columna derecha), no con un corte fijo por la mitad de la
pagina (que dejaba fuera a Valencia ciudad y otros 5 municipios en 4T2022 y 4T2025). CONTROL DE COMPLETITUD: el
conjunto esperado de municipios (union de todas las ediciones + nombres de cabecera del texto completo de cada pagina)
debe coincidir con el extraido en cada edicion; cualquier faltante pone validado=no y se lista en 'detalle'.
"""
from __future__ import annotations

import html
import os
import re
import subprocess
import sys
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
import pdfplumber

sys.path.insert(0, str(Path(__file__).resolve().parent))
from utils_fetch import HEADERS, ROOT, download  # noqa: E402

import requests  # noqa: E402

PDF_DIR = ROOT / "data" / "raw" / "pdf"
ORIG = PDF_DIR / "originales"
VALID_DIR = PDF_DIR / "validacion"
RAW = ROOT / "data" / "raw"
FUENTE_CGN = "Consejo General del Notariado - CIEN"
FUENTE_CNV = "Colegio Notarial de Valencia / Centro Tecnologico del Notariado"
URL_CGN_NEWS = ("https://www.notariado.org/liferay/web/cien/sala-de-prensa/noticias/detalle?p_p_id=CIEN113_WAR_"
                "cienPrensaPlugin&p_p_lifecycle=0&p_p_col_id=column-1&p_p_col_count=1&p_r_p_564233524_NOTARIO_"
                "INFORMA_DETALLE_ID=32452895")
URL_CGN_XLSX = "https://www.notariado.org/liferay/c/document_library/get_file?uuid=125f548a-e4f2-498c-b9ea-3a736e10941a&groupId=2289837"
URL_CGN_PDF = "https://www.notariado.org/liferay/c/document_library/get_file?uuid=14a61d04-68ba-4929-b339-a158b344a0e1&groupId=2289837"
CNV = "https://valencia.notariado.org"
NEWS = CNV + "/portal/noticias/-/asset_publisher/V3kCUm4K8nMX/content/id/{id}"

# --- Revision visual realizada (paginas PNG leidas y comparadas con las tablas extraidas) ---
PAGINAS_REVISADAS = {
    "notariado_cv_prov_trimestral": "notariado_cv_val_extranjeros_4T2025.pdf p2,p10(zoom); notariado_cv_cas_extranjeros_4T2025.pdf p2; notariado_cv_ali_extranjeros_4T2025.pdf p10 (solo estructura, baja resolucion)",
    "notariado_cv_actos_mensual": "notariado_cv_actos_2024-2025.pdf p3; notariado_cv_actos_1T2026.pdf p1",
    "notariado_cv_municipios_anual": "notariado_cv_val_municipios_extranjeros_4T2025.pdf p2; notariado_cv_val_municipios_extranjeros_4T2022.pdf p2 (Valencia ciudad, columna izquierda y derecha)",
}

# ediciones de actos mensuales: (clave, id noticia Colegio Notarial)
ACTOS = [("2019-2020", 1062614), ("2020-2021", 1647219), ("2021-2022", 2433416), ("2022-2023", 3270801),
         ("2023-2024", 4345817), ("2024-2025", 5024924), ("1T2026", 5184030)]
# informes trimestrales de extranjeros por provincia (4T 2025 contiene la serie 2018T1-2025T4)
EXT_NEWS_4T25 = 5024637
# municipios de la provincia de Valencia, ediciones anuales (4T de cada ano)
MUNI = [("4T2021", 1640243), ("4T2022", 2437969), ("4T2023", 3230383), ("4T2024", 4343088), ("4T2025", 5024637)]
PROV = {"Valencia": "val", "Alicante": "ali", "Castellón": "cas"}


# ------------------------------------------------------------------ utilidades
def num_es(s: str) -> float:
    """'7.064' -> 7064 ; '96.481,81 €' -> 96481.81"""
    s = s.replace("€", "").strip().replace(".", "").replace(",", ".")
    return float(s)


def known_names(pdf) -> list[str]:
    """Nombres canonicos de pais (p7: ranking anual, texto limpio) para reponer etiquetas de celdas de varias lineas."""
    names = []
    for line in pdf.pages[6].extract_text().splitlines():
        m = re.match(r"^([^\d]+?) ((?:[\d.]+ ){8}[\d.]+)", line.strip())
        if m:
            names.append(" ".join(m.group(1).split()))
    return names + ["Otras nacionalidades"]


def fix_label(s: str, names: list[str]) -> str:
    """En celdas de etiqueta con salto de linea la extraccion entrelaza las letras de las dos lineas (no se pierde
    ni cambia ningun caracter, solo el orden). Se repone el nombre canonico por igualdad de multiconjunto de letras."""
    t = " ".join(str(s).split())
    key = sorted(re.sub(r"\W", "", t).lower())
    for n in names:
        if sorted(re.sub(r"\W", "", n).lower()) == key:
            return n
    # celda recortada (la 3a linea del nombre no cabe): las letras visibles son un submultiset de un unico nombre
    from collections import Counter
    c = Counter(key)
    cand = [n for n in names if not (c - Counter(re.sub(r"\W", "", n).lower()))]
    cand = [n for n in cand if len(key) >= 0.6 * len(re.sub(r"\W", "", n))]
    return cand[0] if len(cand) == 1 else t


def links(url: str) -> list[tuple[str, str]]:
    """Enlaces (texto, ruta) a /documents/ de una noticia del Colegio Notarial."""
    r = requests.get(url, headers=HEADERS, timeout=60)
    r.raise_for_status()
    out = []
    for m in re.finditer(r'<a [^>]*href="([^"]+)"[^>]*>(.*?)</a>', r.text, re.S):
        u, tx = m.groups()
        if "/documents/" in u:
            out.append((html.unescape(re.sub(r"<[^>]+>", "", tx)).strip(), html.unescape(u)))
    return out


def find_doc(news_id: int, must: list[str], mustnot: list[str] = ()) -> str:
    for tx, u in links(NEWS.format(id=news_id)):
        low = (tx + " " + u).lower()
        if all(re.search(m, low) for m in must) and not any(re.search(m, low) for m in mustnot):
            # el CMS exige las rutas tal cual; se completa si la ruta carece de dominio
            return u if u.startswith("http") else CNV + u
    raise RuntimeError(f"documento no encontrado en noticia {news_id}: {must}")


def first_of_period(sem: str) -> tuple[str, str]:
    """'1S07' -> ('2007-01-01', '2007S1')"""
    h, yy = int(sem[0]), int(sem[2:])
    y = 2000 + yy
    return f"{y}-{'01' if h == 1 else '07'}-01", f"{y}S{h}"


def quarter(label: str) -> tuple[str, str]:
    """'TRIM 3 2018' -> ('2018-07-01', '2018T3')"""
    q, y = re.match(r"TRIM (\d) (\d{4})", label).groups()
    return f"{y}-{(int(q)-1)*3+1:02d}-01", f"{y}T{q}"


MES = {m: i + 1 for i, m in enumerate("ENE FEB MAR ABR MAY JUN JUL AGO SEP OCT NOV DIC".split())}


# ------------------------------------------------------------------ 1) XLSX CGN
def parse_xlsx(path: Path) -> pd.DataFrame:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    keep = {"TABLA 1": ("op_viv_libre", "operaciones"), "TABLA 1C": ("precio_m2", "EUR/m2"),
            "TABLA 2": ("op_viv_libre_extranjeros", "operaciones"), "TABLA 2B": ("precio_m2_extranjeros", "EUR/m2"),
            "TABLA 3": ("op_extranjeros_nacionalidad", "operaciones"), "TABLA 3B": ("precio_m2_extranjeros_nacionalidad", "EUR/m2"),
            "TABLA 4": ("op_ext_residentes_nacionalidad", "operaciones"), "TABLA 4B": ("precio_m2_ext_residentes_nacionalidad", "EUR/m2"),
            "TABLA 5": ("op_ext_no_residentes_nacionalidad", "operaciones"), "TABLA 5B": ("precio_m2_ext_no_residentes_nacionalidad", "EUR/m2")}
    rows = []
    for ws in wb.worksheets[1:]:
        periods = None
        cur = None
        group = ""
        for row in ws.iter_rows():
            cells = {c.column: c.value for c in row if c.value is not None}
            if not cells:
                continue
            vals = {c.column: c.value for c in row}
            lab_cell = next((c for c in row if c.column == 5), None)
            lab = lab_cell.value if lab_cell is not None else None
            if lab == "SEMESTRE":
                periods = {c.column: c.value for c in row if isinstance(c.value, str) and re.fullmatch(r"[12]S\d\d", c.value)}
                continue
            if isinstance(lab, str) and re.match(r"TABLA \d[A-D]?:", lab):
                cur = lab.split(":")[0].strip().upper()
                group = ""
                continue
            if cur not in keep or periods is None:
                continue
            nums = {col: v for col, v in vals.items() if col in periods and isinstance(v, (int, float))}
            if not nums:
                continue
            label = lab if isinstance(lab, str) else "SIN_ETIQUETA_EN_ORIGEN"
            if cur in ("TABLA 1", "TABLA 1C") and label in ("Extranjero", "Nacional", "Total general"):
                group = label
            for col, v in nums.items():
                f, p = first_of_period(periods[col])
                serie, unidad = keep[cur]
                if cur in ("TABLA 1", "TABLA 1C"):
                    cat = label if label in ("Extranjero", "Nacional", "Total general") else f"{group} - {label}"
                    terr, nac = "Espana", ""
                elif cur in ("TABLA 2", "TABLA 2B"):
                    cat, terr, nac = "Extranjero", label, ""
                else:
                    cat = {"TABLA 3": "Extranjero", "TABLA 3B": "Extranjero", "TABLA 4": "Extranjero residente",
                           "TABLA 4B": "Extranjero residente"}.get(cur, "Extranjero no residente")
                    terr, nac = "Espana", label
                rows.append(dict(fecha=f, periodo=p, serie=serie, valor=float(v), unidad=unidad, fuente=FUENTE_CGN,
                                 url=URL_CGN_XLSX, origen="xls", pagina=ws.title, tabla=cur.replace("TABLA ", "T"),
                                 territorio=terr, categoria=cat, nacionalidad=nac))
    return pd.DataFrame(rows)


# ------------------------------------------------------------------ 2) PDF provincia (trimestral)
def parse_prov_pdf(path: Path, prov: str, url: str) -> pd.DataFrame:
    rows = []
    with pdfplumber.open(path) as pdf:
        # p2: viviendas vendidas segun tipo de comprador; p3: cuantia media
        for pno, serie, unidad, rx in [(2, "viv_vendidas", "viviendas", r"^(?:\d{4} )?(TRIM \d \d{4}) ([\d.]+) ([\d.]+)$"),
                                       (3, "cuantia_media", "EUR", r"^(?:\d{4} )?(TRIM \d \d{4}) ([\d.,]+) € ([\d.,]+) ?€?$")]:
            txt = pdf.pages[pno - 1].extract_text()
            for line in txt.splitlines():
                m = re.match(rx, line.strip())
                if not m:
                    # las lineas "Total" anuales se usan en validacion, no se guardan como observacion trimestral
                    continue
                f, p = quarter(m.group(1))
                for tipo, g in (("Españoles", 2), ("Extranjeros", 3)):
                    rows.append(dict(fecha=f, periodo=p, serie=f"{serie}_{'esp' if tipo == 'Españoles' else 'ext'}",
                                     valor=num_es(m.group(g)), unidad=unidad, territorio=prov, nacionalidad="",
                                     origen="pdf", pagina=pno))
        # ultima tabla de nacionalidad (p10): conteo y cuantia, trimestral
        names = known_names(pdf)
        for tb in pdf.pages[9].extract_tables():
            if len(tb) < 5 or len(tb[0]) < 30:
                continue
            title = (tb[0][1] or "").strip()
            hdr = tb[1]
            if "Viviendas Vendidas" in title:
                serie, unidad, parse = "viv_vendidas_pais", "viviendas", num_es
            elif "Cuant" in title:
                serie, unidad, parse = "cuantia_media_pais", "EUR", num_es
            else:
                continue
            for r in tb[2:]:
                lab = fix_label(r[0] or "", names)
                for j in range(1, len(hdr) - 1):          # la ultima columna es 'Total'
                    v = r[j]
                    if v is None or not str(v).strip():
                        continue                          # celda vacia en origen: no hay dato (no se imputa)
                    f, p = quarter(hdr[j])
                    rows.append(dict(fecha=f, periodo=p, serie=serie, valor=parse(v), unidad=unidad,
                                     territorio=prov, nacionalidad=lab, origen="pdf", pagina=10))
    df = pd.DataFrame(rows)
    df["fuente"], df["url"] = FUENTE_CNV, url
    return df


def prov_annual_totals(path: Path) -> pd.DataFrame:
    """Filas 'Total' anuales de p2 (sumas que publica el propio PDF), para cuadrar."""
    out = []
    with pdfplumber.open(path) as pdf:
        txt = pdf.pages[1].extract_text().splitlines()
    year = None
    seq = []
    for line in txt:
        m = re.match(r"^(?:(\d{4}) )?TRIM \d (\d{4}) ", line.strip())
        if m:
            year = int(m.group(2))
        m = re.match(r"^Total ([\d.]+) ([\d.]+)$", line.strip())
        if m:
            seq.append((year, num_es(m.group(1)), num_es(m.group(2))))
    return pd.DataFrame(seq, columns=["anio", "esp", "ext"]).drop_duplicates("anio", keep="first")  # la ultima fila Total es el acumulado 2018-2025


# ------------------------------------------------------------------ 3) PDF actos mensuales
def parse_actos(path: Path, edicion: str, url: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows, tots = [], []
    with pdfplumber.open(path) as pdf:
        for pno, pg in enumerate(pdf.pages, 1):
            for line in pg.extract_text().splitlines():
                line = line.strip()
                m = re.match(r"^(Alicante|Castellón|Valencia) (ENE|FEB|MAR|ABR|MAY|JUN|JUL|AGO|SEP|OCT|NOV|DIC) (\d{4}) ([\d.]+)$", line)
                if m:
                    prov, mes, y, v = m.groups()
                    rows.append(dict(fecha=f"{y}-{MES[mes]:02d}-01", periodo=f"{y}M{MES[mes]:02d}",
                                     serie="actos_compraventa_inmuebles", valor=num_es(v), unidad="actos",
                                     territorio=prov, nacionalidad="", origen="pdf", pagina=pno, edicion=edicion))
                    continue
                m = re.match(r"^TOTAL (ALICANTE|CASTELLÓN|VALENCIA) (?:(1T) )?(\d{4}) ([\d.]+)$", line)
                if m:
                    prov, t1, y, v = m.groups()
                    tots.append(dict(edicion=edicion, territorio=prov.capitalize(), anio=int(y), trimestre1=bool(t1),
                                     total_pdf=num_es(v), pagina=pno))
    df = pd.DataFrame(rows)
    df["fuente"], df["url"] = FUENTE_CNV, url
    return df, pd.DataFrame(tots)


# ------------------------------------------------------------------ 4) PDF municipios
HDR_RX = re.compile(r"\s*- hasta \dT \d{4}\s*")


def mkey(name: str) -> str:
    """clave de comparacion de nombres de municipio (sin acentos/apostrofes/espacios)"""
    n = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", n)


# nombre canonico y codigo INE (5 digitos; verificados con SERPAVI hoja Municipios) por clave normalizada
MUNI_CANON = {"valencia": ("Valencia", "46250"), "oliva": ("Oliva", "46181"), "lliria": ("Llíria", "46147"),
              "gandia": ("Gandia", "46131"), "ontinyent": ("Ontinyent", "46184"), "villalonga": ("Villalonga", "46255"),
              "barx": ("Barx", "46046"), "palmadegandia": ("Palma de Gandía", "46187"), "saguntosagunt": ("Sagunto/Sagunt", "46220"),
              "xativa": ("Xàtiva", "46145"), "canetdenberenguer": ("Canet d'En Berenguer", "46082"), "pucol": ("Puçol", "46205"),
              "cullera": ("Cullera", "46105"), "tavernesdelavalldigna": ("Tavernes de la Valldigna", "46238"),
              "xeraco": ("Xeraco", "46143"), "xeresa": ("Xeresa", "46146"), "bellreguard": ("Bellreguard", "46048"),
              "miramar": ("Miramar", "46168"), "daimus": ("Daimús", "46113")}


def canon(name: str) -> str:
    k = mkey(name)
    if k not in MUNI_CANON:
        raise RuntimeError(f"municipio sin forma canonica ni codigo INE: {name!r}")
    return MUNI_CANON[k][0]


def text_headers(pg) -> list[str]:
    """Nombres de municipio segun el TEXTO COMPLETO de la pagina (independiente de la geometria de columnas)."""
    out = []
    for line in (pg.extract_text() or "").splitlines():
        if "- hasta" in line:
            out += [x.strip() for x in HDR_RX.split(line) if x.strip()]
    return out


def column_bounds(pg) -> list[tuple[float, float]]:
    """Columnas a partir de la geometria: tras cada cabecera '- hasta 4T AAAA' la palabra siguiente de la misma linea
    es el inicio (x0) de la columna derecha. Devuelve [(x_ini, x_fin)] de izquierda a derecha."""
    words = pg.extract_words()
    starts = []
    for k, w in enumerate(words[:-1]):
        if re.fullmatch(r"\d{4}", w["text"]) and k >= 3 and words[k - 2]["text"] == "hasta" and words[k - 1]["text"].endswith("T"):
            nxt = words[k + 1]
            if abs(nxt["top"] - w["top"]) < 3:
                starts.append(nxt["x0"])
    if not starts:
        return [(0, pg.width)]
    xr = min(starts)
    return [(0, xr - 5), (xr - 5, pg.width)]


def parse_muni(path: Path, edicion: str, url: str) -> tuple[pd.DataFrame, pd.DataFrame, list[str]]:
    rows, tots = [], []
    expected = []   # nombres de cabecera vistos en el texto completo de las paginas
    with pdfplumber.open(path) as pdf:
        for pno, pg in enumerate(pdf.pages, 1):
            if pno == 1:
                continue
            expected += text_headers(pg)
            for x0, x1 in column_bounds(pg):
                txt = pg.crop((x0, 0, x1, pg.height)).extract_text() or ""
                cur = None
                for line in txt.splitlines():
                    line = line.strip()
                    m = re.match(r"^(.+?) - hasta (\dT) (\d{4})$", line)
                    if m:
                        cur = m.group(1).strip()
                        year = int(m.group(3))
                        continue
                    if cur is None or line.startswith("Nacionalidad"):
                        continue
                    m = re.match(r"^Total general ([\d.]+)$", line)
                    if m:
                        tots.append(dict(edicion=edicion, municipio=cur, total_pdf=num_es(m.group(1)), pagina=pno))
                        cur = None
                        continue
                    m = re.match(r"^(.+?) ([\d.]+)$", line)
                    if m:
                        rows.append(dict(fecha=f"{year}-01-01", periodo=str(year), serie="viv_compradas_extranjeros_nacionalidad",
                                         valor=num_es(m.group(2)), unidad="viviendas", territorio=cur,
                                         nacionalidad=m.group(1).strip(), origen="pdf", pagina=pno, edicion=edicion))
    # fila 'Total general' publicada (cifra oficial) como nacionalidad='Total general'
    for t in tots:
        rows.append(dict(fecha=f"{year}-01-01", periodo=str(year), serie="viv_compradas_extranjeros_nacionalidad",
                         valor=t["total_pdf"], unidad="viviendas", territorio=t["municipio"], nacionalidad="Total general",
                         origen="pdf", pagina=t["pagina"], edicion=edicion))
    df = pd.DataFrame(rows)
    df["territorio"] = df["territorio"].map(canon)
    df["cod_ine"] = df["territorio"].map(lambda n: MUNI_CANON[mkey(n)][1])
    tots = pd.DataFrame(tots)
    tots["municipio"] = tots["municipio"].map(canon)
    # aditiva: suma de nacionalidades == Total general (por municipio x edicion)
    sm = df[df.nacionalidad != "Total general"].groupby("territorio")["valor"].sum()
    tt = tots.set_index("municipio")["total_pdf"]
    df["aditiva"] = df["territorio"].map(lambda m: bool(sm.get(m, np.nan) == tt.get(m, np.nan)))
    df["fuente"], df["url"] = FUENTE_CNV, url
    return df, tots, expected


# ------------------------------------------------------------------ validacion
def ine_q(prov_label: str, kind: str = "General") -> pd.Series:
    d = pd.read_csv(RAW / "ine_etdp_compraventas.csv")
    d = d[d["nombre"].str.startswith(f"{prov_label}. {kind}. Compraventa")]
    d["periodo_q"] = pd.PeriodIndex(pd.to_datetime(d["fecha"]), freq="Q").astype(str).str.replace("Q", "T")
    return d.groupby("periodo_q")["valor"].sum()


def mivau(file: str, serie: str) -> pd.Series:
    d = pd.read_csv(RAW / file)
    d = d[d["serie"] == serie]
    return d.set_index("periodo")["valor"]


def rel_err(a: pd.Series, b: pd.Series) -> tuple[float, float, int]:
    j = pd.concat([a, b], axis=1, join="inner").dropna()
    if j.empty:
        return np.nan, np.nan, 0
    ab = (j.iloc[:, 0] - j.iloc[:, 1]).abs()
    return float(ab.max()), float((ab / j.iloc[:, 1].abs()).max()), len(j)


def to_sem(q: pd.Series) -> pd.Series:
    """trimestres 'YYYYTq' -> semestres 'YYYYSh' (suma)"""
    idx = [f"{p[:4]}S{1 if int(p[-1]) <= 2 else 2}" for p in q.index]
    s = q.groupby(idx).agg(["sum", "count"])
    return s[s["count"] == 2]["sum"]


OUTPUTS = ["notariado_cgn_extranjeros_semestral.csv", "notariado_cv_prov_trimestral.csv", "notariado_cv_actos_mensual.csv",
           "notariado_cv_municipios_anual.csv", "notariado_validacion.csv"]


def main() -> None:
    if all((PDF_DIR / f).exists() for f in OUTPUTS) and os.environ.get("FORCE") != "1":
        print("[cache] salidas de notariado ya existen; sin red (FORCE=1 para rehacer)")
        return
    ORIG.mkdir(parents=True, exist_ok=True)
    VALID_DIR.mkdir(parents=True, exist_ok=True)
    val = []   # filas de validacion

    def V(tabla, ok, err, metodo, pags, detalle="", err_rel=np.nan):
        val.append(dict(tabla=tabla, validado="si" if ok else "no", error_max=err, error_max_rel=err_rel,
                        metodo=metodo, paginas_revisadas=pags, detalle=detalle))

    # ---------- 1) XLSX CGN
    xlsx = download(URL_CGN_XLSX, ORIG / "notariado_cgn_extranjeros_2S25_anexo.xlsx")
    download(URL_CGN_PDF, ORIG / "notariado_cgn_extranjeros_2S25_informe.pdf")
    cgn = parse_xlsx(xlsx)
    cgn.to_csv(PDF_DIR / "notariado_cgn_extranjeros_semestral.csv", index=False)
    print("[ok] notariado_cgn_extranjeros_semestral.csv", len(cgn))

    t1 = cgn[cgn.tabla == "T1"].pivot(index="periodo", columns="categoria", values="valor")
    e1 = max((t1["Extranjero"] - t1["Extranjero - No residente"] - t1["Extranjero - Residente"]).abs().max(),
             (t1["Nacional"] - t1["Nacional - No residente"] - t1["Nacional - Residente"]).abs().max(),
             (t1["Total general"] - t1["Extranjero"] - t1["Nacional"]).abs().max())
    V("cgn_T1_residencia", e1 == 0, e1, "xls: ext=res+nores, nac=res+nores, total=ext+nac", "n/a (xlsx)")
    t2 = cgn[cgn.tabla == "T2"].pivot(index="periodo", columns="territorio", values="valor")
    nac = t2["Nacional"]
    ccaa_sum = t2.drop(columns=["Nacional"]).sum(axis=1)
    e2 = (ccaa_sum - nac).abs().max()
    V("cgn_T2_ccaa_suma_vs_nacional", e2 == 0, e2, "xls: suma 17 CCAA + fila sin etiqueta = Nacional", "n/a (xlsx)",
      "la fila sin etiqueta del original (entre Cataluna y C. Valenciana; valores pequenos) se conserva como SIN_ETIQUETA_EN_ORIGEN")
    for t, nm in (("T3", "ext"), ("T4", "ext_res"), ("T5", "ext_nores")):
        x = cgn[cgn.tabla == t].pivot(index="periodo", columns="nacionalidad", values="valor")
        tot = x["Total general"]
        s = x.drop(columns=["Total general"]).sum(axis=1)
        e = (s - tot).abs().max()
        V(f"cgn_{t}_nacionalidad_suma_vs_total", e == 0, e, "xls: suma de nacionalidades = Total general (cada hoja)", "n/a (xlsx)")
    # contraste con MIVAU (misma fuente notarial, otra desagregacion): extranjeros CCAA Valencia, Espana
    resid = pd.read_csv(RAW / "mivau_transacciones_residencia.csv")
    for terr, nombre in (("comunitat_valenciana", "Comunidad Valenciana"),):
        a = resid[resid.serie == f"tx_residencia_residentes_extranjeros_ccaa_{terr}"].set_index("periodo")["valor"]
        b = resid[resid.serie == f"tx_residencia_no_residentes_extranjeros_ccaa_{terr}"].set_index("periodo")["valor"]
        mv = to_sem(a + b)
        cg = cgn[(cgn.tabla == "T2") & (cgn.territorio == nombre)].set_index("periodo")["valor"]
        ea, er, n = rel_err(cg, mv)
        V("cgn_T2_CV_extranjeros_vs_MIVAU_residencia", er < 0.15, ea,
          f"comparacion con mivau_transacciones_residencia (res+no res extranjeros, suma 2 trimestres), n={n} semestres; CGN=vivienda libre, MIVAU=todas",
          "n/a (xlsx)", "diferencia de definicion (libre vs total)", er)
    a = resid[resid.serie == "tx_residencia_residentes_extranjeros_nacional"].set_index("periodo")["valor"] if \
        (resid.serie == "tx_residencia_residentes_extranjeros_nacional").any() else None
    nat = sorted(s for s in resid.serie.unique() if s.endswith("_nacional"))
    if nat:
        ext_n = [s for s in nat if "extranjeros" in s and "total" not in s]
        mvn = sum(resid[resid.serie == s].set_index("periodo")["valor"] for s in ext_n if "no_consta" not in s)
        cgn_n = cgn[(cgn.tabla == "T1") & (cgn.categoria == "Extranjero")].set_index("periodo")["valor"]
        ea, er, n = rel_err(cgn_n, to_sem(mvn))
        V("cgn_T1_Espana_extranjeros_vs_MIVAU_residencia", er < 0.15, ea,
          f"comparacion con {ext_n} (suma trimestres), n={n} semestres", "n/a (xlsx)", "libre vs total", er)

    # ---------- 2) PDF provincias (4T 2025)
    prov_frames, prov_tot = [], {}
    for prov, key in PROV.items():
        u = find_doc(EXT_NEWS_4T25, [{"Valencia": r"extranjeros en valencia", "Alicante": r"extranjeros en alicante",
                                      "Castellón": r"extranjeros en castell"}[prov]], [r"municip", r"brit"])
        # en el listado, el enlace de provincia no incluye 'municipios' ni 'británicos'
        path = download(u, ORIG / f"notariado_cv_{key}_extranjeros_4T2025.pdf")
        prov_frames.append(parse_prov_pdf(path, prov, u))
        prov_tot[prov] = (path, prov_annual_totals(path))
    prov_df = pd.concat(prov_frames, ignore_index=True)
    prov_df["edicion"] = "4T2025"
    prov_df.to_csv(PDF_DIR / "notariado_cv_prov_trimestral.csv", index=False)
    print("[ok] notariado_cv_prov_trimestral.csv", len(prov_df))
    COLS = ["fecha", "periodo", "serie", "valor", "unidad", "fuente", "url", "origen", "pagina"]

    for prov, (path, tot) in prov_tot.items():
        d = prov_df[prov_df.territorio == prov].copy()
        d["anio"] = d["periodo"].str[:4].astype(int)
        # (a) trimestres suman al total anual publicado
        errs = []
        for k, col in (("viv_vendidas_esp", "esp"), ("viv_vendidas_ext", "ext")):
            s = d[d.serie == k].groupby("anio")["valor"].sum()
            m = tot.set_index("anio")[col]
            errs.append((s - m.reindex(s.index)).abs().max())
        V(f"prov_{prov}_trimestres_vs_total_anual", max(errs) == 0, max(errs), "suma trimestral = fila Total anual publicada (p2)",
          PAGINAS_REVISADAS["notariado_cv_prov_trimestral"])
        # (b) nacionalidades (incluida 'Otras') suman a la fila Total y esta a extranjeros p2
        c = d[d.serie == "viv_vendidas_pais"].pivot_table(index="periodo", columns="nacionalidad", values="valor", aggfunc="sum")
        tot_row = c["Total"]
        s = c.drop(columns=["Total"]).sum(axis=1)
        e_a = (s - tot_row).abs().max()
        ext = d[d.serie == "viv_vendidas_ext"].set_index("periodo")["valor"]
        e_b = (tot_row - ext).abs().max()
        V(f"prov_{prov}_nacionalidades_suma_vs_total", e_a == 0 and e_b == 0, max(e_a, e_b),
          f"suma de nacionalidades=fila Total (err {e_a:g}); fila Total = extranjeros p2 (err {e_b:g})",
          PAGINAS_REVISADAS["notariado_cv_prov_trimestral"])
        # (c) cuantia: fila Total p10 = extranjeros p3
        q = d[d.serie == "cuantia_media_pais"].pivot_table(index="periodo", columns="nacionalidad", values="valor", aggfunc="sum")
        cm = d[d.serie == "cuantia_media_ext"].set_index("periodo")["valor"]
        e_c = (q["Total"] - cm).abs().max()
        V(f"prov_{prov}_cuantia_total_p10_vs_p3", e_c < 0.01, e_c, "cuantia media extranjeros: fila Total p10 = p3",
          PAGINAS_REVISADAS["notariado_cv_prov_trimestral"])

    # contrastes con INE / MIVAU (viviendas totales = esp + ext, que en origen puede solapar si el acto mixto cuenta en ambos)
    for prov, ine_name, mv_serie in (("Valencia", "Valencia/València", "tx_total_provincia_valencia_valencia"),):
        d = prov_df[prov_df.territorio == prov]
        tot_q = d[d.serie == "viv_vendidas_esp"].set_index("periodo")["valor"].add(
            d[d.serie == "viv_vendidas_ext"].set_index("periodo")["valor"], fill_value=0)
        ine = ine_q(ine_name)
        mv = mivau("mivau_transacciones_total.csv", mv_serie)
        ea, er, n = rel_err(tot_q, ine)
        V(f"prov_{prov}_viv_total_vs_INE_ETDP", er < 0.10, ea, f"esp+ext (notarial) vs INE ETDP 'General' provincia, suma mensual->trimestre, n={n}",
          PAGINAS_REVISADAS["notariado_cv_prov_trimestral"],
          "NO CUADRA: INE ETDP cuenta inscripciones registrales (retraso y estacionalidad propios, ratio notarial/INE 0.95-1.49); esp+ext notarial solapa ~4-6% en actos mixtos", er)
        ea, er, n = rel_err(tot_q, mv)
        V(f"prov_{prov}_viv_total_vs_MIVAU_tx_total", er < 0.10, ea, f"esp+ext (notarial) vs MIVAU transacciones total provincia (misma fuente: notarios), n={n}",
          PAGINAS_REVISADAS["notariado_cv_prov_trimestral"], "", er)
    # extranjeros provincia Valencia vs MIVAU residencia extranjeros (res + no res)
    d = prov_df[(prov_df.territorio == "Valencia") & (prov_df.serie == "viv_vendidas_ext")].set_index("periodo")["valor"]
    a = resid[resid.serie == "tx_residencia_residentes_extranjeros_provincia_valencia_valencia"].set_index("periodo")["valor"]
    b = resid[resid.serie == "tx_residencia_no_residentes_extranjeros_provincia_valencia_valencia"].set_index("periodo")["valor"]
    ea, er, n = rel_err(d, a + b)
    V("prov_Valencia_extranjeros_vs_MIVAU_residencia", er < 0.15, ea, f"extranjeros prov. Valencia vs MIVAU (res+no res extranjeros), n={n}",
      PAGINAS_REVISADAS["notariado_cv_prov_trimestral"],
      "NO CUADRA estrictamente: ratio PDF/MIVAU estable 1.06-1.20 (sesgo sistematico, no ruido); el PDF cuenta viviendas con >=1 comprador extranjero, MIVAU solo residencia informada; usar como serie propia, no mezclar niveles", er)
    # CV agregada (suma 3 provincias) vs xlsx CGN CCAA Valencia (semestral)
    ext_cv = sum(prov_df[(prov_df.territorio == p) & (prov_df.serie == "viv_vendidas_ext")].set_index("periodo")["valor"] for p in PROV)
    cg = cgn[(cgn.tabla == "T2") & (cgn.territorio == "Comunidad Valenciana")].set_index("periodo")["valor"]
    ea, er, n = rel_err(to_sem(ext_cv), cg)
    V("prov_CV_ext_suma3prov_vs_CGN_xlsx_CCAA", er < 0.10, ea,
      f"suma Valencia+Alicante+Castellon (PDF trimestral, a semestre) vs CGN xlsx 'Comunidad Valenciana' (vivienda libre), n={n}",
      PAGINAS_REVISADAS["notariado_cv_prov_trimestral"], "", er)

    # ---------- 3) actos mensuales
    act_frames, act_tots = [], []
    for ed, nid in ACTOS:
        u = find_doc(nid, [r"compraventa", r"inmuebles"], [r"hipotec", r"sociedad", r"herencia", r"donaci", r"testam", r"matrimon", r"divorc"])
        path = download(u, ORIG / f"notariado_cv_actos_{ed}.pdf")
        a, t = parse_actos(path, ed, u)
        act_frames.append(a)
        act_tots.append(t)
    act = pd.concat(act_frames, ignore_index=True)
    tots = pd.concat(act_tots, ignore_index=True)
    act.to_csv(PDF_DIR / "notariado_cv_actos_mensual.csv", index=False)
    print("[ok] notariado_cv_actos_mensual.csv", len(act))
    # (a) meses suman al TOTAL anual (o 1T)
    errs = []
    for _, r in tots.iterrows():
        s = act[(act.edicion == r.edicion) & (act.territorio.str.lower() == r.territorio.lower()) & (act.fecha.str[:4] == str(r.anio))]
        if r.trimestre1:
            s = s[s.fecha.str[5:7].isin(["01", "02", "03"])]
        errs.append(abs(s.valor.sum() - r.total_pdf))
    V("actos_meses_vs_TOTAL_publicado", max(errs) == 0, max(errs), f"suma mensual = fila TOTAL del PDF ({len(errs)} totales, 7 ediciones)",
      PAGINAS_REVISADAS["notariado_cv_actos_mensual"])
    # (b) revisiones entre ediciones solapadas
    ov = act.pivot_table(index=["territorio", "fecha"], columns="edicion", values="valor")
    ed_cols = [c for c in ov.columns]
    rev = []
    for i, c1 in enumerate(ed_cols):
        for c2 in ed_cols[i + 1:]:
            x = ov[[c1, c2]].dropna()
            if len(x):
                rev.append(((x[c1] - x[c2]).abs() / x[c2]).max())
    V("actos_revision_entre_ediciones", max(rev) < 0.02, np.nan, "cambio relativo maximo del mismo mes entre ediciones solapadas (revision de datos)",
      PAGINAS_REVISADAS["notariado_cv_actos_mensual"], "los valores difieren entre ediciones: se conservan todas (columna edicion); usar la mas reciente", max(rev))
    # (c) contraste con INE ETDP provincia (viviendas) mensual: la edicion mas reciente de cada mes
    # orden cronologico de ediciones (v5, BK-045): el orden de texto ponia '1T2026' antes que '2019-2020'
    orden_ed = {k: i for i, (k, _) in enumerate(ACTOS)}
    last = act.assign(_o=act.edicion.map(orden_ed)).sort_values("_o").groupby(["territorio", "fecha"]).tail(1).drop(columns="_o")
    for prov, ine_name in (("Valencia", "Valencia/València"), ("Alicante", "Alicante/Alacant"), ("Castellón", "Castellón/Castelló")):
        d = pd.read_csv(RAW / "ine_etdp_compraventas.csv")
        d = d[d["nombre"].str.startswith(f"{ine_name}. General. Compraventa")].set_index("fecha")["valor"]
        n_ = last[last.territorio == prov].set_index("fecha")["valor"]
        ea, er, n = rel_err(n_, d)
        V(f"actos_{prov}_vs_INE_ETDP_mensual", False, ea,
          f"actos notariales de compraventa de inmuebles (no solo vivienda) vs INE compraventas de viviendas (registrales), n={n} meses",
          PAGINAS_REVISADAS["notariado_cv_actos_mensual"], "conceptos distintos (inmuebles vs viviendas; escritura vs inscripcion): solo orden de magnitud, no cuadra", er)

    # ---------- 4) municipios
    mu_frames, mu_tots, mu_exp = [], [], {}
    for ed, nid in MUNI:
        u = find_doc(nid, [r"extranjeros", r"valencia", r"municip"], [r"brit", r"alicante", r"castell"])
        path = download(u, ORIG / f"notariado_cv_val_municipios_extranjeros_{ed}.pdf")
        a, t, ex = parse_muni(path, ed, u)
        mu_frames.append(a)
        mu_tots.append(t)
        mu_exp[ed] = ex
    mu = pd.concat(mu_frames, ignore_index=True)
    mt = pd.concat(mu_tots, ignore_index=True)
    mu.to_csv(PDF_DIR / "notariado_cv_municipios_anual.csv", index=False)
    print("[ok] notariado_cv_municipios_anual.csv", len(mu))
    PM = PAGINAS_REVISADAS["notariado_cv_municipios_anual"]
    s = mu[mu.nacionalidad != "Total general"].groupby(["edicion", "territorio"])["valor"].sum().reset_index().merge(mt, left_on=["edicion", "territorio"], right_on=["edicion", "municipio"])
    e = (s["valor"] - s["total_pdf"]).abs()
    nad = mu[["edicion", "territorio", "aditiva"]].drop_duplicates()
    nad = nad[~nad.aditiva]
    V("municipios_no_aditivos_lista", nad.empty, float(len(nad)),
      "combinaciones municipio x edicion con suma de nacionalidades != Total general (columna 'aditiva' del CSV)", PM,
      "ninguna" if nad.empty else "; ".join(f"{r.edicion} {r.territorio}" for r in nad.itertuples()) +
      " (usar la fila nacionalidad='Total general', que es la cifra oficial)")
    V("municipios_19_territorios", mu.territorio.nunique() == 19 and mu.groupby("edicion").territorio.nunique().eq(19).all(),
      float(mu.territorio.nunique()), "19 territorios canonicos con cod_ine en cada edicion (nombres normalizados por mapeo explicito)", PM)
    V("municipios_nacionalidades_vs_Total_general", e.max() == 0 and len(s) == len(mt), e.max(),
      f"suma de nacionalidades = 'Total general' de cada municipio ({len(s)} bloques de {len(mt)} totales, {len(MUNI)} ediciones)", PM,
      "" if e.max() == 0 else "NO CUADRA EN ORIGEN: " + "; ".join(
          f"{r.edicion} {r.territorio}: suma {r.valor:.0f} vs Total general {r.total_pdf:.0f}" for r in s[e > 0].itertuples())
      + ". En Valencia 4T2022 la fila 'Otras nacionalidades' (995) no es aditiva: las demas nacionalidades suman exactamente 2.171 = Total "
        "(comprobado en el PNG de p2); se marca, no se corrige")
    # CONTROL DE COMPLETITUD: esperado = union de ediciones (texto completo) ; cada edicion debe tenerlos todos
    names = {}
    for ed, ex in mu_exp.items():
        for n in ex:
            names.setdefault(mkey(n), n)
    exp_all = set(names)
    falt, partes, hdr_ok = [], [], True
    for ed, _ in MUNI:
        got = set(mt[mt.edicion == ed].municipio.map(mkey))
        got_rows = set(mu[mu.edicion == ed].territorio.map(mkey))
        txt_ed = {mkey(n) for n in mu_exp[ed]}
        miss = (exp_all | txt_ed) - got
        miss_rows = (exp_all | txt_ed) - got_rows
        n_hdr = len(mu_exp[ed])
        if miss or miss_rows or n_hdr != len(mt[mt.edicion == ed]):
            hdr_ok = False
        falt += [f"{ed}:{names[k]}" for k in sorted(miss | miss_rows)]
        partes.append(f"{ed}: {len(got)}/{len(exp_all | txt_ed)} municipios ({n_hdr} cabeceras en texto, {len(mt[mt.edicion == ed])} totales)")
    V("municipios_completitud", hdr_ok and not falt, float(len(falt)),
      "municipios esperados (union de ediciones + cabeceras '- hasta' del texto completo de cada pagina) = extraidos en cada edicion",
      PM, "; ".join(partes) + ("" if not falt else " | FALTAN: " + ", ".join(falt)))
    # ciudad de Valencia <= provincia; falla si falta algun ano
    ann = prov_df[(prov_df.territorio == "Valencia") & (prov_df.serie == "viv_vendidas_ext")].assign(anio=lambda x: x.periodo.str[:4]).groupby("anio")["valor"].sum()
    vm = mt[mt.municipio.map(mkey) == "valencia"].assign(anio=lambda x: x.edicion.str[2:]).set_index("anio")["total_pdf"]
    anios = [ed[2:] for ed, _ in MUNI]
    j = pd.concat([vm.rename("ciudad"), ann.rename("prov")], axis=1).reindex(anios)
    ok = bool(j.notna().all().all() and (j.ciudad <= j.prov).all())
    V("municipios_Valencia_ciudad_le_provincia", ok, 0 if ok else float("nan"),
      "ciudad de Valencia <= provincia (extranjeros, anual), falla si falta algun ano: " + "; ".join(
          f"{k}: {a:.0f}<={b:.0f}" if pd.notna(a) else f"{k}: FALTA" for k, (a, b) in j.iterrows()), PM)
    # suma de municipios vs total provincial (no hay total provincial publicado en el PDF de municipios: se usa p2 del PDF de provincia)
    cov = []
    okc = True
    for ed, _ in MUNI:
        sm = mt[mt.edicion == ed].total_pdf.sum()
        pv = ann.get(ed[2:], np.nan)
        okc &= bool(pd.notna(pv) and sm <= pv)
        cov.append(f"{ed[2:]}: municipios {sm:.0f} / provincia {pv:.0f} = {sm / pv:.1%}")
    V("municipios_suma_vs_total_provincial", okc, 0 if okc else np.nan,
      "suma de 'Total general' de los 19 municipios <= extranjeros provincia Valencia (PDF provincia, anual); el PDF de municipios no publica total provincial (cobertura, no cuadre)",
      PM, "; ".join(cov))

    vdf = pd.DataFrame(val)
    vdf.to_csv(PDF_DIR / "notariado_validacion.csv", index=False)
    print(vdf[["tabla", "validado", "error_max", "error_max_rel"]].to_string())

    # ---------- renders PNG para revision visual
    import pypdfium2 as pdfium
    for fn, pages in (("notariado_cv_val_extranjeros_4T2025.pdf", [2, 3, 10]), ("notariado_cv_ali_extranjeros_4T2025.pdf", [10]),
                      ("notariado_cv_cas_extranjeros_4T2025.pdf", [2]), ("notariado_cv_actos_2024-2025.pdf", [3]),
                      ("notariado_cv_actos_1T2026.pdf", [1]),
                      ("notariado_cv_val_municipios_extranjeros_4T2025.pdf", [2]),
                      ("notariado_cv_val_municipios_extranjeros_4T2022.pdf", [2])):
        pdf = pdfium.PdfDocument(str(ORIG / fn))
        for p in pages:
            out = VALID_DIR / f"{Path(fn).stem}_p{p}.png"
            if not out.exists():
                pdf[p - 1].render(scale=1.6).to_pil().save(out)


if __name__ == "__main__":
    main()
